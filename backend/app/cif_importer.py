from __future__ import annotations

import hashlib
import csv
import json
import re
import threading
from datetime import datetime, timezone
from pathlib import Path

import xlrd
from openpyxl import load_workbook
from sqlalchemy import distinct, func, text

from app.database import SessionLocal
from app.models import (
    CifCustomer,
    CifCustomerIdentifier,
    CifIdentityConflict,
    CifImportBatch,
    CifImportError,
    CifSourceRecord,
)


REQUIRED_COLUMNS = {
    "custno",
    "nm",
    "nmloc",
    "custtpcd",
    "custdtltpcd",
    "regno",
    "stscd",
}
DATE_COLUMNS = ("issuedt1", "issuedt2", "issuedt3", "issuedt4", "issuedt5", "issuedt6", "incrdt")
IMPORTER_VERSION = "2.1-multi-branch"
REVIEW_FIELDS = (
    "customer_name", "customer_name_ascii", "short_name", "customer_type",
    "customer_detail_type", "registration_number", "passport_number",
    "driver_license_number", "tax_number", "telephone", "address_type",
    "full_address", "province", "district", "commune_ward", "nationality_code",
    "birth_date", "gender_code", "establishment_date", "occupation",
    "source_status", "normalized_status", "operator_user",
)
STATUS_MAP = {
    "bình thường": "active",
    "không hoạt động": "inactive",
    "không hợp lệ": "invalid",
    "không xác nhận": "unverified",
}
IDENTIFIER_COMPARE_FIELDS = (
    "customer_name",
    "customer_name_ascii",
    "short_name",
    "customer_type",
    "customer_detail_type",
    "registration_number",
    "passport_number",
    "driver_license_number",
    "tax_number",
    "telephone",
    "address_type",
    "full_address",
    "province",
    "district",
    "commune_ward",
    "nationality_code",
    "source_status",
    "normalized_status",
    "operator_user",
    "raw_data",
)
MASTER_IDENTITY_FIELDS = (
    "registration_number",
    "passport_number",
    "tax_number",
)


def clean(value) -> str:
    if value is None:
        return ""
    result = str(value).strip()
    if result.endswith(".0") and result[:-2].isdigit():
        result = result[:-2]
    return result


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def chunked(values: list[str], size: int = 5000):
    for start in range(0, len(values), size):
        yield values[start:start + size]


def parse_yyyymmdd(value: str) -> bool:
    if not re.fullmatch(r"\d{8}", value):
        return False
    try:
        parsed = datetime.strptime(value, "%Y%m%d").date()
    except ValueError:
        return False
    return 1900 <= parsed.year <= datetime.now().year


def date_yyyymmdd(value: str):
    if not parse_yyyymmdd(value):
        return None
    return datetime.strptime(value, "%Y%m%d").date()


def normalized_status(value: str) -> str:
    return STATUS_MAP.get(value.strip().lower(), "unknown")


def normalized_identity(value: str) -> str | None:
    result = re.sub(r"\s+", "", value or "").upper()
    return result or None


def full_address(row: dict) -> str | None:
    parts = [row.get("addr1loc"), row.get("addr2loc"), row.get("addr3loc")]
    result = ", ".join(part for part in parts if part)
    return result or None


def telephone(row: dict) -> str | None:
    primary = clean(row.get("name_4"))
    if primary:
        return primary
    parts = [row.get("telnoctry"), row.get("telnoarea"), row.get("telno"), row.get("telnoextn")]
    result = "".join(part for part in parts if part)
    return result or None


def infer_branch_code(filename: str) -> str | None:
    match = re.search(r"(?:^|[_\-])(\d{4})\s*CIF|CIF[_\-]?(\d{4})", filename, re.IGNORECASE)
    return next((value for value in match.groups() if value), None) if match else None


def row_hash(raw: dict) -> str:
    payload = json.dumps(raw, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def iter_cif_sources(path: Path):
    """Yield (sheet name, headers, row iterator) for CSV, XLS and XLSX."""
    suffix = path.suffix.lower()
    if suffix == ".xls":
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        if not workbook.sheet_names():
            raise ValueError("Workbook CIF không có sheet dữ liệu")
        for sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)
            headers = [clean(sheet.cell_value(0, column)).lower() for column in range(sheet.ncols)]
            rows = (
                (row_index + 1, [sheet.cell_value(row_index, column) for column in range(sheet.ncols)])
                for row_index in range(1, sheet.nrows)
            )
            yield sheet_name, headers, rows
            workbook.unload_sheet(sheet_name)
        return
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        if not workbook.sheetnames:
            raise ValueError("Workbook CIF không có sheet dữ liệu")
        for worksheet in workbook.worksheets:
            values = worksheet.iter_rows(values_only=True)
            first = next(values, None)
            if first is None:
                continue
            headers = [clean(value).lower() for value in first]
            rows = ((row_number, list(values_row)) for row_number, values_row in enumerate(values, 2))
            yield worksheet.title, headers, rows
        workbook.close()
        return
    if suffix == ".csv":
        encoding = "utf-8-sig"
        probe = path.open("rb").read(65536)
        try:
            probe.decode(encoding, errors="strict")
        except UnicodeDecodeError:
            encoding = "cp1258"
        source = path.open("r", encoding=encoding, newline="")
        sample = source.read(65536)
        source.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(source, dialect)
        first = next(reader, None)
        if first is None:
            source.close()
            raise ValueError("File CSV CIF không có dữ liệu")
        headers = [clean(value).lower() for value in first]
        rows = ((row_number, values_row) for row_number, values_row in enumerate(reader, 2))
        yield "CSV", headers, rows
        source.close()
        return
    raise ValueError("Kho CIF chỉ hỗ trợ CSV, XLS và XLSX")


def _refresh_customer_counts(db) -> None:
    db.execute(text("""
        WITH counts AS (
            SELECT customer_id, count(*) identifier_count,
                   count(DISTINCT branch_code) branch_count
            FROM cif_customer_identifiers
            GROUP BY customer_id
        )
        UPDATE cif_customers customer
        SET identifier_count=counts.identifier_count,
            branch_count=counts.branch_count
        FROM counts
        WHERE customer.id=counts.customer_id
    """))


def _refresh_registration_conflicts(db, batch_id: int) -> int:
    db.query(CifIdentityConflict).filter(
        CifIdentityConflict.conflict_type == "duplicate_registration_number",
        CifIdentityConflict.status == "pending",
    ).delete(synchronize_session=False)
    duplicate_values = (
        db.query(CifCustomerIdentifier.registration_number)
        .filter(
            CifCustomerIdentifier.registration_number.isnot(None),
            func.trim(CifCustomerIdentifier.registration_number) != "",
        )
        .group_by(CifCustomerIdentifier.registration_number)
        .having(func.count(distinct(CifCustomerIdentifier.customer_id)) > 1)
        .all()
    )
    mappings = []
    for (identity_value,) in duplicate_values:
        rows = (
            db.query(
                CifCustomerIdentifier.customer_id,
                CifCustomerIdentifier.full_cif_code,
                CifCustomerIdentifier.customer_name,
                CifCustomerIdentifier.branch_code,
            )
            .filter(CifCustomerIdentifier.registration_number == identity_value)
            .all()
        )
        mappings.append(
            {
                "import_batch_id": batch_id,
                "conflict_type": "duplicate_registration_number",
                "identity_value": identity_value,
                "customer_ids": sorted({int(row.customer_id) for row in rows}),
                "full_cif_codes": sorted({row.full_cif_code for row in rows}),
                "details": {
                    "records": [
                        {
                            "full_cif_code": row.full_cif_code,
                            "branch_code": row.branch_code,
                            "customer_name": row.customer_name,
                        }
                        for row in rows
                    ]
                },
                "status": "pending",
            }
        )
    if mappings:
        db.bulk_insert_mappings(CifIdentityConflict, mappings)
    return len(mappings)


def json_value(value):
    return value.isoformat() if hasattr(value, "isoformat") else value


def set_import_stage(db, batch, key: str, label: str, progress: int, details: dict | None = None) -> None:
    """Persist a durable, user-visible checkpoint for a background import."""
    now = datetime.now(timezone.utc).isoformat()
    history = list(batch.stage_history or [])
    if history and history[-1].get("key") == key:
        history[-1] = {**history[-1], "label": label, "progress": progress, "details": details or history[-1].get("details"), "at": now}
    else:
        if history and history[-1].get("status") == "processing":
            history[-1] = {**history[-1], "status": "completed", "finished_at": now}
        history.append({"key": key, "label": label, "status": "processing", "progress": progress, "details": details or {}, "at": now})
    batch.stage_history = history
    batch.stage = label
    batch.progress_percent = progress
    batch.heartbeat_at = datetime.now(timezone.utc)
    db.commit()


def resume_pending_cif_imports() -> None:
    """Resume durable CIF jobs after an application/container restart."""
    def worker():
        db = SessionLocal()
        try:
            pending = (
                db.query(CifImportBatch)
                .filter(CifImportBatch.status.in_(["queued", "processing"]))
                .order_by(CifImportBatch.uploaded_at, CifImportBatch.id)
                .all()
            )
            jobs = [(item.file_path, item.original_filename, item.uploaded_by or "system-recovery") for item in pending]
        finally:
            db.close()
        for path, filename, uploaded_by in jobs:
            if not Path(path).is_file():
                continue
            try:
                import_cif_file(path, filename, uploaded_by)
            except Exception:
                # import_cif_file persists the exact failure on the batch.
                continue

    threading.Thread(target=worker, name="cif-recovery-worker", daemon=True).start()


def import_cif_file(file_path: str | Path, original_filename: str | None = None, uploaded_by: str | None = None) -> int:
    path = Path(file_path)
    filename = original_filename or path.name
    checksum = file_sha256(path)
    db = SessionLocal()
    batch = None
    try:
        existing_batch = db.query(CifImportBatch).filter(CifImportBatch.content_sha256 == checksum).first()
        if existing_batch and existing_batch.status == "success":
            return existing_batch.id
        if existing_batch:
            batch = existing_batch
            batch.status = "processing"
            batch.error_message = None
            batch.uploaded_by = uploaded_by or batch.uploaded_by
            batch.importer_version = IMPORTER_VERSION
        else:
            format_warning = None
            if path.suffix.lower() == ".xls" and path.stat().st_size > 512 and (path.stat().st_size - 512) % 512:
                format_warning = "Kích thước OLE không khớp biên sector 512 byte; workbook vẫn được phép đọc nếu cấu trúc hợp lệ."
            batch = CifImportBatch(
                original_filename=filename,
                stored_filename=path.name,
                file_path=str(path),
                content_sha256=checksum,
                file_size=path.stat().st_size,
                actual_format=path.suffix.lower().lstrip("."),
                branch_code=infer_branch_code(filename),
                status="processing",
                stage=f"Đang mở file {path.suffix.upper().lstrip('.')}",
                started_at=datetime.now(timezone.utc),
                format_warning=format_warning,
                uploaded_by=uploaded_by,
                importer_version=IMPORTER_VERSION,
                attempt_count=1,
                heartbeat_at=datetime.now(timezone.utc),
            )
            db.add(batch)
        db.commit()
        db.refresh(batch)
        set_import_stage(db, batch, "checksum", "Đã kiểm tra checksum và định dạng file", 5, {"format": path.suffix.lower().lstrip("."), "file_size": path.stat().st_size})
        db.query(CifImportError).filter(CifImportError.import_batch_id == batch.id).delete(
            synchronize_session=False
        )
        db.query(CifSourceRecord).filter(CifSourceRecord.import_batch_id == batch.id).delete(
            synchronize_session=False
        )
        db.query(CifIdentityConflict).filter(
            CifIdentityConflict.import_batch_id == batch.id,
            CifIdentityConflict.conflict_type == "core_identity_mismatch",
            CifIdentityConflict.status == "pending",
        ).delete(synchronize_session=False)
        for field in (
            "processed_rows",
            "accepted_rows",
            "warning_rows",
            "rejected_rows",
            "new_customers",
            "new_identifiers",
            "updated_identifiers",
            "unchanged_identifiers",
            "duplicate_rows",
            "multi_branch_identifiers",
            "review_rows",
            "conflict_count",
        ):
            setattr(batch, field, 0)
        db.commit()

        set_import_stage(db, batch, "read_source", f"Đang đọc dữ liệu từ file {path.suffix.upper().lstrip('.')}", 10)

        parsed_rows: list[dict] = []
        source_records: list[dict] = []
        errors: list[dict] = []
        seen_full_codes: dict[str, int] = {}
        detected_branches: set[str] = set()
        column_nonblank: dict[str, int] = {}
        column_max_length: dict[str, int] = {}
        source_headers: list[str] = []
        total_rows = 0
        source_count = 0
        for sheet_name, headers, source_rows in iter_cif_sources(path):
            source_count += 1
            if not source_headers:
                source_headers = headers
                column_nonblank = {header: 0 for header in headers}
                column_max_length = {header: 0 for header in headers}
            missing = sorted(REQUIRED_COLUMNS - set(headers))
            if missing:
                raise ValueError(f"Sheet {sheet_name} thiếu cột bắt buộc: {', '.join(missing)}")
            duplicate_headers = sorted({header for header in headers if header and headers.count(header) > 1})
            if duplicate_headers:
                raise ValueError(f"Sheet {sheet_name} có tiêu đề cột trùng: {', '.join(duplicate_headers)}")
            batch.sheet_name = sheet_name if not batch.sheet_name else f"{batch.sheet_name}; {sheet_name}"
            for row_number, values in source_rows:
                total_rows += 1
                raw = {
                    header: clean(values[column] if column < len(values) else None)
                    for column, header in enumerate(headers)
                }
                if not any(raw.values()):
                    total_rows -= 1
                    continue
                for column, value in raw.items():
                    if value:
                        column_nonblank[column] = column_nonblank.get(column, 0) + 1
                        column_max_length[column] = max(column_max_length.get(column, 0), len(value))
                full_code = raw.get("custno", "")
                row_errors = []
                if not re.fullmatch(r"\d{13}", full_code):
                    row_errors.append(("error", "invalid_cif_code", "custno", full_code, "Mã CIF phải gồm đúng 13 chữ số"))
                duplicate_of = seen_full_codes.get(full_code) if full_code else None
                if duplicate_of is not None:
                    row_errors.append(("warning", "duplicate_cif_skipped", "custno", full_code, "Mã CIF đã xuất hiện trước đó trong file; dòng này được tự động bỏ qua"))
                elif full_code:
                    seen_full_codes[full_code] = row_number
                branch_code = full_code[:4] if len(full_code) >= 4 else ""
                customer_core_code = full_code[4:] if len(full_code) == 13 else ""
                if branch_code:
                    detected_branches.add(branch_code)
                for date_column in DATE_COLUMNS:
                    value = raw.get(date_column, "")
                    if value and not parse_yyyymmdd(value):
                        row_errors.append(
                            ("warning", "invalid_date", date_column, value, f"{date_column} không phải ngày YYYYMMDD hợp lệ trong khoảng 1900 đến hiện tại")
                        )
                for severity, code, field, raw_value, message in row_errors:
                    errors.append(
                        {
                            "import_batch_id": batch.id,
                            "source_row_number": row_number,
                            "full_cif_code": full_code or None,
                            "severity": severity,
                            "error_code": code,
                            "field_name": field,
                            "raw_value": raw_value,
                            "message": message,
                        }
                    )
                validation_status = "duplicate_skipped" if duplicate_of is not None else (
                    "rejected" if any(item[0] == "error" for item in row_errors) else (
                        "warning" if row_errors else "accepted"
                    )
                )
                source_records.append({
                    "import_batch_id": batch.id,
                    "source_sheet": sheet_name,
                    "source_row_number": row_number,
                    "full_cif_code": full_code or None,
                    "branch_code": branch_code or None,
                    "customer_core_code": customer_core_code or None,
                    "row_hash": row_hash(raw),
                    "validation_status": validation_status,
                    "duplicate_of_row_number": duplicate_of,
                    "raw_data": raw,
                })
                if len(source_records) >= 2000:
                    db.bulk_insert_mappings(CifSourceRecord, source_records)
                    db.commit()
                    source_records.clear()
                    batch.processed_rows = total_rows
                    set_import_stage(db, batch, "read_source", f"Đã đọc và tiền kiểm {total_rows:,} dòng", min(28, 10 + total_rows // 10000), {"rows_read": total_rows, "sheet": sheet_name})
                if duplicate_of is not None:
                    continue
                if any(item[0] == "error" for item in row_errors):
                    continue
                parsed_rows.append(
                    {
                        "full_cif_code": full_code,
                        "source_sheet": sheet_name,
                        "branch_code": branch_code,
                        "customer_core_code": customer_core_code,
                        "source_row_number": row_number,
                        "customer_name": raw.get("nmloc") or raw.get("nm") or None,
                        "customer_name_ascii": raw.get("nm") or None,
                        "short_name": raw.get("shrtnmloc") or raw.get("shrtnm") or None,
                        "customer_type": raw.get("custtpcd") or None,
                        "customer_detail_type": raw.get("custdtltpcd") or None,
                        "registration_number": normalized_identity(raw.get("regno", "")),
                        "passport_number": normalized_identity(raw.get("passno", "")),
                        "driver_license_number": normalized_identity(raw.get("dlno", "")),
                        "tax_number": normalized_identity(raw.get("taxno", "")),
                        "telephone": telephone(raw),
                        "address_type": raw.get("addrtpcd") or None,
                        "full_address": full_address(raw),
                        "province": raw.get("province") or None,
                        "district": raw.get("district") or None,
                        "commune_ward": raw.get("commune_ward") or None,
                        "nationality_code": raw.get("ctrycdnatl") or None,
                        "birth_date": date_yyyymmdd(raw.get("name_1", "")),
                        "gender_code": raw.get("name_3") or None,
                        "establishment_date": date_yyyymmdd(raw.get("incrdt", "")),
                        "occupation": raw.get("profnm") or None,
                        "source_status": raw.get("stscd") or None,
                        "normalized_status": normalized_status(raw.get("stscd", "")),
                        "operator_user": raw.get("usridop1") or None,
                        "raw_data": raw,
                        "has_warning": any(item[0] == "warning" for item in row_errors),
                    }
                )
        if source_count > 1:
            batch.format_warning = (batch.format_warning or "") + " File có nhiều sheet; importer đã xử lý toàn bộ sheet cùng cấu trúc."
        if source_records:
            db.bulk_insert_mappings(CifSourceRecord, source_records)
            db.commit()
            source_records.clear()

        filename_branch = infer_branch_code(filename)
        if not detected_branches:
            raise ValueError("Không xác định được mã chi nhánh từ cột CUSTNO")
        sorted_branches = sorted(detected_branches)
        is_multi_branch = len(sorted_branches) > 1
        detected_branch = sorted_branches[0] if not is_multi_branch else None
        if filename_branch and not is_multi_branch and filename_branch != detected_branch:
            raise ValueError(f"Chi nhánh trong tên file {filename_branch} không khớp dữ liệu {detected_branch}")
        batch.branch_code = "MULTI" if is_multi_branch else detected_branch
        if is_multi_branch:
            multi_note = f"File gộp {len(sorted_branches)} chi nhánh: {', '.join(sorted_branches)}."
            batch.format_warning = " ".join(filter(None, [batch.format_warning, multi_note]))
        batch.total_rows = total_rows
        batch.processed_rows = total_rows
        batch.column_stats = {
            "total_columns": len(source_headers),
            "branch_codes": sorted_branches,
            "branch_count": len(sorted_branches),
            "blank_columns": [column for column in source_headers if column_nonblank.get(column, 0) == 0],
            "columns": [
                {
                    "name": column,
                    "nonblank": column_nonblank.get(column, 0),
                    "blank": total_rows - column_nonblank.get(column, 0),
                    "coverage_percent": round(column_nonblank.get(column, 0) * 100 / total_rows, 2) if total_rows else 0,
                    "max_length": column_max_length.get(column, 0),
                }
                for column in source_headers
            ],
        }
        branch_row_counts = {
            branch: sum(1 for row in parsed_rows if row["branch_code"] == branch)
            for branch in sorted_branches
        }
        set_import_stage(db, batch, "validate_structure", "Đã kiểm tra cấu trúc, chi nhánh và chất lượng cột", 32, {"total_rows": total_rows, "total_columns": len(source_headers), "branch_code": batch.branch_code, "branch_codes": sorted_branches, "branch_row_counts": branch_row_counts})
        set_import_stage(db, batch, "match_customer", "Đang chuẩn hóa mã lõi và đối chiếu khách hàng hiện có", 38, {"unique_core_codes": len({row['customer_core_code'] for row in parsed_rows})})

        core_codes = sorted({row["customer_core_code"] for row in parsed_rows})
        customer_by_core = {}
        for group in chunked(core_codes):
            for customer in db.query(CifCustomer).filter(CifCustomer.customer_core_code.in_(group)).all():
                customer_by_core[customer.customer_core_code] = customer
        first_by_core = {}
        for row in parsed_rows:
            first_by_core.setdefault(row["customer_core_code"], row)
        new_customer_mappings = []
        for core_code in core_codes:
            if core_code in customer_by_core:
                continue
            row = first_by_core[core_code]
            new_customer_mappings.append(
                {
                    "customer_core_code": core_code,
                    "customer_name": row["customer_name"],
                    "customer_name_ascii": row["customer_name_ascii"],
                    "customer_type": row["customer_type"],
                    "customer_detail_type": row["customer_detail_type"],
                    "registration_number": row["registration_number"],
                    "passport_number": row["passport_number"],
                    "tax_number": row["tax_number"],
                    "telephone": row["telephone"],
                    "full_address": row["full_address"],
                    "nationality_code": row["nationality_code"],
                    "birth_date": row["birth_date"],
                    "gender_code": row["gender_code"],
                    "establishment_date": row["establishment_date"],
                    "occupation": row["occupation"],
                    "status": row["normalized_status"],
                    "first_import_batch_id": batch.id,
                    "last_import_batch_id": batch.id,
                }
            )
        set_import_stage(db, batch, "write_customers", f"Chuẩn bị ghi {len(new_customer_mappings):,} khách hàng mới", 45, {"written": 0, "total": len(new_customer_mappings)})
        for start in range(0, len(new_customer_mappings), 5000):
            db.bulk_insert_mappings(CifCustomer, new_customer_mappings[start:start + 5000])
            db.commit()
            set_import_stage(db, batch, "write_customers", f"Đang ghi khách hàng mới: {min(start + 5000, len(new_customer_mappings)):,}/{len(new_customer_mappings):,}", 45 + int(10 * min(start + 5000, len(new_customer_mappings)) / max(len(new_customer_mappings), 1)), {"written": min(start + 5000, len(new_customer_mappings)), "total": len(new_customer_mappings)})
        for group in chunked(core_codes):
            for customer in db.query(CifCustomer).filter(CifCustomer.customer_core_code.in_(group)).all():
                customer_by_core[customer.customer_core_code] = customer

        full_codes = [row["full_cif_code"] for row in parsed_rows]
        existing_branch_codes = {
            code for (code,) in (
                db.query(CifCustomerIdentifier.full_cif_code)
                .filter(CifCustomerIdentifier.branch_code.in_(sorted_branches))
                .all()
            )
        }
        missing_from_latest = len(existing_branch_codes - set(full_codes))
        source_record_ids = {
            (sheet, row_number): source_id
            for source_id, sheet, row_number in (
                db.query(
                    CifSourceRecord.id,
                    CifSourceRecord.source_sheet,
                    CifSourceRecord.source_row_number,
                )
                .filter(CifSourceRecord.import_batch_id == batch.id)
                .all()
            )
        }
        identifier_by_code = {}
        for group in chunked(full_codes):
            for identifier in db.query(CifCustomerIdentifier).filter(CifCustomerIdentifier.full_cif_code.in_(group)).all():
                identifier_by_code[identifier.full_cif_code] = identifier
        branches_by_customer: dict[int, set[str]] = {}
        customer_ids = [customer.id for customer in customer_by_core.values()]
        for group in chunked(customer_ids):
            branch_rows = (
                db.query(
                    CifCustomerIdentifier.customer_id,
                    CifCustomerIdentifier.branch_code,
                )
                .filter(CifCustomerIdentifier.customer_id.in_(group))
                .all()
            )
            for customer_id, branch_code in branch_rows:
                branches_by_customer.setdefault(customer_id, set()).add(branch_code)

        insert_mappings = []
        updated_identifiers = 0
        unchanged_identifiers = 0
        pending_changes = 0
        new_source_records = 0
        no_change_records = 0
        source_review_updates: list[dict] = []
        multi_branch_identifiers = 0
        review_source_rows: set[int] = set()
        profile_conflicts = []
        existing_profile_conflicts = {
            value
            for (value,) in (
                db.query(CifIdentityConflict.identity_value)
                .filter(
                    CifIdentityConflict.conflict_type == "core_identity_mismatch",
                    CifIdentityConflict.status == "pending",
                )
                .all()
            )
        }
        for row in parsed_rows:
            customer = customer_by_core[row["customer_core_code"]]
            identifier = identifier_by_code.get(row["full_cif_code"])
            payload = {key: value for key, value in row.items() if key != "has_warning"}
            payload.pop("source_sheet", None)
            payload.update({"customer_id": customer.id, "import_batch_id": batch.id})
            source_record_id = source_record_ids.get((row["source_sheet"], row["source_row_number"]))
            if identifier:
                unchanged_identifiers += 1
                differences = {
                    field: {
                        "current": json_value(getattr(identifier, field)),
                        "incoming": json_value(row.get(field)),
                    }
                    for field in REVIEW_FIELDS
                    if json_value(getattr(identifier, field)) != json_value(row.get(field))
                }
                if source_record_id:
                    source_review_updates.append({
                        "id": source_record_id,
                        "current_snapshot": {field: json_value(getattr(identifier, field)) for field in REVIEW_FIELDS},
                        "changed_fields": differences or None,
                        "comparison_status": "changed" if differences else "no_change",
                        "review_status": "pending" if differences else "not_required",
                    })
                if differences:
                    pending_changes += 1
                else:
                    no_change_records += 1
                continue
            else:
                insert_mappings.append(payload)
                if source_record_id:
                    source_review_updates.append({
                        "id": source_record_id,
                        "comparison_status": "new",
                        "review_status": "applied",
                    })
                new_source_records += 1
                known_branches = branches_by_customer.setdefault(customer.id, set())
                if known_branches and row["branch_code"] not in known_branches:
                    multi_branch_identifiers += 1
                known_branches.add(row["branch_code"])

            differences = {
                field: {
                    "master": getattr(customer, field),
                    "source": row[field],
                }
                for field in MASTER_IDENTITY_FIELDS
                if getattr(customer, field) and row[field] and getattr(customer, field) != row[field]
            }
            if differences:
                review_source_rows.add(row["source_row_number"])
                if row["customer_core_code"] not in existing_profile_conflicts:
                    profile_conflicts.append(
                        {
                            "import_batch_id": batch.id,
                            "conflict_type": "core_identity_mismatch",
                            "identity_value": row["customer_core_code"],
                            "customer_ids": [customer.id],
                            "full_cif_codes": [row["full_cif_code"]],
                            "details": {
                                "source_row_number": row["source_row_number"],
                                "branch_code": row["branch_code"],
                                "differences": differences,
                            },
                            "status": "pending",
                        }
                    )
                    existing_profile_conflicts.add(row["customer_core_code"])
            customer.last_import_batch_id = batch.id
            if not customer.customer_name and row["customer_name"]:
                customer.customer_name = row["customer_name"]
        set_import_stage(db, batch, "write_identifiers", f"Chuẩn bị ghi {len(insert_mappings):,} mã CIF", 62, {"written": 0, "total": len(insert_mappings)})
        for start in range(0, len(insert_mappings), 5000):
            db.bulk_insert_mappings(CifCustomerIdentifier, insert_mappings[start:start + 5000])
            db.commit()
            set_import_stage(db, batch, "write_identifiers", f"Đang ghi mã CIF: {min(start + 5000, len(insert_mappings)):,}/{len(insert_mappings):,}", 62 + int(13 * min(start + 5000, len(insert_mappings)) / max(len(insert_mappings), 1)), {"written": min(start + 5000, len(insert_mappings)), "total": len(insert_mappings)})
        for start in range(0, len(source_review_updates), 5000):
            db.bulk_update_mappings(CifSourceRecord, source_review_updates[start:start + 5000])
            db.commit()
        if errors:
            for start in range(0, len(errors), 5000):
                db.bulk_insert_mappings(CifImportError, errors[start:start + 5000])
                db.commit()

        set_import_stage(db, batch, "identity_conflicts", "Đang kiểm tra trùng và xung đột định danh", 85)
        _refresh_customer_counts(db)
        registration_conflict_count = _refresh_registration_conflicts(db, batch.id)
        if profile_conflicts:
            db.bulk_insert_mappings(CifIdentityConflict, profile_conflicts)
        warning_rows = len({item["source_row_number"] for item in errors if item["severity"] == "warning"})
        rejected_rows = total_rows - len(parsed_rows)
        duplicate_rows = len(
            {
                item["source_row_number"]
                for item in errors
                if item["error_code"] == "duplicate_cif_skipped"
            }
        )
        batch.processed_rows = total_rows
        batch.accepted_rows = len(parsed_rows)
        batch.warning_rows = warning_rows
        batch.rejected_rows = rejected_rows
        batch.new_customers = len(new_customer_mappings)
        batch.new_identifiers = len(insert_mappings)
        batch.updated_identifiers = updated_identifiers
        batch.unchanged_identifiers = unchanged_identifiers
        batch.duplicate_rows = duplicate_rows
        batch.multi_branch_identifiers = multi_branch_identifiers
        batch.review_rows = len(review_source_rows)
        batch.conflict_count = registration_conflict_count + len(profile_conflicts)
        batch.comparison_summary = {
            "new": new_source_records,
            "no_change": no_change_records,
            "pending_changes": pending_changes,
            "missing_from_latest": missing_from_latest,
        }
        set_import_stage(db, batch, "quality_summary", "Đã tổng hợp kết quả và chất lượng dữ liệu", 95, {"accepted": len(parsed_rows), "rejected": rejected_rows, "conflicts": registration_conflict_count + len(profile_conflicts)})
        batch.status = "success"
        history = list(batch.stage_history or [])
        if history and history[-1].get("status") == "processing":
            history[-1] = {**history[-1], "status": "completed", "finished_at": datetime.now(timezone.utc).isoformat()}
        history.append({"key": "completed", "label": "Hoàn thành", "status": "completed", "progress": 100, "at": datetime.now(timezone.utc).isoformat(), "details": {"processed_rows": total_rows}})
        batch.stage_history = history
        batch.stage = "Hoàn thành"
        batch.progress_percent = 100
        batch.finished_at = datetime.now(timezone.utc)
        db.commit()
        return batch.id
    except Exception as exc:
        db.rollback()
        if batch and batch.id:
            failed = db.query(CifImportBatch).filter(CifImportBatch.id == batch.id).first()
            if failed:
                failed.status = "error"
                failed.stage = "Import thất bại"
                history = list(failed.stage_history or [])
                if history and history[-1].get("status") == "processing":
                    history[-1] = {**history[-1], "status": "error", "finished_at": datetime.now(timezone.utc).isoformat(), "error": str(exc)}
                history.append({"key": "failed", "label": "Import thất bại", "status": "error", "progress": failed.progress_percent, "at": datetime.now(timezone.utc).isoformat(), "error": str(exc)})
                failed.stage_history = history
                failed.error_message = str(exc)
                failed.finished_at = datetime.now(timezone.utc)
                db.commit()
        raise
    finally:
        db.close()
