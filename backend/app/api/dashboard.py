from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from time import monotonic

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import Integer, and_, case, desc, exists, func, literal, or_, select, text
from sqlalchemy.orm import Query as OrmQuery, Session
from sqlalchemy.orm import aliased

from app.auth.branch_scope import BranchScope
from app.auth.dependencies import get_branch_scope, get_current_user, require_any_permission
from app.auth.schemas import CurrentUser
from app.database import get_db
from app.security_audit import record_security_event
from app.export_progress import begin_export, update_export
from app.fee_rules import FEE_CANDIDATE_PREFIXES, FEE_CATEGORY_LABELS, FEE_CATEGORY_PREFIXES, FEE_FIELDS
from app.analysis_cache import get_shared_analysis_cache, set_shared_analysis_cache
from app.models import (
    CN05CustomerService,
    DP01DepositAccount,
    CustomerPeriodBranchDetail,
    CustomerPeriodProfile,
    CustomerProcessingJob,
    ImportFile,
    KH02CustomerTransaction,
    LN01Loan,
    OrgBranch,
    OrgDepartment,
    PF14AccountBalance,
    ReportSourceStatus,
    RR01HandledRiskLoan,
    SystemUser,
)


router = APIRouter(prefix="/api/dashboard", tags=["dashboard"])
_INSIGHTS_CACHE: dict[tuple[str, str | None, str | None], tuple[float, dict]] = {}
_SUMMARY_CACHE: dict[tuple[str, str | None, str | None], tuple[float, dict]] = {}
_BUSINESS_CACHE: dict[tuple[str, str | None, str | None], tuple[float, dict]] = {}
_BUSINESS_TREND_CACHE: dict[tuple[int, str | None, str | None], tuple[float, dict]] = {}
_INSIGHTS_CACHE_TTL_SECONDS = 60

ACTIVE_SERVICE_KEYS = [
    "thau_chi",
    "tk_so_dep",
    "agribank_plus",
    "tin_nhan_ott",
    "e_banking",
    "sms_nhac_no_vay",
    "sms_tien_gui",
    "the_ghi_no_noi_dia",
    "the_ghi_no_quoc_te",
    "the_td_noi_dia",
    "the_td_quoc_te",
    "the_td_loc_viet",
    "bao_lanh",
    "loa_bien_dong_so_du",
    "phat_hanh_lc",
    "ttqt",
    "thuho_dien",
    "thuho_nuoc",
    "thuho_dt",
    "hkd_tk",
    "abic_batk",
    "abic_bathe",
    "pos",
]


def _account_prefix_condition(column, prefixes):
    normalized = func.trim(func.coalesce(column, ""))
    return or_(*(normalized.like(f"{prefix}%") for prefix in prefixes))


def _fee_category_condition(column, category: str):
    prefixes = FEE_CATEGORY_PREFIXES.get(category)
    return _account_prefix_condition(column, prefixes) if prefixes else None


def _all_classified_fee_condition(column):
    return or_(*(
        _fee_category_condition(column, category)
        for category in FEE_CATEGORY_PREFIXES
    ))


def _fee_category_expression(column):
    return case(*(
        (_fee_category_condition(column, category), category)
        for category in FEE_CATEGORY_PREFIXES
    ), else_="unclassified")


def _fee_reconciliation(db: Session, period_key: str, scope: BranchScope, customer_ids, profile_total: float) -> dict:
    """Reconcile KH02 fee candidates once, without counting an account code twice."""
    candidate = _account_prefix_condition(KH02CustomerTransaction.account_code, FEE_CANDIDATE_PREFIXES)
    classified = _all_classified_fee_condition(KH02CustomerTransaction.account_code)
    eligible_customers = customer_ids
    if eligible_customers is None:
        eligible_customers = db.query(CustomerPeriodProfile.ma_kh).filter(
            CustomerPeriodProfile.period_key == period_key
        ).distinct().subquery()
    query = db.query(KH02CustomerTransaction).filter(
        KH02CustomerTransaction.period_key == period_key,
        candidate,
        func.trim(KH02CustomerTransaction.customer_code).in_(select(eligible_customers.c.ma_kh)),
    )
    if scope.ma_cn:
        query = query.filter(func.trim(KH02CustomerTransaction.branch_code) == scope.ma_cn)
    credit = func.coalesce(KH02CustomerTransaction.credit_amount, 0)
    debit = func.coalesce(KH02CustomerTransaction.debit_amount, 0)
    net = credit - debit
    row = query.with_entities(
        func.count(KH02CustomerTransaction.id),
        func.count(func.distinct(func.trim(KH02CustomerTransaction.customer_code))),
        func.coalesce(func.sum(credit), 0),
        func.coalesce(func.sum(debit), 0),
        func.coalesce(func.sum(net), 0),
        func.coalesce(func.sum(case((classified, credit), else_=0)), 0),
        func.coalesce(func.sum(case((classified, debit), else_=0)), 0),
        func.coalesce(func.sum(case((classified, net), else_=0)), 0),
        func.sum(case((classified, 1), else_=0)),
        func.count(func.distinct(case((classified, func.trim(KH02CustomerTransaction.customer_code)), else_=None))),
        func.coalesce(func.sum(case((~classified, credit), else_=0)), 0),
        func.coalesce(func.sum(case((~classified, debit), else_=0)), 0),
        func.coalesce(func.sum(case((~classified, net), else_=0)), 0),
        func.sum(case((~classified, 1), else_=0)),
        func.count(func.distinct(case((~classified, func.trim(KH02CustomerTransaction.customer_code)), else_=None))),
    ).one()
    source_net = float(row[4] or 0)
    classified_net = float(row[7] or 0)
    unclassified_net = float(row[12] or 0)
    total_records = int(row[0] or 0)
    classified_records = int(row[8] or 0)
    balance_difference = source_net - classified_net - unclassified_net
    return {
        "candidate_prefixes": list(FEE_CANDIDATE_PREFIXES),
        "source": {"records": total_records, "customers": int(row[1] or 0), "credit": float(row[2] or 0), "debit": float(row[3] or 0), "net": source_net},
        "classified": {"records": classified_records, "customers": int(row[9] or 0), "credit": float(row[5] or 0), "debit": float(row[6] or 0), "net": classified_net},
        "unclassified": {"records": int(row[13] or 0), "customers": int(row[14] or 0), "credit": float(row[10] or 0), "debit": float(row[11] or 0), "net": unclassified_net},
        "record_coverage_pct": round(classified_records * 100 / total_records, 2) if total_records else 0,
        "balance_difference": balance_difference,
        "is_balanced": abs(balance_difference) <= 0.01,
        "profile_total": float(profile_total or 0),
        "profile_difference": float(profile_total or 0) - classified_net,
        "profile_is_balanced": abs(float(profile_total or 0) - classified_net) <= 0.01,
    }


def _advanced_filters(
    keyword: str | None = Query(default=None), customer_type: str | None = Query(default=None),
    loan_type: str | None = Query(default=None), officer_code: str | None = Query(default=None),
    multi_branch: bool | None = Query(default=None), no_service: bool = Query(default=False),
    has_deposit: bool | None = Query(default=None), has_loan: bool | None = Query(default=None),
    min_deposit: float | None = Query(default=None, ge=0), max_deposit: float | None = Query(default=None, ge=0),
    min_loan: float | None = Query(default=None, ge=0), max_loan: float | None = Query(default=None, ge=0),
    min_casa: float | None = Query(default=None, ge=0), max_casa: float | None = Query(default=None, ge=0),
    service_codes: str | None = Query(default=None), min_service_count: int | None = Query(default=None, ge=0, le=30),
    missing_phone: bool | None = Query(default=None),
    missing_officer: bool | None = Query(default=None),
    unclear_primary_branch: bool | None = Query(default=None),
    new_in_period: bool | None = Query(default=None),
):
    values = {key: value for key, value in locals().items() if value is not None and value != ""}
    if not no_service:
        values.pop("no_service", None)
    return values


def _matching_customer_ids(db: Session, period_key: str, scope: BranchScope, filters: dict):
    """Một tập mã KH duy nhất dùng chung cho KPI, bảng tổng hợp, cảnh báo và drill-down."""
    from app.api.customer_processing import apply_profile_filters

    query = db.query(CustomerPeriodProfile.ma_kh).filter(CustomerPeriodProfile.period_key == period_key)
    query = apply_profile_filters(
        query,
        period_key=period_key,
        branch_code=scope.ma_cn,
        pgd_code=scope.ma_pgd,
        **filters,
    )
    return query.distinct().subquery()


def _scope_filters(scope: BranchScope, filters: dict) -> dict:
    """Prevent a client request from widening an OWN user's assigned-customer scope."""
    scoped = dict(filters or {})
    if scope.officer_code:
        scoped["officer_code"] = scope.officer_code
    return scoped


def _restrict_to_matching_customers(query, model, customer_ids):
    return query if customer_ids is None else query.filter(model.ma_kh.in_(select(customer_ids.c.ma_kh)))


def _customer_filter_ids(db: Session, period_key: str, scope: BranchScope, filters: dict):
    """Avoid a redundant self-IN subquery for the common province/no-filter view."""
    if not scope.ma_cn and not scope.ma_pgd and not filters:
        return None
    return _matching_customer_ids(db, period_key, scope, filters)


def _analysis_data_version(db: Session, period_key: str | None = None) -> int:
    query = db.query(func.max(CustomerProcessingJob.id)).filter(CustomerProcessingJob.status == "success")
    if period_key:
        query = query.filter(CustomerProcessingJob.period_key <= period_key)
    return int(query.scalar() or 0)

# Đồng bộ với quy tắc cơ hội bán chéo / loại KH thực tế từ DP01.
RETAIL_CUSTOMER_TYPES = ("Cá nhân", "KHCN", "Tư nhân")

SERVICE_LABELS = {
    "thau_chi": "Thấu chi",
    "tk_so_dep": "TK số đẹp",
    "agribank_plus": "Agribank Plus",
    "tin_nhan_ott": "Tin nhắn OTT",
    "e_banking": "E-Banking",
    "sms_nhac_no_vay": "SMS nhắc nợ vay",
    "sms_tien_gui": "SMS tiền gửi",
    "the_ghi_no_noi_dia": "Thẻ ghi nợ nội địa",
    "the_ghi_no_quoc_te": "Thẻ ghi nợ quốc tế",
    "the_td_noi_dia": "Thẻ TD nội địa",
    "the_td_quoc_te": "Thẻ TD quốc tế",
    "the_td_loc_viet": "Thẻ TD Lộc Việt",
    "bao_lanh": "Bảo lãnh",
    "loa_bien_dong_so_du": "Loa Thần Tài",
    "phat_hanh_lc": "Phát hành LC",
    "ttqt": "LC/TTQT/KDNT",
    "thuho_dien": "Thu hộ tiền điện",
    "thuho_nuoc": "Thu hộ tiền nước",
    "thuho_dt": "Thu hộ điện thoại/viễn thông",
    "hkd_tk": "Tài khoản hộ kinh doanh",
    "abic_batk": "Bảo an tài khoản",
    "abic_bathe": "Bảo an chủ thẻ",
    "pos": "Đơn vị chấp nhận thẻ POS",
}

SERVICE_GROUPS = {
    "thau_chi": "Tài khoản",
    "tk_so_dep": "Tài khoản",
    "agribank_plus": "Digital",
    "tin_nhan_ott": "Digital",
    "e_banking": "Digital",
    "sms_nhac_no_vay": "Digital",
    "sms_tien_gui": "Digital",
    "the_ghi_no_noi_dia": "Thẻ",
    "the_ghi_no_quoc_te": "Thẻ",
    "the_td_noi_dia": "Thẻ",
    "the_td_quoc_te": "Thẻ",
    "the_td_loc_viet": "Thẻ",
    "bao_lanh": "Bảo lãnh/TTQT",
    "loa_bien_dong_so_du": "Digital",
    "phat_hanh_lc": "Bảo lãnh/TTQT",
    "ttqt": "Bảo lãnh/TTQT",
    "thuho_dien": "Thanh toán",
    "thuho_nuoc": "Thanh toán",
    "thuho_dt": "Thanh toán",
    "hkd_tk": "Tài khoản",
    "abic_batk": "Bảo hiểm",
    "abic_bathe": "Bảo hiểm",
    "pos": "Thẻ",
}

CAMPAIGN_GROUP_PRIORITY = {
    "KHCN": ["Digital", "Thẻ", "Thanh toán", "Tài khoản", "Bảo hiểm", "Bảo lãnh/TTQT", "Khác"],
    "KHDN": ["Bảo lãnh/TTQT", "Digital", "Thanh toán", "Tài khoản", "Thẻ", "Bảo hiểm", "Khác"],
}

LOAN_TYPE_LABELS = ("Thấu chi", "Ngắn", "Trung", "Dài", "Khác")

# Mẫu số thâm nhập đặc thù (mặc định = toàn bộ KH trong phạm vi).
SERVICE_PENETRATION_BASE = {
    "sms_nhac_no_vay": "loan",           # / số KH có dư nợ vay
    "tin_nhan_ott": "agribank_plus",     # / số KH đăng ký Agribank Plus
}


def _loan_type_category_expr():
    return case(
        (LN01Loan.loan_type == "Thấu chi trên TK khách hàng", "Thấu chi"),
        (LN01Loan.loan_type == "Vay ngắn hạn (TK 211)", "Ngắn"),
        (LN01Loan.loan_type == "Vay trung hạn (TK 212)", "Trung"),
        (LN01Loan.loan_type == "Vay dài hạn (TK 213)", "Dài"),
        else_="Khác",
    )


def _build_service_penetration_rows(*, count_fn, loan_base: int, plus_base: int, total_customers: int) -> list[dict]:
    """count_fn(key, base_kind) -> số KH đã dùng dịch vụ trong đúng tập mẫu số."""
    rows = []
    for key in ACTIVE_SERVICE_KEYS:
        base_kind = SERVICE_PENETRATION_BASE.get(key)
        if base_kind == "loan":
            base = loan_base
            base_label = "KH vay"
            count = count_fn(key, "loan")
        elif base_kind == "agribank_plus":
            base = plus_base
            base_label = "KH Agribank Plus"
            count = count_fn(key, "agribank_plus")
        else:
            base = total_customers
            base_label = "KH"
            count = count_fn(key, "all")
        rows.append({
            "key": key,
            "label": SERVICE_LABELS[key],
            "group": SERVICE_GROUPS.get(key, "Khác"),
            "count": count,
            "base": base,
            "base_label": base_label,
            "pct": round((count / base) * 100, 2) if base else 0,
        })
    rows.sort(key=lambda item: item["pct"], reverse=True)
    return rows


def _count_kh_with_payment_account(
    db: Session,
    period_key: str,
    ma_cn: str | None,
    ma_pgd: str | None,
    customer_ids=None,
) -> int:
    """KH có TK thanh toán theo CN05.TKTT_SO_TK > 0, trong tập KH của phạm vi dashboard."""
    q = db.query(func.count(func.distinct(CN05CustomerService.ma_kh))).filter(
        CN05CustomerService.period_key == period_key,
        CN05CustomerService.ma_kh.isnot(None),
        func.coalesce(CN05CustomerService.tktt_so_tk, 0) > 0,
    )
    if customer_ids is not None:
        q = q.filter(CN05CustomerService.ma_kh.in_(db.query(customer_ids.c.ma_kh)))
    if ma_cn:
        scoped_kh = (
            _branch_detail_query(db, period_key, ma_cn, ma_pgd)
            .with_entities(CustomerPeriodBranchDetail.ma_kh)
            .distinct()
            .subquery()
        )
        q = q.filter(
            CN05CustomerService.ma_cn == ma_cn.strip(),
            CN05CustomerService.ma_kh.in_(db.query(scoped_kh.c.ma_kh)),
        )
    else:
        scoped_kh = (
            _base_query(db, period_key, None, ma_pgd)
            .with_entities(CustomerPeriodProfile.ma_kh)
            .subquery()
        )
        q = q.filter(CN05CustomerService.ma_kh.in_(db.query(scoped_kh.c.ma_kh)))
    return int(q.scalar() or 0)


def _with_payment_account_penetration(
    rows: list[dict],
    *,
    db: Session,
    period_key: str,
    ma_cn: str | None,
    ma_pgd: str | None,
    total_customers: int,
    customer_ids=None,
) -> list[dict]:
    """Thêm chỉ tiêu độ phủ TK thanh toán / tổng KH (không ghi vào ACTIVE_SERVICE_KEYS)."""
    count = _count_kh_with_payment_account(db, period_key, ma_cn, ma_pgd, customer_ids)
    rows = [
        *rows,
        {
            "key": "tk_thanh_toan",
            "label": "Tài khoản thanh toán",
            "group": "Tài khoản",
            "count": count,
            "base": total_customers,
            "base_label": "KH",
            "pct": round((count / total_customers) * 100, 2) if total_customers else 0,
        },
    ]
    rows.sort(key=lambda item: item["pct"], reverse=True)
    return rows


def _ln01_query(db: Session, period_key: str, ma_cn: str | None, ma_pgd: str | None) -> OrmQuery:
    query = db.query(LN01Loan).filter(LN01Loan.period_key == period_key)
    if ma_cn:
        query = query.filter(LN01Loan.brcd == ma_cn.strip())
    if ma_pgd:
        profile_subq = (
            db.query(CustomerPeriodProfile.ma_kh)
            .filter(CustomerPeriodProfile.period_key == period_key)
            .filter(CustomerPeriodProfile.pgd_codes.ilike(f"%{ma_pgd}%"))
        )
        if ma_cn:
            profile_subq = profile_subq.filter(CustomerPeriodProfile.branch_codes.ilike(f"%{ma_cn}%"))
        query = query.filter(LN01Loan.custseq.in_(profile_subq))
    return query


def _build_loan_type_breakdown(
    db: Session, period_key: str, ma_cn: str | None, ma_pgd: str | None, customer_ids=None
) -> list[dict]:
    category_expr = _loan_type_category_expr()
    source_query = _ln01_query(db, period_key, ma_cn, ma_pgd)
    if customer_ids is not None:
        source_query = source_query.filter(LN01Loan.custseq.in_(select(customer_ids.c.ma_kh)))
    loan_type_rows = (
        source_query
        .with_entities(
            category_expr.label("loan_category"),
            func.coalesce(func.sum(LN01Loan.du_no), 0),
        )
        .group_by(category_expr)
        .all()
    )
    total_ln_amt = sum(float(amt or 0) for _, amt in loan_type_rows)
    order_map = {label: index for index, label in enumerate(LOAN_TYPE_LABELS)}
    loan_type_breakdown = [
        {
            "type": loan_category,
            "amt": float(amt or 0),
            "pct": round((float(amt or 0) / total_ln_amt) * 100) if total_ln_amt else 0,
        }
        for loan_category, amt in loan_type_rows
        if float(amt or 0) > 0
    ]
    loan_type_breakdown.sort(key=lambda item: (order_map.get(item["type"], 99), -item["amt"]))
    return loan_type_breakdown


def serialize_value(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _base_query(db: Session, period_key: str, ma_cn: str | None, ma_pgd: str | None):
    query = db.query(CustomerPeriodProfile).filter(CustomerPeriodProfile.period_key == period_key)
    if ma_cn:
        query = query.filter(CustomerPeriodProfile.branch_codes.ilike(f"%{ma_cn}%"))
    if ma_pgd:
        query = query.filter(CustomerPeriodProfile.pgd_codes.ilike(f"%{ma_pgd}%"))
    return query


def _branch_detail_query(db: Session, period_key: str, ma_cn: str | None, ma_pgd: str | None):
    query = db.query(CustomerPeriodBranchDetail).filter(CustomerPeriodBranchDetail.period_key == period_key)
    if ma_cn:
        query = query.filter(CustomerPeriodBranchDetail.branch_code == ma_cn.strip())
    if ma_pgd:
        query = (
            query.join(
                SystemUser,
                and_(
                    SystemUser.is_active.is_(True),
                    or_(
                        SystemUser.credit_officer_code == CustomerPeriodBranchDetail.ma_cb,
                        SystemUser.employee_code == CustomerPeriodBranchDetail.ma_cb,
                        SystemUser.employee_code == CustomerPeriodBranchDetail.officer_employee_code,
                    ),
                ),
            )
            .join(OrgDepartment, OrgDepartment.id == SystemUser.department_id)
            .join(OrgBranch, OrgBranch.id == SystemUser.branch_id)
            .filter(
                OrgDepartment.department_code == ma_pgd.strip(),
                OrgDepartment.status == "active",
                OrgBranch.status == "active",
            )
        )
        if ma_cn:
            query = query.filter(OrgBranch.branch_code == ma_cn.strip())
    return query


def _customer_branch_agg_subquery(db: Session, period_key: str, ma_cn: str, ma_pgd: str | None):
    """Gom theo mã KH trong một chi nhánh (cộng nhiều PGD cùng CN)."""
    detail_query = _branch_detail_query(db, period_key, ma_cn, ma_pgd)
    service_max = [
        func.max(getattr(CustomerPeriodBranchDetail, key)).label(key)
        for key in ACTIVE_SERVICE_KEYS
    ]
    return (
        detail_query.with_entities(
            CustomerPeriodBranchDetail.ma_kh.label("ma_kh"),
            func.max(CustomerPeriodBranchDetail.ten_kh).label("ten_kh"),
            func.max(CustomerPeriodBranchDetail.loai_khach_hang).label("loai_khach_hang"),
            func.max(CustomerPeriodBranchDetail.ma_pgd).label("ma_pgd"),
            func.max(CustomerPeriodBranchDetail.ten_pgd).label("ten_pgd"),
            func.max(CustomerPeriodBranchDetail.loai_vay).label("loai_vay"),
            func.max(CustomerPeriodBranchDetail.ma_cb).label("ma_cb"),
            func.max(CustomerPeriodBranchDetail.ten_can_bo).label("ten_can_bo"),
            func.coalesce(func.sum(CustomerPeriodBranchDetail.so_du_tien_vay), 0).label("so_du_tien_vay"),
            func.coalesce(func.sum(CustomerPeriodBranchDetail.so_du_tien_gui), 0).label("so_du_tien_gui"),
            func.coalesce(func.sum(CustomerPeriodBranchDetail.so_du_tgtt_binh_quan), 0).label("so_du_tgtt_binh_quan"),
            func.coalesce(func.sum(CustomerPeriodBranchDetail.doanh_so_cramt), 0).label("doanh_so_cramt"),
            *service_max,
        )
        .group_by(CustomerPeriodBranchDetail.ma_kh)
        .subquery()
    )


def _build_officer_leaderboard(
    db: Session,
    period_key: str,
    ma_cn: str | None,
    ma_pgd: str | None,
    customer_ids=None,
    previous_period: str | None = None,
) -> list[dict]:
    """Top cán bộ theo dư nợ. Có CN → branch_details; chọn tất cả → profiles toàn tỉnh (không trùng KH)."""
    if ma_cn:
        source_query = _branch_detail_query(db, period_key, ma_cn, ma_pgd)
        model = CustomerPeriodBranchDetail
        cust_count_expr = func.count(func.distinct(model.ma_kh))
    else:
        source_query = _base_query(db, period_key, None, ma_pgd)
        model = CustomerPeriodProfile
        cust_count_expr = func.count(model.ma_kh)

    source_query = _restrict_to_matching_customers(source_query, model, customer_ids)
    active_users = db.query(SystemUser).filter(SystemUser.is_active.is_(True)).all()
    credit_codes = {str(user.credit_officer_code).strip() for user in active_users if user.credit_officer_code}
    employee_codes = {str(user.employee_code).strip() for user in active_users if user.employee_code}
    if not credit_codes and not employee_codes:
        return []

    loan_expr = func.coalesce(func.sum(model.so_du_tien_vay), 0)
    deposit_expr = func.coalesce(func.sum(model.so_du_tien_gui), 0)
    casa_expr = func.coalesce(func.sum(model.so_du_tgtt_binh_quan), 0)
    fee_expr = func.coalesce(func.sum(sum(func.coalesce(getattr(model, key), 0) for key in FEE_FIELDS)), 0)
    service_exprs = [
        func.coalesce(func.sum(getattr(model, key)), 0).label(key)
        for key in ACTIVE_SERVICE_KEYS
    ]
    emp_code_expr = getattr(model, "officer_employee_code", model.ma_cb)
    officer_rows = (
        source_query.filter(
            model.ma_cb.isnot(None),
            or_(model.ma_cb.in_(credit_codes | employee_codes), emp_code_expr.in_(employee_codes)),
        )
        .with_entities(
            model.ma_cb,
            model.ten_can_bo,
            emp_code_expr,
            cust_count_expr,
            loan_expr,
            deposit_expr,
            casa_expr,
            fee_expr,
            *service_exprs,
        )
        .group_by(model.ma_cb, model.ten_can_bo, emp_code_expr)
        .all()
    )

    previous_by_officer: dict[str, dict] = {}
    new_customers_by_officer: dict[str, int] = {}
    transferred_customers_by_officer: dict[str, int] = {}
    current_assignments = source_query.filter(
        model.ma_cb.isnot(None),
        or_(model.ma_cb.in_(credit_codes | employee_codes), emp_code_expr.in_(employee_codes)),
    ).with_entities(
        model.ma_cb.label("officer_code"), model.ma_kh.label("customer_code"),
    ).distinct().subquery()

    if previous_period:
        if ma_cn:
            previous_query = _branch_detail_query(db, previous_period, ma_cn, ma_pgd)
            previous_model = CustomerPeriodBranchDetail
            previous_count_expr = func.count(func.distinct(previous_model.ma_kh))
        else:
            previous_query = _base_query(db, previous_period, None, ma_pgd)
            previous_model = CustomerPeriodProfile
            previous_count_expr = func.count(previous_model.ma_kh)
        previous_query = _restrict_to_matching_customers(previous_query, previous_model, customer_ids)
        previous_employee_expr = getattr(previous_model, "officer_employee_code", previous_model.ma_cb)
        previous_valid_condition = or_(
            previous_model.ma_cb.in_(credit_codes | employee_codes),
            previous_employee_expr.in_(employee_codes),
        )
        previous_rows = previous_query.filter(
            previous_model.ma_cb.isnot(None), previous_valid_condition,
        ).with_entities(
            previous_model.ma_cb,
            previous_count_expr,
            func.coalesce(func.sum(previous_model.so_du_tien_gui), 0),
            func.coalesce(func.sum(previous_model.so_du_tgtt_binh_quan), 0),
            func.coalesce(func.sum(previous_model.so_du_tien_vay), 0),
        ).group_by(previous_model.ma_cb).all()
        previous_by_officer = {
            str(code): {
                "custCount": int(customer_count or 0),
                "totalDeposit": float(term_deposit or 0),
                "totalCASA": float(casa or 0),
                "totalLoan": float(loan or 0),
            }
            for code, customer_count, term_deposit, casa, loan in previous_rows
        }
        previous_assignments = previous_query.filter(
            previous_model.ma_cb.isnot(None), previous_valid_condition,
        ).with_entities(
            previous_model.ma_cb.label("officer_code"), previous_model.ma_kh.label("customer_code"),
        ).distinct().subquery()
        same_assignment = and_(
            previous_assignments.c.officer_code == current_assignments.c.officer_code,
            previous_assignments.c.customer_code == current_assignments.c.customer_code,
        )
        new_rows = db.query(
            current_assignments.c.officer_code, func.count(),
        ).outerjoin(previous_assignments, same_assignment).filter(
            previous_assignments.c.customer_code.is_(None),
        ).group_by(current_assignments.c.officer_code).all()
        transferred_rows = db.query(
            previous_assignments.c.officer_code, func.count(),
        ).outerjoin(current_assignments, same_assignment).filter(
            current_assignments.c.customer_code.is_(None),
        ).group_by(previous_assignments.c.officer_code).all()
        new_customers_by_officer = {str(code): int(count or 0) for code, count in new_rows}
        transferred_customers_by_officer = {str(code): int(count or 0) for code, count in transferred_rows}

    officer_leaderboard = []
    for row in officer_rows:
        ma_cb, ten_can_bo, officer_employee_code, cust_count, total_loan_amt, total_deposit, total_casa, total_fee = row[:8]
        used_services_total = sum(float(value or 0) for value in row[8:])
        officer_key = str(ma_cb)
        previous = previous_by_officer.get(officer_key, {})
        current_funding = float(total_deposit or 0) + float(total_casa or 0)
        previous_funding = float(previous.get("totalDeposit", 0)) + float(previous.get("totalCASA", 0))
        officer_leaderboard.append({
            "code": ma_cb,
            "employeeCode": officer_employee_code,
            "name": ten_can_bo or ma_cb,
            "custCount": cust_count,
            "totalLoan": float(total_loan_amt or 0),
            "totalDeposit": float(total_deposit or 0),
            "totalCASA": float(total_casa or 0),
            "totalFee": float(total_fee or 0),
            "avgCrossSell": round(used_services_total / cust_count, 1) if cust_count else 0,
            "previousPeriod": previous_period,
            "previousCustCount": int(previous.get("custCount", 0)),
            "newCustomers": new_customers_by_officer.get(officer_key, 0),
            "transferredCustomers": transferred_customers_by_officer.get(officer_key, 0),
            "depositChange": current_funding - previous_funding,
            "loanChange": float(total_loan_amt or 0) - float(previous.get("totalLoan", 0)),
        })
    officer_leaderboard.sort(key=lambda item: item["totalLoan"], reverse=True)
    return officer_leaderboard


def _count_used_from_mapping(row) -> int:
    return sum(1 for key in ACTIVE_SERVICE_KEYS if _as_number(getattr(row, key, 0)) > 0)


def _as_number(value) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _unused_services_from_mapping(row) -> list[dict]:
    unused = []
    loai = getattr(row, "loai_khach_hang", None) or ""
    priority = CAMPAIGN_GROUP_PRIORITY.get(
        "KHCN" if loai in RETAIL_CUSTOMER_TYPES else "KHDN",
        CAMPAIGN_GROUP_PRIORITY["KHCN"],
    )
    for key in ACTIVE_SERVICE_KEYS:
        if _as_number(getattr(row, key, 0)) == 0:
            unused.append({"key": key, "label": SERVICE_LABELS[key], "group": SERVICE_GROUPS.get(key, "Khác")})
    unused.sort(key=lambda item: priority.index(item["group"]) if item["group"] in priority else 99)
    return unused


def _count_used(row: CustomerPeriodProfile) -> int:
    return sum(1 for key in ACTIVE_SERVICE_KEYS if getattr(row, key, 0) > 0)


def _unused_services(row: CustomerPeriodProfile) -> list[dict]:
    unused = []
    for key in ACTIVE_SERVICE_KEYS:
        if getattr(row, key, 0) == 0:
            unused.append({"key": key, "label": SERVICE_LABELS[key], "group": SERVICE_GROUPS.get(key, "Khác")})
    priority_key = "KHCN" if (row.loai_khach_hang or "") in RETAIL_CUSTOMER_TYPES else "KHDN"
    priority = CAMPAIGN_GROUP_PRIORITY.get(priority_key, CAMPAIGN_GROUP_PRIORITY["KHCN"])
    unused.sort(key=lambda item: priority.index(item["group"]) if item["group"] in priority else 99)
    return unused


def _used_count_expr(model=CustomerPeriodProfile):
    return sum(getattr(model, key) for key in ACTIVE_SERVICE_KEYS)


def _asset_expr(model=CustomerPeriodProfile):
    return (
        func.coalesce(model.so_du_tien_vay, 0)
        + func.coalesce(model.so_du_tien_gui, 0)
        + func.coalesce(model.so_du_tgtt_binh_quan, 0)
    )


def _no_service_condition(model=CustomerPeriodProfile):
    return and_(*(getattr(model, key) == 0 for key in ACTIVE_SERVICE_KEYS))


def _row_to_dict(row: CustomerPeriodProfile) -> dict:
    data = {
        "period_key": row.period_key,
        "ma_kh_chuan": row.ma_kh,
        "ma_cn": row.branch_codes,
        "ma_pgd": row.pgd_codes,
        "ma_kh": row.ma_kh,
        "ten_kh": row.ten_kh,
        "loai_khach_hang": row.loai_khach_hang,
        "so_du_tien_vay": row.so_du_tien_vay,
        "so_du_tien_gui_ckh": row.so_du_tien_gui,
        "loai_vay": row.loai_vay,
        "doanh_so_chuyen_tien_ve_tai_khoan": row.doanh_so_chuyen_tien_ve_tk,
        "so_du_tgtt_binh_quan": row.so_du_tgtt_binh_quan,
        "ma_cb": row.ma_cb,
        "ten_can_bo": row.ten_can_bo,
        "officer_employee_code": row.officer_employee_code,
        "telephone": row.telephone,
    }
    for key in ACTIVE_SERVICE_KEYS:
        data[key] = getattr(row, key)
    return {field: serialize_value(value) for field, value in data.items()}


def _branch_row_to_campaign(period_key: str, branch_code: str, row) -> dict:
    total_assets = (
        _as_number(row.so_du_tien_vay)
        + _as_number(row.so_du_tien_gui)
        + _as_number(row.so_du_tgtt_binh_quan)
    )
    used_count = _count_used_from_mapping(row)
    unused = _unused_services_from_mapping(row)
    data = {
        "period_key": period_key,
        "ma_kh_chuan": row.ma_kh,
        "ma_cn": branch_code,
        "ma_pgd": row.ma_pgd,
        "ma_kh": row.ma_kh,
        "ten_kh": row.ten_kh,
        "loai_khach_hang": row.loai_khach_hang,
        "so_du_tien_vay": row.so_du_tien_vay,
        "so_du_tien_gui_ckh": row.so_du_tien_gui,
        "loai_vay": row.loai_vay,
        "doanh_so_chuyen_tien_ve_tai_khoan": row.doanh_so_cramt,
        "so_du_tgtt_binh_quan": row.so_du_tgtt_binh_quan,
        "ma_cb": row.ma_cb,
        "ten_can_bo": row.ten_can_bo,
        "telephone": None,
        "totalAssets": total_assets,
        "usedCount": used_count,
        "unused": unused,
        "crossSellPct": round((used_count / len(ACTIVE_SERVICE_KEYS)) * 100),
    }
    for key in ACTIVE_SERVICE_KEYS:
        data[key] = getattr(row, key)
    return {field: serialize_value(value) if field not in {"unused"} else value for field, value in data.items()}


def _summary_from_profiles(db: Session, period_key: str, ma_pgd: str | None, customer_ids=None) -> dict:
    """Toàn tỉnh: 1 mã KH = 1 hồ sơ, không cộng trùng theo chi nhánh."""
    query = _base_query(db, period_key, None, ma_pgd)
    if customer_ids is not None:
        query = query.filter(CustomerPeriodProfile.ma_kh.in_(select(customer_ids.c.ma_kh)))
    retail_filter = CustomerPeriodProfile.loai_khach_hang.in_(RETAIL_CUSTOMER_TYPES)
    aggregate_columns = [
        func.count(CustomerPeriodProfile.id),
        func.coalesce(func.sum(CustomerPeriodProfile.so_du_tien_vay), 0),
        func.coalesce(func.sum(CustomerPeriodProfile.so_du_tien_gui), 0),
        func.coalesce(func.sum(CustomerPeriodProfile.so_du_tgtt_binh_quan), 0),
        func.sum(case((CustomerPeriodProfile.so_du_tien_vay > 0, 1), else_=0)),
        func.sum(case((CustomerPeriodProfile.agribank_plus > 0, 1), else_=0)),
        func.sum(case((retail_filter, 1), else_=0)),
        func.coalesce(func.sum(case((retail_filter, CustomerPeriodProfile.so_du_tien_vay), else_=0)), 0),
        func.sum(case((_no_service_condition(), 1), else_=0)),
    ]
    for key in ACTIVE_SERVICE_KEYS:
        base_kind = SERVICE_PENETRATION_BASE.get(key, "all")
        condition = getattr(CustomerPeriodProfile, key) > 0
        if base_kind == "loan":
            condition = and_(condition, CustomerPeriodProfile.so_du_tien_vay > 0)
        elif base_kind == "agribank_plus":
            condition = and_(condition, CustomerPeriodProfile.agribank_plus > 0)
        aggregate_columns.append(func.sum(case((condition, 1), else_=0)))
    aggregate = query.with_entities(*aggregate_columns).one()

    total_customers = int(aggregate[0] or 0)
    totals = aggregate[1:4]
    loan_base = int(aggregate[4] or 0)
    plus_base = int(aggregate[5] or 0)
    service_counts = {
        key: int(aggregate[9 + index] or 0)
        for index, key in enumerate(ACTIVE_SERVICE_KEYS)
    }

    service_penetration = _build_service_penetration_rows(
        count_fn=lambda key, _base_kind: service_counts[key],
        loan_base=loan_base,
        plus_base=plus_base,
        total_customers=total_customers,
    )
    service_penetration = _with_payment_account_penetration(
        service_penetration,
        db=db,
        period_key=period_key,
        ma_cn=None,
        ma_pgd=ma_pgd,
        total_customers=total_customers,
        customer_ids=customer_ids,
    )

    cn_count = int(aggregate[6] or 0)
    dn_count = max(total_customers - cn_count, 0)
    cn_loan = aggregate[7] or 0
    total_loan_amt = float(totals[0] or 0)
    dn_loan = max(total_loan_amt - float(cn_loan or 0), 0)
    no_service_count = int(aggregate[8] or 0)

    campaign_candidates = []
    campaign_rows = (
        query.filter(_asset_expr() >= 80_000_000)
        .filter(_used_count_expr() <= 2)
        .order_by(desc(_asset_expr()), _used_count_expr())
        .limit(50)
        .all()
    )
    for row in campaign_rows:
        total_assets = float(row.so_du_tien_vay or 0) + float(row.so_du_tien_gui or 0) + float(row.so_du_tgtt_binh_quan or 0)
        used_count = _count_used(row)
        unused = _unused_services(row)
        campaign_candidates.append({
            **_row_to_dict(row),
            "totalAssets": total_assets,
            "usedCount": used_count,
            "unused": unused,
            "crossSellPct": round((used_count / len(ACTIVE_SERVICE_KEYS)) * 100),
        })
    campaign_candidates.sort(key=lambda item: (-item["totalAssets"], item["usedCount"]))

    return {
        "kpis": {
            "total_customers": total_customers,
            "total_loan": float(totals[0] or 0),
            "total_deposit": float(totals[1] or 0),
            "total_casa": float(totals[2] or 0),
            "no_service_count": no_service_count,
        },
        "service_penetration": service_penetration,
        "segment": {
            "cn": cn_count,
            "dn": dn_count,
            "cn_loan": float(cn_loan or 0),
            "dn_loan": float(dn_loan or 0),
        },
        "campaign_top5": campaign_candidates[:5],
    }


def _summary_from_branch(db: Session, period_key: str, ma_cn: str, ma_pgd: str | None, customer_ids=None) -> dict:
    """Theo chi nhánh: số liệu lấy từ branch_details, không lấy tổng hồ sơ đa CN."""
    agg = _customer_branch_agg_subquery(db, period_key, ma_cn, ma_pgd)
    if customer_ids is not None:
        agg = db.query(agg).filter(agg.c.ma_kh.in_(select(customer_ids.c.ma_kh))).subquery()

    no_service_condition = and_(*(getattr(agg.c, key) == 0 for key in ACTIVE_SERVICE_KEYS))
    aggregate_columns = [
        func.count(),
        func.coalesce(func.sum(agg.c.so_du_tien_vay), 0),
        func.coalesce(func.sum(agg.c.so_du_tien_gui), 0),
        func.coalesce(func.sum(agg.c.so_du_tgtt_binh_quan), 0),
        func.coalesce(func.sum(case((agg.c.so_du_tien_vay > 0, 1), else_=0)), 0),
        func.coalesce(func.sum(case((agg.c.agribank_plus > 0, 1), else_=0)), 0),
        func.coalesce(func.sum(case((agg.c.loai_khach_hang.in_(RETAIL_CUSTOMER_TYPES), 1), else_=0)), 0),
        func.coalesce(func.sum(case((agg.c.loai_khach_hang.in_(RETAIL_CUSTOMER_TYPES), agg.c.so_du_tien_vay), else_=0)), 0),
        func.coalesce(func.sum(case((no_service_condition, 1), else_=0)), 0),
    ]
    for key in ACTIVE_SERVICE_KEYS:
        condition = getattr(agg.c, key) > 0
        base_kind = SERVICE_PENETRATION_BASE.get(key)
        if base_kind == "loan":
            condition = and_(condition, agg.c.so_du_tien_vay > 0)
        elif base_kind == "agribank_plus":
            condition = and_(condition, agg.c.agribank_plus > 0)
        aggregate_columns.append(func.coalesce(func.sum(case((condition, 1), else_=0)), 0))

    aggregate = db.query(*aggregate_columns).select_from(agg).one()
    total_customers = int(aggregate[0] or 0)
    totals = aggregate[1:4]
    loan_base = int(aggregate[4] or 0)
    plus_base = int(aggregate[5] or 0)
    cn_count = int(aggregate[6] or 0)
    cn_loan = aggregate[7] or 0
    no_service_count = int(aggregate[8] or 0)
    service_counts = {
        key: int(aggregate[9 + index] or 0)
        for index, key in enumerate(ACTIVE_SERVICE_KEYS)
    }

    def _branch_service_count(key: str, _base_kind: str) -> int:
        return service_counts.get(key, 0)

    service_penetration = _build_service_penetration_rows(
        count_fn=_branch_service_count,
        loan_base=loan_base,
        plus_base=plus_base,
        total_customers=total_customers,
    )
    service_penetration = _with_payment_account_penetration(
        service_penetration,
        db=db,
        period_key=period_key,
        ma_cn=ma_cn,
        ma_pgd=ma_pgd,
        total_customers=total_customers,
        customer_ids=customer_ids,
    )

    dn_count = max(total_customers - cn_count, 0)
    total_loan_amt = float(totals[0] or 0)
    dn_loan = max(total_loan_amt - float(cn_loan or 0), 0)

    used_expr = sum(getattr(agg.c, key) for key in ACTIVE_SERVICE_KEYS)
    asset_expr = (
        func.coalesce(agg.c.so_du_tien_vay, 0)
        + func.coalesce(agg.c.so_du_tien_gui, 0)
        + func.coalesce(agg.c.so_du_tgtt_binh_quan, 0)
    )
    campaign_rows = (
        db.query(agg)
        .filter(asset_expr >= 80_000_000)
        .filter(used_expr <= 2)
        .order_by(desc(asset_expr), used_expr)
        .limit(50)
        .all()
    )
    campaign_candidates = [_branch_row_to_campaign(period_key, ma_cn, row) for row in campaign_rows]
    campaign_candidates.sort(key=lambda item: (-item["totalAssets"], item["usedCount"]))

    return {
        "kpis": {
            "total_customers": total_customers,
            "total_loan": float(totals[0] or 0),
            "total_deposit": float(totals[1] or 0),
            "total_casa": float(totals[2] or 0),
            "no_service_count": no_service_count,
        },
        "service_penetration": service_penetration,
        "segment": {
            "cn": cn_count,
            "dn": dn_count,
            "cn_loan": float(cn_loan or 0),
            "dn_loan": float(dn_loan or 0),
        },
        "campaign_top5": campaign_candidates[:5],
    }


@router.get("/summary", dependencies=[Depends(require_any_permission("dashboard:view"))])
def dashboard_summary(
    period_key: str = Query(...),
    filters: dict = Depends(_advanced_filters),
    scope: BranchScope = Depends(get_branch_scope),
    db: Session = Depends(get_db),
):
    # Heavy dashboard requests are queued by the client; allow PostgreSQL to
    # parallelize each aggregate instead of forcing a slow single-worker scan.
    db.execute(text("SET LOCAL max_parallel_workers_per_gather = 2"))
    filters = _scope_filters(scope, filters)
    filter_key = tuple(sorted((key, str(value)) for key, value in filters.items()))
    cache_key = (period_key, scope.ma_cn, f"{scope.ma_pgd or ''}:{filter_key}")
    cached = _SUMMARY_CACHE.get(cache_key)
    if cached and monotonic() - cached[0] < _INSIGHTS_CACHE_TTL_SECONDS:
        return cached[1]

    customer_ids = _customer_filter_ids(db, period_key, scope, filters)
    if scope.ma_cn:
        payload = _summary_from_branch(db, period_key, scope.ma_cn, scope.ma_pgd, customer_ids)
    else:
        payload = _summary_from_profiles(db, period_key, scope.ma_pgd, customer_ids)

    loan_type_breakdown = _build_loan_type_breakdown(
        db, period_key, scope.ma_cn, scope.ma_pgd, customer_ids
    )
    officer_leaderboard = _build_officer_leaderboard(
        db, period_key, scope.ma_cn, scope.ma_pgd, customer_ids
    )

    result = {
        "period_key": period_key,
        "ma_cn": scope.ma_cn,
        "ma_pgd": scope.ma_pgd,
        "kpis": payload["kpis"],
        "service_penetration": payload["service_penetration"],
        "officer_leaderboard": officer_leaderboard,
        "loan_type_breakdown": loan_type_breakdown,
        "segment": payload["segment"],
        "campaign_top5": payload["campaign_top5"],
    }
    _SUMMARY_CACHE[cache_key] = (monotonic(), result)
    return result


def _period_trend_totals(
    db: Session,
    period_key: str,
    ma_cn: str | None,
    ma_pgd: str | None,
    customer_ids=None,
) -> tuple[float, float]:
    """Tổng dư nợ / CASA theo kỳ. Khi chọn chi nhánh thì lấy số liệu trong branch_details."""
    if ma_cn:
        detail_query = _branch_detail_query(db, period_key, ma_cn, ma_pgd)
        detail_query = _restrict_to_matching_customers(
            detail_query, CustomerPeriodBranchDetail, customer_ids
        )
        totals = detail_query.with_entities(
            func.coalesce(func.sum(CustomerPeriodBranchDetail.so_du_tien_vay), 0),
            func.coalesce(func.sum(CustomerPeriodBranchDetail.so_du_tgtt_binh_quan), 0),
        ).one()
        return float(totals[0] or 0), float(totals[1] or 0)

    query = _base_query(db, period_key, None, ma_pgd)
    query = _restrict_to_matching_customers(query, CustomerPeriodProfile, customer_ids)
    totals = query.with_entities(
        func.coalesce(func.sum(CustomerPeriodProfile.so_du_tien_vay), 0),
        func.coalesce(func.sum(CustomerPeriodProfile.so_du_tgtt_binh_quan), 0),
    ).one()
    return float(totals[0] or 0), float(totals[1] or 0)


@router.get("/trends", dependencies=[Depends(require_any_permission("dashboard:view"))])
def dashboard_trends(
    periods: int = Query(default=6, ge=1, le=12),
    scope: BranchScope = Depends(get_branch_scope),
    db: Session = Depends(get_db),
):
    period_keys = [
        item.period_key
        for item in db.query(CustomerPeriodProfile.period_key)
        .distinct()
        .order_by(desc(CustomerPeriodProfile.period_key))
        .limit(periods)
        .all()
    ]
    period_keys.reverse()

    labels = []
    loans = []
    casa_values = []

    for index, period_key in enumerate(period_keys):
        customer_ids = (
            _matching_customer_ids(db, period_key, scope, {"officer_code": scope.officer_code})
            if scope.officer_code else None
        )
        total_loan, total_casa = _period_trend_totals(
            db, period_key, scope.ma_cn, scope.ma_pgd, customer_ids
        )
        is_current = index == len(period_keys) - 1
        month_label = f"T{period_key[4:6]}" if len(period_key) >= 6 else period_key
        if is_current:
            month_label = f"{month_label} (Hiện tại)"
        labels.append(month_label)
        loans.append(round(total_loan / 1e9, 2))
        casa_values.append(round(total_casa / 1e9, 2))

    return {
        "period_keys": period_keys,
        "labels": labels,
        "loans": loans,
        "casa": casa_values,
        "ma_cn": scope.ma_cn,
        "ma_pgd": scope.ma_pgd,
    }


def _analytics_source(db: Session, period_key: str, ma_cn: str | None, ma_pgd: str | None):
    """Return the correctly scoped, non-duplicated source for executive analytics."""
    if ma_cn:
        return _branch_detail_query(db, period_key, ma_cn, ma_pgd), CustomerPeriodBranchDetail
    return _base_query(db, period_key, None, ma_pgd), CustomerPeriodProfile


def _funding_breakdown(
    db: Session,
    period_key: str | None,
    branch_code: str | None,
    customer_ids=None,
) -> dict:
    """Nguồn vốn cuối kỳ từ DP01, loại số dư âm vì đó là thấu chi.

    DP01 được dùng cho tổng nguồn vốn vì bao phủ cả tài khoản OSB; PF14 dùng để
    đối soát cơ cấu kỳ hạn. Số dư ngoại tệ được nhân TY_GIA ngay trên từng dòng.
    """
    empty = {
        "total": 0.0, "demand": 0.0, "non_term": 0.0, "payment": 0.0,
        "term": 0.0, "term_under_6": 0.0, "term_6_to_under_12": 0.0,
        "term_12_plus": 0.0,
    }
    if not period_key:
        return empty
    term_months = func.cast(
        func.coalesce(
            func.nullif(func.regexp_replace(func.coalesce(DP01DepositAccount.month_term, ""), "[^0-9]", "", "g"), ""),
            "0",
        ),
        Integer,
    )
    ccy = func.upper(func.trim(func.coalesce(DP01DepositAccount.ccy, "VND")))
    exchange_rate = case(
        (ccy == "VND", 1),
        else_=func.coalesce(func.nullif(DP01DepositAccount.tygia, 0), 1),
    )
    amount = case(
        (func.coalesce(DP01DepositAccount.current_balance, 0) > 0,
         DP01DepositAccount.current_balance * exchange_rate),
        else_=0,
    )
    type_name = func.lower(func.coalesce(DP01DepositAccount.dp_type_name, ""))
    type_code = func.upper(func.trim(func.coalesce(DP01DepositAccount.dp_type_code, "")))
    non_term_condition = and_(
        term_months == 0,
        or_(
            type_name.like("%kkh%"),
            type_name.like("%không kỳ hạn%"),
            type_name.like("%khong ky han%"),
            type_code == "401",
        ),
    )
    query = db.query(DP01DepositAccount).filter(DP01DepositAccount.period_key == period_key)
    if branch_code:
        query = query.filter(func.trim(DP01DepositAccount.ma_cn) == branch_code)
    if customer_ids is not None:
        query = query.filter(DP01DepositAccount.ma_kh.in_(select(customer_ids.c.ma_kh)))
    row = query.with_entities(
        func.coalesce(func.sum(amount), 0),
        func.coalesce(func.sum(case((term_months == 0, amount), else_=0)), 0),
        func.coalesce(func.sum(case((non_term_condition, amount), else_=0)), 0),
        func.coalesce(func.sum(case((and_(term_months == 0, ~non_term_condition), amount), else_=0)), 0),
        func.coalesce(func.sum(case((term_months > 0, amount), else_=0)), 0),
        func.coalesce(func.sum(case((and_(term_months > 0, term_months < 6), amount), else_=0)), 0),
        func.coalesce(func.sum(case((and_(term_months >= 6, term_months < 12), amount), else_=0)), 0),
        func.coalesce(func.sum(case((term_months >= 12, amount), else_=0)), 0),
    ).one()
    return {
        "total": float(row[0] or 0), "demand": float(row[1] or 0),
        "non_term": float(row[2] or 0), "payment": float(row[3] or 0),
        "term": float(row[4] or 0), "term_under_6": float(row[5] or 0),
        "term_6_to_under_12": float(row[6] or 0), "term_12_plus": float(row[7] or 0),
    }


def _change_payload(current: dict, previous: dict) -> dict:
    result = {}
    for key, value in current.items():
        current_value = float(value or 0)
        previous_value = float(previous.get(key) or 0)
        result[key] = {
            "current": current_value,
            "previous": previous_value,
            "change": current_value - previous_value,
            "change_pct": ((current_value - previous_value) / abs(previous_value) * 100) if previous_value else None,
        }
    return result


def _rr01_recovery(
    db: Session,
    period_key: str | None,
    branch_code: str | None,
    customer_ids=None,
) -> dict:
    if not period_key:
        return {"balance": 0.0, "principal_period": 0.0, "interest_period": 0.0, "period_total": 0.0, "principal_cumulative": 0.0, "interest_cumulative": 0.0, "cumulative": 0.0}
    base = db.query(RR01HandledRiskLoan)
    if branch_code:
        base = base.filter(func.trim(RR01HandledRiskLoan.branch_code) == branch_code)
    if customer_ids is not None:
        base = base.filter(RR01HandledRiskLoan.customer_code.in_(select(customer_ids.c.ma_kh)))
    current = base.filter(RR01HandledRiskLoan.period_key == period_key).with_entities(
        func.coalesce(func.sum(RR01HandledRiskLoan.current_principal), 0),
        func.coalesce(func.sum(RR01HandledRiskLoan.recovered_principal_period), 0),
        func.coalesce(func.sum(RR01HandledRiskLoan.recovered_interest_period), 0),
        func.coalesce(func.sum(
            func.coalesce(RR01HandledRiskLoan.recovered_principal_before_period, 0)
            + func.coalesce(RR01HandledRiskLoan.recovered_principal_period, 0)
        ), 0),
        func.coalesce(func.sum(func.greatest(
            func.coalesce(RR01HandledRiskLoan.original_accrued_interest, 0)
            - func.coalesce(RR01HandledRiskLoan.current_interest, 0),
            0,
        )), 0),
    ).one()
    balance = float(current[0] or 0)
    principal = float(current[1] or 0)
    interest = float(current[2] or 0)
    principal_cumulative = float(current[3] or 0)
    interest_cumulative = float(current[4] or 0)
    return {
        "balance": balance, "principal_period": principal, "interest_period": interest,
        "period_total": principal + interest,
        "principal_cumulative": principal_cumulative, "interest_cumulative": interest_cumulative,
        "cumulative": principal_cumulative + interest_cumulative,
    }


@router.get("/business-analytics", dependencies=[Depends(require_any_permission("analytics:view", "dashboard:view"))])
def dashboard_business_analytics(
    period_key: str = Query(...),
    include_rankings: bool = Query(default=True),
    filters: dict = Depends(_advanced_filters),
    scope: BranchScope = Depends(get_branch_scope),
    db: Session = Depends(get_db),
):
    """Compact aggregate shared by the executive dashboard and four domain pages."""
    db.execute(text("SET LOCAL max_parallel_workers_per_gather = 2"))
    filters = _scope_filters(scope, filters)
    filter_key = tuple(sorted((key, str(value)) for key, value in filters.items()))
    cache_key = (period_key, scope.ma_cn, f"v9-pos:{scope.ma_pgd or ''}:{int(include_rankings)}:{filter_key}")
    cached = _BUSINESS_CACHE.get(cache_key)
    if cached and monotonic() - cached[0] < _INSIGHTS_CACHE_TTL_SECONDS:
        return cached[1]
    shared_key = (*cache_key, _analysis_data_version(db, period_key))
    shared = get_shared_analysis_cache("business-analytics", shared_key)
    if shared is not None:
        _BUSINESS_CACHE[cache_key] = (monotonic(), shared)
        return shared
    query, model = _analytics_source(db, period_key, scope.ma_cn, scope.ma_pgd)
    customer_ids = _customer_filter_ids(db, period_key, scope, filters)
    query = _restrict_to_matching_customers(query, model, customer_ids)
    previous_period = (
        db.query(func.max(CustomerPeriodProfile.period_key))
        .filter(CustomerPeriodProfile.period_key < period_key)
        .scalar()
    )
    # Chỉ tiêu nguồn vốn cấp chi nhánh phải khớp tổng sổ nguồn DP01, kể cả mã
    # đang chờ đối chiếu CIF. Khi người dùng lọc theo KH/PGD thì mới giới hạn
    # tập mã theo hồ sơ C360 tương ứng.
    source_customer_ids = customer_ids if filters or scope.ma_pgd else None
    funding = _funding_breakdown(db, period_key, scope.ma_cn, source_customer_ids)
    previous_funding = _funding_breakdown(db, previous_period, scope.ma_cn, source_customer_ids)
    fee_expr = sum(func.coalesce(getattr(model, field), 0) for field in FEE_FIELDS)
    service_expr = sum(func.coalesce(getattr(model, field), 0) for field in ACTIVE_SERVICE_KEYS)
    service_count_exprs = []
    for key in ACTIVE_SERVICE_KEYS:
        condition = getattr(model, key) > 0
        base_kind = SERVICE_PENETRATION_BASE.get(key)
        if base_kind == "loan":
            condition = and_(condition, func.coalesce(model.so_du_tien_vay, 0) > 0)
        elif base_kind == "agribank_plus":
            condition = and_(condition, model.agribank_plus > 0)
        service_count_exprs.append(func.sum(case((condition, 1), else_=0)))
    fee_count_exprs = [
        func.sum(case((func.coalesce(getattr(model, field), 0) != 0, 1), else_=0))
        for field in FEE_FIELDS
    ]
    totals = query.with_entities(
        func.count(func.distinct(model.ma_kh)),
        func.coalesce(func.sum(model.so_du_tien_gui), 0),
        func.coalesce(func.sum(model.so_du_tgtt_binh_quan), 0),
        func.coalesce(func.sum(model.doanh_so_chuyen_tien_ve_tk if model is CustomerPeriodProfile else model.doanh_so_cramt), 0),
        func.coalesce(func.sum(model.so_du_tien_vay), 0),
        func.coalesce(func.sum(model.du_no_ngan_han), 0),
        func.coalesce(func.sum(model.du_no_trung_dai_han), 0),
        func.coalesce(func.sum(model.du_no_thau_chi), 0),
        func.coalesce(func.sum(model.dprr_chung_tt), 0),
        func.coalesce(func.sum(model.dprr_chung_lk), 0),
        func.coalesce(func.sum(model.dprr_cuthe_tt), 0),
        func.coalesce(func.sum(model.dprr_cuthe_lk), 0),
        func.coalesce(func.sum(getattr(model, "du_no_xlrr", 0)), 0) if model is CustomerPeriodProfile else 0,
        func.coalesce(func.sum(getattr(model, "ds_thu_no_xlrr", 0)), 0) if model is CustomerPeriodProfile else 0,
        *(func.coalesce(func.sum(getattr(model, field)), 0) for field in FEE_FIELDS),
        func.sum(case((service_expr > 0, 1), else_=0)),
        func.sum(case((func.coalesce(model.so_du_tien_gui, 0) + func.coalesce(model.so_du_tgtt_binh_quan, 0) > 0, 1), else_=0)),
        func.sum(case((func.coalesce(model.so_du_tien_vay, 0) > 0, 1), else_=0)),
        *service_count_exprs,
        func.sum(case((model.agribank_plus > 0, 1), else_=0)),
        *fee_count_exprs,
        func.sum(case((and_(
            func.coalesce(model.so_du_tien_gui, 0) + func.coalesce(model.so_du_tgtt_binh_quan, 0) <= 0,
            func.coalesce(model.so_du_tien_vay, 0) <= 0,
            service_expr <= 0,
            fee_expr == 0,
        ), 1), else_=0)),
    ).one()

    retail_condition = model.loai_khach_hang.in_(RETAIL_CUSTOMER_TYPES)
    legal_condition = or_(model.loai_khach_hang.is_(None), ~retail_condition)
    credit_segments = query.with_entities(
        func.coalesce(func.sum(case((retail_condition, model.so_du_tien_vay), else_=0)), 0),
        func.coalesce(func.sum(case((legal_condition, model.so_du_tien_vay), else_=0)), 0),
    ).one()
    previous_credit = {
        "total": 0.0, "individual": 0.0, "legal": 0.0,
        "short_term": 0.0, "medium_long_term": 0.0, "overdraft": 0.0,
    }
    previous_risk = {"general": 0.0, "specific": 0.0, "provision_after_reversal": 0.0}
    if previous_period:
        previous_query, previous_model = _analytics_source(db, previous_period, scope.ma_cn, scope.ma_pgd)
        previous_query = _restrict_to_matching_customers(previous_query, previous_model, customer_ids)
        previous_retail = previous_model.loai_khach_hang.in_(RETAIL_CUSTOMER_TYPES)
        previous_legal = or_(previous_model.loai_khach_hang.is_(None), ~previous_retail)
        previous_values = previous_query.with_entities(
            func.coalesce(func.sum(previous_model.so_du_tien_vay), 0),
            func.coalesce(func.sum(case((previous_retail, previous_model.so_du_tien_vay), else_=0)), 0),
            func.coalesce(func.sum(case((previous_legal, previous_model.so_du_tien_vay), else_=0)), 0),
            func.coalesce(func.sum(previous_model.du_no_ngan_han), 0),
            func.coalesce(func.sum(previous_model.du_no_trung_dai_han), 0),
            func.coalesce(func.sum(previous_model.du_no_thau_chi), 0),
            func.coalesce(func.sum(previous_model.dprr_chung_lk), 0),
            func.coalesce(func.sum(previous_model.dprr_cuthe_lk), 0),
        ).one()
        previous_credit = {
            "total": float(previous_values[0] or 0), "individual": float(previous_values[1] or 0),
            "legal": float(previous_values[2] or 0), "short_term": float(previous_values[3] or 0),
            "medium_long_term": float(previous_values[4] or 0), "overdraft": float(previous_values[5] or 0),
        }
        previous_risk = {
            "general": float(previous_values[6] or 0), "specific": float(previous_values[7] or 0),
            "provision_after_reversal": float((previous_values[6] or 0) + (previous_values[7] or 0)),
        }
    current_credit = {
        "total": float(totals[4] or 0), "individual": float(credit_segments[0] or 0),
        "legal": float(credit_segments[1] or 0), "short_term": float(totals[5] or 0),
        "medium_long_term": float(totals[6] or 0), "overdraft": float(totals[7] or 0),
    }
    rr_recovery = _rr01_recovery(db, period_key, scope.ma_cn, source_customer_ids)
    previous_rr_recovery = _rr01_recovery(db, previous_period, scope.ma_cn, source_customer_ids)
    provision_after_reversal = float((totals[9] or 0) + (totals[11] or 0))

    fee_value_start = 14
    service_any_index = fee_value_start + len(FEE_FIELDS)
    with_deposit_index = service_any_index + 1
    with_loan_index = service_any_index + 2
    service_count_start = service_any_index + 3
    plus_base_index = service_count_start + len(ACTIVE_SERVICE_KEYS)
    fee_count_start = plus_base_index + 1
    without_relationship_index = fee_count_start + len(FEE_FIELDS)
    fee_rows = [
        {
            "key": field,
            "label": FEE_CATEGORY_LABELS[field],
            "value": float(totals[fee_value_start + index] or 0),
            "customers": int(totals[fee_count_start + index] or 0),
        }
        for index, field in enumerate(FEE_FIELDS)
    ]
    for item in fee_rows:
        item["average"] = item["value"] / item["customers"] if item["customers"] else 0
    # Cơ cấu thu phí chỉ dùng phần thu dương làm mẫu số. Khoản âm là hoàn/giảm
    # phí và được giữ nguyên giá trị, nhưng không được làm co sai tỷ trọng các
    # nguồn thu còn lại.
    fee_composition_total = sum(max(float(item["value"] or 0), 0.0) for item in fee_rows)
    for item in fee_rows:
        positive_value = max(float(item["value"] or 0), 0.0)
        item["pct"] = round(positive_value / fee_composition_total * 100, 2) if fee_composition_total else 0

    configured_officer_codes = {
        str(code).strip()
        for user in db.query(SystemUser).filter(SystemUser.is_active.is_(True)).all()
        for code in (user.credit_officer_code, user.employee_code)
        if code and str(code).strip()
    }
    branch_rows_query = db.query(
        CustomerPeriodBranchDetail.branch_code,
        func.count(func.distinct(CustomerPeriodBranchDetail.ma_kh)),
        func.coalesce(func.sum(CustomerPeriodBranchDetail.so_du_tien_gui), 0),
        func.coalesce(func.sum(CustomerPeriodBranchDetail.so_du_tgtt_binh_quan), 0),
        func.coalesce(func.sum(CustomerPeriodBranchDetail.so_du_tien_vay), 0),
        func.coalesce(func.sum(sum(
            func.coalesce(getattr(CustomerPeriodBranchDetail, field), 0)
            for field in FEE_FIELDS
        )), 0),
        func.count(func.distinct(case(
            (CustomerPeriodBranchDetail.ma_cb.in_(configured_officer_codes), CustomerPeriodBranchDetail.ma_cb),
            else_=None,
        ))),
    ).filter(CustomerPeriodBranchDetail.period_key == period_key)
    if scope.ma_cn:
        branch_rows_query = branch_rows_query.filter(CustomerPeriodBranchDetail.branch_code == scope.ma_cn)
    if scope.ma_pgd:
        branch_rows_query = branch_rows_query.filter(CustomerPeriodBranchDetail.ma_pgd == scope.ma_pgd)
    if customer_ids is not None:
        branch_rows_query = branch_rows_query.filter(CustomerPeriodBranchDetail.ma_kh.in_(select(customer_ids.c.ma_kh)))
    branch_rows = branch_rows_query.group_by(CustomerPeriodBranchDetail.branch_code).all()

    loan_group_query = _ln01_query(db, period_key, scope.ma_cn, scope.ma_pgd)
    if customer_ids is not None:
        loan_group_query = loan_group_query.filter(LN01Loan.custseq.in_(select(customer_ids.c.ma_kh)))
    loan_group_rows = (
        loan_group_query
        .with_entities(
            LN01Loan.debt_group,
            func.count(func.distinct(LN01Loan.custseq)),
            func.coalesce(func.sum(LN01Loan.du_no), 0),
        )
        .group_by(LN01Loan.debt_group)
        .order_by(LN01Loan.debt_group)
        .all()
    )

    service_counts = {key: int(totals[service_count_start + index] or 0) for index, key in enumerate(ACTIVE_SERVICE_KEYS)}
    pos_status = query.with_entities(
        func.sum(case((func.coalesce(model.pos_moi, 0) > 0, 1), else_=0)),
        func.sum(case((func.coalesce(model.pos_khong_hoat_dong, 0) > 0, 1), else_=0)),
        func.sum(case((func.coalesce(model.pos_ngung_hoat_dong, 0) > 0, 1), else_=0)),
        func.coalesce(func.sum(model.so_thiet_bi_pos), 0),
    ).one()
    profile_fee_total = float(sum(item["value"] for item in fee_rows))
    fee_reconciliation = _fee_reconciliation(db, period_key, scope, customer_ids, profile_fee_total)
    result = {
        "period_key": period_key,
        "scope": {"branch_code": scope.ma_cn, "pgd_code": scope.ma_pgd},
        "customer": {
            "total": int(totals[0] or 0), "with_deposit": int(totals[with_deposit_index] or 0),
            "with_loan": int(totals[with_loan_index] or 0), "with_service": int(totals[service_any_index] or 0),
            "without_relationship": int(totals[without_relationship_index] or 0),
        },
        "funding": {
            **funding,
            "previous_period": previous_period,
            "previous": previous_funding,
            "changes": _change_payload(funding, previous_funding),
            "identity_check": {
                "components_total": float(funding["term"] + funding["non_term"] + funding["payment"]),
                "difference": float(funding["total"] - funding["term"] - funding["non_term"] - funding["payment"]),
            },
        },
        "deposit": {"term": float(totals[1] or 0), "casa_average": float(totals[2] or 0), "payment_turnover": float(totals[3] or 0)},
        "credit": {
            **current_credit,
            "previous_period": previous_period,
            "previous": previous_credit,
            "changes": _change_payload(current_credit, previous_credit),
            "identity_check": {
                "term_components_total": float((totals[5] or 0) + (totals[6] or 0) + (totals[7] or 0)),
                "term_components_difference": float((totals[4] or 0) - (totals[5] or 0) - (totals[6] or 0) - (totals[7] or 0)),
                "customer_components_total": float((credit_segments[0] or 0) + (credit_segments[1] or 0)),
                "customer_components_difference": float((totals[4] or 0) - (credit_segments[0] or 0) - (credit_segments[1] or 0)),
            },
        },
        "risk": {
            "general_period": float(totals[8] or 0), "general_accumulated": float(totals[9] or 0),
            "specific_period": float(totals[10] or 0), "specific_accumulated": float(totals[11] or 0),
            "provision_after_reversal_accumulated": provision_after_reversal,
            "provision_after_reversal_previous": previous_risk["provision_after_reversal"],
            "provision_after_reversal_change": provision_after_reversal - previous_risk["provision_after_reversal"],
            "written_off_balance": rr_recovery["balance"],
            "written_off_recovery": rr_recovery["period_total"],
            "written_off_recovery_principal": rr_recovery["principal_period"],
            "written_off_recovery_interest": rr_recovery["interest_period"],
            "written_off_recovery_accumulated": rr_recovery["cumulative"],
            "written_off_recovery_principal_accumulated": rr_recovery["principal_cumulative"],
            "written_off_recovery_interest_accumulated": rr_recovery["interest_cumulative"],
            "written_off_recovery_previous_accumulated": previous_rr_recovery["cumulative"],
            "written_off_recovery_accumulated_change": rr_recovery["cumulative"] - previous_rr_recovery["cumulative"],
            "debt_groups": [{"group": str(group or "Chưa xác định"), "customers": int(count or 0), "balance": float(balance or 0)} for group, count, balance in loan_group_rows],
        },
        "income": {
            "total_fee": profile_fee_total,
            "composition_total": float(fee_composition_total),
            "fees": fee_rows,
            "reconciliation": fee_reconciliation,
        },
        "services": _build_service_penetration_rows(
            count_fn=lambda key, _base_kind: service_counts[key],
            loan_base=int(totals[with_loan_index] or 0), plus_base=int(totals[plus_base_index] or 0),
            total_customers=int(totals[0] or 0),
        ),
        "pos": {
            "customers": service_counts.get("pos", 0),
            "devices": int(pos_status[3] or 0),
            "new_customers": int(pos_status[0] or 0),
            "inactive_customers": int(pos_status[1] or 0),
            "stopped_customers": int(pos_status[2] or 0),
        },
        "branches": [{
            "branch_code": row[0], "customers": int(row[1] or 0), "deposit": float(row[2] or 0),
            "casa": float(row[3] or 0), "loan": float(row[4] or 0), "fee": float(row[5] or 0),
            "officers": int(row[6] or 0),
        } for row in branch_rows],
        "officers": _build_officer_leaderboard(
            db, period_key, scope.ma_cn, scope.ma_pgd, source_customer_ids, previous_period
        ) if include_rankings else [],
        "metric_definitions": {
            "term_deposit": {"source": "PF14", "columns": "MONTHLYENDBALANCE, CCY", "formula": "SUM(MONTHLYENDBALANCE × tỷ giá)", "currency": "Quy đổi VNĐ bằng tỷ giá DP01 của kỳ"},
            "funding": {"source": "DP01", "columns": "CURRENT_BALANCE, MONTH_TERM, DP_TYPE_NAME, DP_TYPE_CODE, CCY, TYGIA", "formula": "SUM(CURRENT_BALANCE dương × TYGIA) = CKH + KKH + TGTT; loại số dư âm vì là thấu chi", "currency": "Quy đổi VNĐ trên từng tài khoản bằng TYGIA của DP01"},
            "demand_deposit": {"source": "DP01", "columns": "CURRENT_BALANCE, MONTH_TERM, DP_TYPE_NAME", "formula": "MONTH_TERM = 0; tách KKH theo tên loại tiền gửi, phần còn lại là TGTT", "currency": "Quy đổi VNĐ bằng TYGIA"},
            "deposit": {"source": "PF14/DP01", "columns": "MONTHLYENDBALANCE, CURRENT_BALANCE, AVGBAL/AVERAGEBALANCE", "formula": "Tiền gửi CKH cuối kỳ + TGTT bình quân; DP01 CURRENT_BALANCE âm bị loại vì là thấu chi", "currency": "Quy đổi VNĐ bằng tỷ giá DP01 của kỳ"},
            "casa": {"source": "PF14/DP01", "columns": "AVERAGEBALANCE/AVGBAL, CURRENT_BALANCE, CCY", "formula": "SUM(số dư TKTT bình quân × tỷ giá); không cộng CURRENT_BALANCE âm", "currency": "Quy đổi VNĐ bằng tỷ giá DP01 của kỳ"},
            "payment_turnover": {"source": "GL02", "columns": "LOCAC, CRAMOUNT, CUSTOMER", "formula": "SUM(CRAMOUNT), LOCAC = 421101, giao dịch hợp lệ", "currency": "Giá trị đã chuẩn hóa về VNĐ; không dùng tỷ giá nếu nguồn không có ngoại tệ"},
            "loan": {"source": "PF10/LN01", "columns": "EOMBAL, DU_NO, LNTYPE, CCY", "formula": "SUM(dư nợ ngắn hạn + trung dài hạn + thấu chi)", "currency": "Quy đổi VNĐ bằng tỷ giá DP01 của kỳ"},
            "short_loan": {"source": "PF10/LN01", "columns": "EOMBAL/DU_NO, LNTYPE", "formula": "SUM dư nợ với LNTYPE = 100 hoặc khoản vay ngắn hạn", "currency": "Quy đổi VNĐ bằng tỷ giá DP01 của kỳ"},
            "medium_long_loan": {"source": "PF10/LN01", "columns": "EOMBAL/DU_NO, LNTYPE", "formula": "SUM dư nợ với LNTYPE IN (110,120)", "currency": "Quy đổi VNĐ bằng tỷ giá DP01 của kỳ"},
            "overdraft": {"source": "PF10/LN01", "columns": "EOMBAL/DU_NO, LNTYPE", "formula": "SUM dư nợ với LNTYPE = 241", "currency": "Quy đổi VNĐ bằng tỷ giá DP01 của kỳ"},
            "provision": {"source": "LN01/BC29", "columns": "DU_NO, NHOM_NO, SO_TRICH_LAP_TRONG_KY", "formula": "DPRR chung lũy kế + DPRR cụ thể lũy kế", "currency": "VNĐ theo dữ liệu nguồn"},
            "written_off": {"source": "RR01", "columns": "DUNO_GOC_HIENTAI, DOC_DAUKY_DA_THU_HT, THU_GOC, DUNO_LAI_TICHLUY_BD, DUNO_LAI_HIENTAI", "formula": "Dư nợ XLRR = SUM(DUNO_GOC_HIENTAI); thu gốc LK = DOC_DAUKY_DA_THU_HT + THU_GOC; thu lãi LK = MAX(DUNO_LAI_TICHLUY_BD − DUNO_LAI_HIENTAI, 0)", "currency": "VNĐ theo dữ liệu nguồn"},
            "general_provision": {"source": "LN01", "columns": "DU_NO, NHOM_NO", "formula": "Nhóm 1–4: biến động dư nợ × 0,75%; lũy kế = dư nợ cuối kỳ × 0,75%", "currency": "VNĐ theo dữ liệu nguồn"},
            "specific_provision": {"source": "BC29", "columns": "NHOM_NO, SO_TRICH_LAP_TRONG_KY", "formula": "Nhóm 2–5; trong tháng = kỳ này − kỳ trước, lũy kế = số trích lập cuối kỳ", "currency": "VNĐ theo dữ liệu nguồn"},
            "principal_due": {"source": "LN01", "columns": "NEXT_REPAYMENT_DATE, NEXT_REPAYMENT_AMOUNT", "formula": "SUM gốc có ngày trả nợ trong tháng kế tiếp", "currency": "VNĐ theo dữ liệu nguồn"},
            "interest_due": {"source": "LN01", "columns": "NEXT_INTEREST_REPAYMENT_DATE, TOTAL_INTEREST_REPAYMENT_AMOUNT", "formula": "SUM lãi có ngày trả lãi trong tháng kế tiếp", "currency": "VNĐ theo dữ liệu nguồn"},
            "overdue_interest": {"source": "LN01", "columns": "PASTDUE_INTEREST_AMOUNT", "formula": "SUM lãi quá hạn theo khách hàng trong phạm vi", "currency": "VNĐ theo dữ liệu nguồn"},
            "fee": {"source": "KH02", "columns": "ACCTCD, CRAMT, DRAMT, CUSTSEQ", "formula": "SUM(CRAMT) − SUM(DRAMT) theo nhóm tài khoản cấu hình", "currency": "VNĐ theo dữ liệu nguồn"},
            "service": {"source": "CN05/DP01/Bill Payment/KH02", "columns": "Cờ sản phẩm và mã dịch vụ cấu hình", "formula": "Đếm duy nhất khách hàng có ít nhất một sản phẩm hợp lệ", "currency": "Không áp dụng"},
            "customers": {"source": "Kho CIF + kết quả C360", "columns": "CUSTNO/MA_KH lõi", "formula": "COUNT(DISTINCT mã KH lõi) theo phạm vi", "currency": "Không áp dụng"},
            "total_scale": {"source": "PF14/DP01/PF10/LN01", "columns": "Tiền gửi CKH, TGTT bình quân, dư nợ", "formula": "Tiền gửi CKH + TGTT bình quân + tổng dư nợ", "currency": "Các nguồn ngoại tệ đã quy đổi VNĐ bằng tỷ giá DP01 của kỳ"},
        },
    }
    _BUSINESS_CACHE[cache_key] = (monotonic(), result)
    set_shared_analysis_cache("business-analytics", shared_key, result)
    return result


@router.get("/business-trends", dependencies=[Depends(require_any_permission("analytics:view", "dashboard:view"))])
def dashboard_business_trends(
    periods: int = Query(default=6, ge=2, le=24),
    period_key: str | None = Query(default=None),
    filters: dict = Depends(_advanced_filters),
    scope: BranchScope = Depends(get_branch_scope),
    db: Session = Depends(get_db),
):
    filters = _scope_filters(scope, filters)
    filter_key = tuple(sorted((key, str(value)) for key, value in filters.items()))
    cache_key = ("v4-business-charts-fees", periods, period_key, scope.ma_cn, scope.ma_pgd, filter_key)
    cached = _BUSINESS_TREND_CACHE.get(cache_key)
    if cached and monotonic() - cached[0] < _INSIGHTS_CACHE_TTL_SECONDS:
        return cached[1]
    shared_key = (*cache_key, _analysis_data_version(db, period_key))
    shared = get_shared_analysis_cache("business-trends", shared_key)
    if shared is not None:
        _BUSINESS_TREND_CACHE[cache_key] = (monotonic(), shared)
        return shared
    period_keys = [row[0] for row in (
        db.query(CustomerPeriodProfile.period_key)
        .filter(CustomerPeriodProfile.period_key <= period_key if period_key else True)
        .distinct().order_by(desc(CustomerPeriodProfile.period_key)).limit(periods).all()
    )]
    period_keys.reverse()
    model = CustomerPeriodBranchDetail if scope.ma_cn else CustomerPeriodProfile
    query = db.query(model).filter(model.period_key.in_(period_keys))
    if period_key:
        customer_ids = _customer_filter_ids(db, period_key, scope, filters)
        if customer_ids is not None:
            query = query.filter(model.ma_kh.in_(select(customer_ids.c.ma_kh)))
    if scope.ma_cn:
        query = query.filter(CustomerPeriodBranchDetail.branch_code == scope.ma_cn)
    if scope.ma_pgd:
        query = query.filter(
            CustomerPeriodBranchDetail.ma_pgd == scope.ma_pgd
            if model is CustomerPeriodBranchDetail
            else CustomerPeriodProfile.pgd_codes.ilike(f"%{scope.ma_pgd}%")
        )
    active_users = db.query(SystemUser).filter(SystemUser.is_active.is_(True)).all()
    valid_officer_codes = {
        str(code).strip()
        for user in active_users
        for code in (user.credit_officer_code, user.employee_code)
        if code and str(code).strip()
    }
    employee_code_column = getattr(model, "officer_employee_code", model.ma_cb)
    valid_officer = or_(model.ma_cb.in_(valid_officer_codes), employee_code_column.in_(valid_officer_codes)) if valid_officer_codes else text("1=0")
    aggregate_rows = query.with_entities(
        model.period_key,
        func.count(func.distinct(model.ma_kh)),
        func.coalesce(func.sum(model.so_du_tien_gui), 0),
        func.coalesce(func.sum(model.so_du_tgtt_binh_quan), 0),
        func.coalesce(func.sum(model.so_du_tien_vay), 0),
        *(func.coalesce(func.sum(getattr(model, field)), 0) for field in FEE_FIELDS),
        func.coalesce(func.sum(model.dprr_chung_lk), 0),
        func.coalesce(func.sum(model.dprr_cuthe_lk), 0),
        func.count(func.distinct(case((valid_officer, model.ma_cb), else_=None))),
        func.count(func.distinct(case((valid_officer, model.ma_kh), else_=None))),
        func.coalesce(func.sum(model.du_no_ngan_han), 0),
        func.coalesce(func.sum(model.du_no_trung_dai_han), 0),
        func.coalesce(func.sum(model.du_no_thau_chi), 0),
    ).group_by(model.period_key).all()
    values_by_period = {row[0]: row for row in aggregate_rows}
    source_rows = db.query(ReportSourceStatus.period_key, ReportSourceStatus.source_code, ReportSourceStatus.status).filter(
        ReportSourceStatus.period_key.in_(period_keys)
    ).all()
    ready_sources: dict[str, set[str]] = {}
    for source_period, source_code, source_status in source_rows:
        if source_status in {"ready", "success", "partial"}:
            ready_sources.setdefault(source_period, set()).add(source_code)
    rows = []
    for key in period_keys:
        values = values_by_period.get(key)
        if not values:
            continue
        trend_fee_start = 5
        trend_provision_start = trend_fee_start + len(FEE_FIELDS)
        trend_officer_start = trend_provision_start + 2
        trend_loan_start = trend_officer_start + 2
        officer_count = int(values[trend_officer_start] or 0)
        managed_customers = int(values[trend_officer_start + 1] or 0)
        total_scale = float((values[2] or 0) + (values[3] or 0) + (values[4] or 0))
        total_fee = float(sum(float(value or 0) for value in values[trend_fee_start:trend_provision_start]))
        trend_fees = {field: float(values[trend_fee_start + index] or 0) for index, field in enumerate(FEE_FIELDS)}
        rows.append({
            "period_key": key, "customers": int(values[1] or 0),
            "deposit": float(values[2] or 0), "casa": float(values[3] or 0),
            "loan": float(values[4] or 0),
            "fee": total_fee,
            "fee_guarantee": trend_fees["phi_bao_lanh"],
            "fee_transfer": trend_fees["phi_chuyen_tien"],
            "fee_digital": trend_fees["phi_nhdt"],
            "fee_abic": trend_fees["abic_batd"],
            "fee_fx": trend_fees["phi_kdnt"],
            "fee_lc": trend_fees["phi_lc"],
            "fee_international": trend_fees["phi_ttqt"],
            "fee_card": trend_fees["phi_the"],
            "fee_other": trend_fees["phi_khac"],
            "provision": float((values[trend_provision_start] or 0) + (values[trend_provision_start + 1] or 0)),
            "officers": officer_count,
            "managed_customers": managed_customers,
            "customers_per_officer": managed_customers / officer_count if officer_count else 0,
            "scale_per_officer": total_scale / officer_count if officer_count else 0,
            "deposit_per_officer": float((values[2] or 0) + (values[3] or 0)) / officer_count if officer_count else 0,
            "loan_per_officer": float(values[4] or 0) / officer_count if officer_count else 0,
            "fee_per_officer": total_fee / officer_count if officer_count else 0,
            "short_loan": float(values[trend_loan_start] or 0),
            "medium_long_loan": float(values[trend_loan_start + 1] or 0),
            "overdraft": float(values[trend_loan_start + 2] or 0),
            "availability": {
                "deposit": "PF14" in ready_sources.get(key, set()),
                "casa": "PF14" in ready_sources.get(key, set()),
                "loan": bool({"LN01", "PF10"} & ready_sources.get(key, set())),
                "fee": "KH02" in ready_sources.get(key, set()),
                "provision": bool({"LN01", "BC29"} & ready_sources.get(key, set())),
                "customers": True,
                "officers": True,
                "managed_customers": True,
                "customers_per_officer": True,
                "scale_per_officer": True,
                "deposit_per_officer": True,
                "loan_per_officer": True,
                "fee_per_officer": "KH02" in ready_sources.get(key, set()),
                "short_loan": bool({"LN01", "PF10"} & ready_sources.get(key, set())),
                "medium_long_loan": bool({"LN01", "PF10"} & ready_sources.get(key, set())),
                "overdraft": bool({"LN01", "PF10"} & ready_sources.get(key, set())),
            },
        })
    result = {"items": rows, "scope": {"branch_code": scope.ma_cn, "pgd_code": scope.ma_pgd}}
    _BUSINESS_TREND_CACHE[cache_key] = (monotonic(), result)
    set_shared_analysis_cache("business-trends", shared_key, result)
    return result


BUSINESS_DRILLDOWN_LABELS = {
    "all": "Toàn bộ khách hàng trong phạm vi",
    "deposit": "Khách hàng có tiền gửi",
    "term_deposit": "Khách hàng có tiền gửi có kỳ hạn",
    "funding": "Khách hàng cấu thành tổng nguồn vốn huy động",
    "demand_deposit": "Khách hàng có tiền gửi không kỳ hạn",
    "non_term_deposit": "Khách hàng có tiền gửi KKH",
    "payment_deposit": "Khách hàng có tiền gửi thanh toán",
    "term_under_12": "Khách hàng có tiền gửi CKH dưới 12 tháng",
    "term_under_6": "Khách hàng có tiền gửi CKH dưới 6 tháng",
    "term_6_12": "Khách hàng có tiền gửi CKH từ 6 đến dưới 12 tháng",
    "term_12_plus": "Khách hàng có tiền gửi CKH từ 12 tháng trở lên",
    "casa": "Khách hàng có TGTT bình quân",
    "payment_turnover": "Khách hàng có doanh số TKTT",
    "loan": "Khách hàng có dư nợ",
    "individual_loan": "Khách hàng cá nhân có dư nợ",
    "legal_loan": "Khách hàng pháp nhân có dư nợ",
    "short_loan": "Khách hàng có dư nợ ngắn hạn",
    "medium_long_loan": "Khách hàng có dư nợ trung dài hạn",
    "overdraft": "Khách hàng có dư nợ thấu chi",
    "fee": "Khách hàng phát sinh phí",
    "guarantee_fee": "Khách hàng phát sinh phí bảo lãnh",
    "transfer_fee": "Khách hàng phát sinh phí chuyển tiền",
    "digital_fee": "Khách hàng phát sinh phí NHĐT",
    "abic_fee": "Khách hàng phát sinh phí ABIC/BATĐ",
    "fx_fee": "Khách hàng phát sinh phí kinh doanh ngoại tệ",
    "lc_fee": "Khách hàng phát sinh phí LC",
    "international_fee": "Khách hàng phát sinh phí thanh toán quốc tế",
    "card_fee": "Khách hàng phát sinh phí thẻ",
    "other_fee": "Khách hàng phát sinh phí khác",
    "service": "Khách hàng sử dụng sản phẩm",
    "multi_branch": "Khách hàng có quan hệ đa chi nhánh",
    "no_service": "Khách hàng chưa sử dụng sản phẩm",
    "no_relationship": "Khách hàng chưa phát sinh tiền gửi, tiền vay, sản phẩm dịch vụ hoặc phí",
    "risk": "Khách hàng có dự phòng",
    "provision_accumulated": "Khách hàng có DPRR lũy kế sau hoàn nhập",
    "written_off": "Khách hàng có dư nợ XLRR",
    "general_provision_period": "Chi tiết tăng/giảm DPRR chung trong kỳ",
    "specific_provision_period": "Chi tiết tăng/giảm DPRR cụ thể trong kỳ",
    "new_deposit_account": "Khách hàng có tài khoản tiền gửi mới trong kỳ",
    "closed_deposit_account": "Khách hàng có tài khoản tất toán/ngừng trong kỳ",
    "deposit_drop": "Khách hàng giảm tiền gửi từ 30%",
    "pos_new": "Khách hàng có POS mới trong kỳ",
    "pos_inactive": "Khách hàng có POS không phát sinh giao dịch trong kỳ",
    "pos_stopped": "Khách hàng có POS ngừng hoạt động trong kỳ",
}


def _business_drilldown_condition(model, metric: str, db: Session | None = None, period_key: str | None = None):
    deposit = func.coalesce(model.so_du_tien_gui, 0) + func.coalesce(model.so_du_tgtt_binh_quan, 0)
    fee = sum(func.coalesce(getattr(model, field), 0) for field in FEE_FIELDS)
    services = sum(func.coalesce(getattr(model, field), 0) for field in ACTIVE_SERVICE_KEYS)
    conditions = {
        "all": text("1=1"),
        "deposit": deposit > 0,
        "term_deposit": func.coalesce(model.so_du_tien_gui, 0) > 0,
        "casa": func.coalesce(model.so_du_tgtt_binh_quan, 0) > 0,
        "payment_turnover": func.coalesce(
            model.doanh_so_chuyen_tien_ve_tk if model is CustomerPeriodProfile else model.doanh_so_cramt, 0
        ) > 0,
        "loan": func.coalesce(model.so_du_tien_vay, 0) > 0,
        "individual_loan": and_(model.loai_khach_hang.in_(RETAIL_CUSTOMER_TYPES), func.coalesce(model.so_du_tien_vay, 0) > 0),
        "legal_loan": and_(or_(model.loai_khach_hang.is_(None), ~model.loai_khach_hang.in_(RETAIL_CUSTOMER_TYPES)), func.coalesce(model.so_du_tien_vay, 0) > 0),
        "short_loan": func.coalesce(model.du_no_ngan_han, 0) > 0,
        "medium_long_loan": func.coalesce(model.du_no_trung_dai_han, 0) > 0,
        "overdraft": func.coalesce(model.du_no_thau_chi, 0) > 0,
        "fee": fee != 0,
        "guarantee_fee": func.coalesce(model.phi_bao_lanh, 0) != 0,
        "transfer_fee": func.coalesce(model.phi_chuyen_tien, 0) != 0,
        "digital_fee": func.coalesce(model.phi_nhdt, 0) != 0,
        "abic_fee": func.coalesce(model.abic_batd, 0) != 0,
        "fx_fee": func.coalesce(model.phi_kdnt, 0) != 0,
        "lc_fee": func.coalesce(model.phi_lc, 0) != 0,
        "international_fee": func.coalesce(model.phi_ttqt, 0) != 0,
        "card_fee": func.coalesce(model.phi_the, 0) != 0,
        "other_fee": func.coalesce(model.phi_khac, 0) != 0,
        "service": services > 0,
        "no_service": services == 0,
        "no_relationship": and_(
            deposit <= 0,
            func.coalesce(model.so_du_tien_vay, 0) <= 0,
            services <= 0,
            fee == 0,
        ),
        "risk": (func.abs(func.coalesce(model.dprr_chung_lk, 0)) + func.abs(func.coalesce(model.dprr_cuthe_lk, 0))) > 0,
        "provision_accumulated": (func.abs(func.coalesce(model.dprr_chung_lk, 0)) + func.abs(func.coalesce(model.dprr_cuthe_lk, 0))) > 0,
        "general_provision_period": func.coalesce(model.dprr_chung_tt, 0) != 0,
        "specific_provision_period": func.coalesce(model.dprr_cuthe_tt, 0) != 0,
        "pos_new": func.coalesce(model.pos_moi, 0) > 0,
        "pos_inactive": func.coalesce(model.pos_khong_hoat_dong, 0) > 0,
        "pos_stopped": func.coalesce(model.pos_ngung_hoat_dong, 0) > 0,
    }
    if model is CustomerPeriodProfile:
        conditions["written_off"] = func.coalesce(model.du_no_xlrr, 0) > 0
        conditions["multi_branch"] = func.coalesce(model.branch_count, 0) > 1
    elif metric == "multi_branch" and db is not None and period_key:
        conditions["multi_branch"] = model.ma_kh.in_(
            db.query(CustomerPeriodProfile.ma_kh).filter(
                CustomerPeriodProfile.period_key == period_key,
                CustomerPeriodProfile.branch_count > 1,
            )
        )
    if metric.startswith("service:"):
        key = metric.split(":", 1)[1]
        return func.coalesce(getattr(model, key), 0) > 0 if key in ACTIVE_SERVICE_KEYS else services > 0
    if metric.startswith("no_service:"):
        key = metric.split(":", 1)[1]
        return func.coalesce(getattr(model, key), 0) == 0 if key in ACTIVE_SERVICE_KEYS else services == 0
    if metric.startswith("debt_group:") and period_key:
        group = metric.split(":", 1)[1].strip()
        debt_group_conditions = [
            LN01Loan.period_key == period_key,
            LN01Loan.custseq == model.ma_kh,
        ]
        if group == "Chưa xác định":
            debt_group_conditions.append(or_(LN01Loan.debt_group.is_(None), func.trim(LN01Loan.debt_group) == ""))
        else:
            debt_group_conditions.append(func.trim(LN01Loan.debt_group) == group)
        # Khi xem trong một chi nhánh, chỉ nhận nhóm nợ phát sinh đúng tại chi
        # nhánh đang hiển thị, không kéo nhóm nợ từ quan hệ ở đơn vị khác sang.
        if model is CustomerPeriodBranchDetail:
            debt_group_conditions.append(LN01Loan.brcd == model.branch_code)
        return exists().where(and_(*debt_group_conditions))
    return conditions.get(metric, or_(deposit > 0, func.coalesce(model.so_du_tien_vay, 0) > 0))


FUNDING_DRILLDOWN_METRICS = {
    "funding", "demand_deposit", "non_term_deposit", "payment_deposit",
    "term_under_12", "term_under_6", "term_6_12", "term_12_plus",
}
ACCOUNT_MOVEMENT_METRICS = {"new_deposit_account", "closed_deposit_account"}


def _funding_drilldown(
    db: Session, period_key: str, metric: str, scope: BranchScope, customer_ids,
    detail_keyword: str | None, page: int, page_size: int,
) -> dict:
    term_months = func.cast(func.coalesce(func.nullif(func.regexp_replace(
        func.coalesce(DP01DepositAccount.month_term, ""), "[^0-9]", "", "g"
    ), ""), "0"), Integer)
    ccy = func.upper(func.trim(func.coalesce(DP01DepositAccount.ccy, "VND")))
    rate = case((ccy == "VND", 1), else_=func.coalesce(func.nullif(DP01DepositAccount.tygia, 0), 1))
    amount = case((func.coalesce(DP01DepositAccount.current_balance, 0) > 0,
                   DP01DepositAccount.current_balance * rate), else_=0)
    name = func.lower(func.coalesce(DP01DepositAccount.dp_type_name, ""))
    code = func.upper(func.trim(func.coalesce(DP01DepositAccount.dp_type_code, "")))
    is_kkh = and_(term_months == 0, or_(
        name.like("%kkh%"), name.like("%không kỳ hạn%"), name.like("%khong ky han%"), code == "401",
    ))
    source = db.query(DP01DepositAccount).filter(
        DP01DepositAccount.period_key == period_key,
        func.coalesce(DP01DepositAccount.current_balance, 0) > 0,
    )
    if scope.ma_cn:
        source = source.filter(func.trim(DP01DepositAccount.ma_cn) == scope.ma_cn)
    if customer_ids is not None:
        source = source.filter(DP01DepositAccount.ma_kh.in_(select(customer_ids.c.ma_kh)))
    source = source.with_entities(
        DP01DepositAccount.ma_kh.label("ma_kh"),
        func.trim(DP01DepositAccount.ma_cn).label("branch_code"),
        func.max(DP01DepositAccount.ten_kh).label("ten_kh"),
        func.sum(amount).label("funding"),
        func.sum(case((term_months == 0, amount), else_=0)).label("demand"),
        func.sum(case((is_kkh, amount), else_=0)).label("non_term"),
        func.sum(case((and_(term_months == 0, ~is_kkh), amount), else_=0)).label("payment"),
        func.sum(case((and_(term_months > 0, term_months < 12), amount), else_=0)).label("term_under_12"),
        func.sum(case((and_(term_months > 0, term_months < 6), amount), else_=0)).label("term_under_6"),
        func.sum(case((and_(term_months >= 6, term_months < 12), amount), else_=0)).label("term_6_12"),
        func.sum(case((term_months >= 12, amount), else_=0)).label("term_12_plus"),
        func.count(func.distinct(DP01DepositAccount.so_tai_khoan)).label("account_count"),
    ).group_by(DP01DepositAccount.ma_kh, func.trim(DP01DepositAccount.ma_cn)).subquery()
    value_column = getattr(source.c, {
        "funding": "funding", "demand_deposit": "demand",
        "non_term_deposit": "non_term", "payment_deposit": "payment",
        "term_under_12": "term_under_12", "term_under_6": "term_under_6",
        "term_6_12": "term_6_12", "term_12_plus": "term_12_plus",
    }[metric])
    query = db.query(source).filter(value_column > 0)
    if detail_keyword:
        pattern = f"%{detail_keyword.strip()}%"
        query = query.filter(or_(source.c.ma_kh.ilike(pattern), source.c.ten_kh.ilike(pattern)))
    total = int(query.count())
    total_value = float(query.with_entities(func.coalesce(func.sum(value_column), 0)).scalar() or 0)
    rows = query.order_by(desc(value_column), source.c.ma_kh).offset((page - 1) * page_size).limit(page_size).all()
    codes = [row.ma_kh for row in rows]
    details = {
        row.ma_kh: row for row in db.query(CustomerPeriodBranchDetail).filter(
            CustomerPeriodBranchDetail.period_key == period_key,
            CustomerPeriodBranchDetail.branch_code == scope.ma_cn if scope.ma_cn else text("1=1"),
            CustomerPeriodBranchDetail.ma_kh.in_(codes or ["__NONE__"]),
        ).order_by(CustomerPeriodBranchDetail.ma_kh).all()
    }
    items = []
    for row in rows:
        detail = details.get(row.ma_kh)
        items.append({
            "ma_kh": row.ma_kh, "ten_kh": row.ten_kh or (detail.ten_kh if detail else None),
            "customer_type": detail.loai_khach_hang if detail else None,
            "branch_code": row.branch_code, "officer_code": detail.ma_cb if detail else None,
            "officer_name": detail.ten_can_bo if detail else None,
            "funding": float(row.funding or 0), "deposit": float((row.funding or 0) - (row.demand or 0)),
            "casa": float(row.demand or 0), "non_term": float(row.non_term or 0),
            "payment": float(row.payment or 0), "term_under_12": float(row.term_under_12 or 0),
            "term_under_6": float(row.term_under_6 or 0), "term_6_12": float(row.term_6_12 or 0),
            "term_12_plus": float(row.term_12_plus or 0), "account_count": int(row.account_count or 0),
            "loan": float(detail.so_du_tien_vay or 0) if detail else 0, "fee": 0, "provision": 0,
            "service_count": 0, "written_off": 0, "active_services": [],
            "primary_branch_code": row.branch_code, "branch_codes": row.branch_code, "branch_count": 1,
        })
    return {
        "metric": metric, "label": BUSINESS_DRILLDOWN_LABELS[metric], "total": total,
        "total_value": total_value, "page": page, "page_size": page_size,
        "scope": {"branch_code": scope.ma_cn, "pgd_code": scope.ma_pgd}, "items": items,
    }


def _account_movement_drilldown(
    db: Session, period_key: str, metric: str, scope: BranchScope, customer_ids,
    detail_keyword: str | None, page: int, page_size: int,
) -> dict:
    """Chi tiết khách hàng và số tài khoản PF14 mới/tất toán giữa hai kỳ gần nhất."""
    previous_period = db.query(func.max(PF14AccountBalance.period_key)).filter(
        PF14AccountBalance.period_key < period_key
    ).scalar()
    label = BUSINESS_DRILLDOWN_LABELS[metric]
    if not previous_period:
        return {"metric": metric, "label": label, "total": 0, "total_value": 0, "page": page, "page_size": page_size, "items": []}

    current_accounts = db.query(
        func.coalesce(PF14AccountBalance.trbrcd, PF14AccountBalance.branch_code).label("branch"),
        PF14AccountBalance.custseq.label("customer"),
        PF14AccountBalance.accountno.label("account"),
    ).filter(
        PF14AccountBalance.period_key == period_key,
        PF14AccountBalance.custseq.isnot(None), PF14AccountBalance.accountno.isnot(None),
    )
    previous_accounts = db.query(
        func.coalesce(PF14AccountBalance.trbrcd, PF14AccountBalance.branch_code).label("branch"),
        PF14AccountBalance.custseq.label("customer"),
        PF14AccountBalance.accountno.label("account"),
    ).filter(
        PF14AccountBalance.period_key == previous_period,
        PF14AccountBalance.custseq.isnot(None), PF14AccountBalance.accountno.isnot(None),
    )
    if scope.ma_cn:
        current_accounts = current_accounts.filter(func.coalesce(PF14AccountBalance.trbrcd, PF14AccountBalance.branch_code) == scope.ma_cn)
        previous_accounts = previous_accounts.filter(func.coalesce(PF14AccountBalance.trbrcd, PF14AccountBalance.branch_code) == scope.ma_cn)
    if customer_ids is not None:
        current_accounts = current_accounts.filter(PF14AccountBalance.custseq.in_(select(customer_ids.c.ma_kh)))
        previous_accounts = previous_accounts.filter(PF14AccountBalance.custseq.in_(select(customer_ids.c.ma_kh)))
    current_accounts = current_accounts.distinct().subquery()
    previous_accounts = previous_accounts.distinct().subquery()
    source, target = (current_accounts, previous_accounts) if metric == "new_deposit_account" else (previous_accounts, current_accounts)
    changed = db.query(source.c.branch, source.c.customer, source.c.account).select_from(source).outerjoin(
        target,
        and_(target.c.branch == source.c.branch, target.c.customer == source.c.customer, target.c.account == source.c.account),
    ).filter(target.c.account.is_(None)).subquery()
    grouped = db.query(
        changed.c.customer.label("customer"),
        func.string_agg(func.distinct(changed.c.branch), literal(", ")).label("branches"),
        func.array_agg(func.distinct(changed.c.account)).label("accounts"),
        func.count(func.distinct(changed.c.account)).label("account_count"),
    ).group_by(changed.c.customer).subquery()
    query = db.query(grouped)
    if detail_keyword:
        pattern = f"%{detail_keyword.strip()}%"
        matching_names = db.query(CustomerPeriodProfile.ma_kh).filter(
            CustomerPeriodProfile.period_key == period_key,
            or_(CustomerPeriodProfile.ma_kh.ilike(pattern), CustomerPeriodProfile.ten_kh.ilike(pattern)),
        )
        query = query.filter(grouped.c.customer.in_(matching_names))
    total = int(query.count())
    total_value = int(query.with_entities(func.coalesce(func.sum(grouped.c.account_count), 0)).scalar() or 0)
    rows = query.order_by(desc(grouped.c.account_count), grouped.c.customer).offset((page - 1) * page_size).limit(page_size).all()
    codes = [row.customer for row in rows]
    model = CustomerPeriodBranchDetail if scope.ma_cn else CustomerPeriodProfile
    current_rows = db.query(model).filter(model.period_key == period_key, model.ma_kh.in_(codes or ["__NONE__"]))
    previous_rows = db.query(model).filter(model.period_key == previous_period, model.ma_kh.in_(codes or ["__NONE__"]))
    if scope.ma_cn:
        current_rows = current_rows.filter(CustomerPeriodBranchDetail.branch_code == scope.ma_cn)
        previous_rows = previous_rows.filter(CustomerPeriodBranchDetail.branch_code == scope.ma_cn)
    current_by_customer = {row.ma_kh: row for row in current_rows.all()}
    previous_by_customer = {row.ma_kh: row for row in previous_rows.all()}
    items = []
    for row in rows:
        current = current_by_customer.get(row.customer)
        previous = previous_by_customer.get(row.customer)
        deposit_current = float((getattr(current, "so_du_tien_gui", 0) or 0) + (getattr(current, "so_du_tgtt_binh_quan", 0) or 0))
        deposit_previous = float((getattr(previous, "so_du_tien_gui", 0) or 0) + (getattr(previous, "so_du_tgtt_binh_quan", 0) or 0))
        items.append({
            "ma_kh": row.customer, "ten_kh": getattr(current, "ten_kh", None) or getattr(previous, "ten_kh", None),
            "branch_code": row.branches, "officer_code": getattr(current, "ma_cb", None),
            "officer_name": getattr(current, "ten_can_bo", None), "movement_accounts": sorted(row.accounts or []),
            "movement_account_count": int(row.account_count or 0), "deposit": deposit_current,
            "previous_deposit": deposit_previous, "change": deposit_current - deposit_previous,
        })
    return {
        "metric": metric, "label": label, "total": total, "total_value": total_value,
        "page": page, "page_size": page_size, "previous_period": previous_period, "items": items,
    }


def _deposit_drop_drilldown(
    db: Session, period_key: str, scope: BranchScope, customer_ids,
    detail_keyword: str | None, page: int, page_size: int,
) -> dict:
    previous_period = db.query(func.max(CustomerPeriodProfile.period_key)).filter(
        CustomerPeriodProfile.period_key < period_key
    ).scalar()
    label = BUSINESS_DRILLDOWN_LABELS["deposit_drop"]
    if not previous_period:
        return {"metric": "deposit_drop", "label": label, "total": 0, "total_value": 0, "page": page, "page_size": page_size, "items": []}
    model = CustomerPeriodBranchDetail if scope.ma_cn else CustomerPeriodProfile
    current = aliased(model); previous = aliased(model)
    current_deposit = func.coalesce(current.so_du_tien_gui, 0) + func.coalesce(current.so_du_tgtt_binh_quan, 0)
    previous_deposit = func.coalesce(previous.so_du_tien_gui, 0) + func.coalesce(previous.so_du_tgtt_binh_quan, 0)
    join_conditions = [previous.period_key == previous_period, previous.ma_kh == current.ma_kh]
    if scope.ma_cn:
        join_conditions.append(previous.branch_code == current.branch_code)
        join_conditions.append(func.coalesce(previous.ma_pgd, "") == func.coalesce(current.ma_pgd, ""))
    query = db.query(current, previous).join(previous, and_(*join_conditions)).filter(
        current.period_key == period_key, previous_deposit > 0, current_deposit <= previous_deposit * 0.7,
    )
    if scope.ma_cn:
        query = query.filter(current.branch_code == scope.ma_cn)
    if scope.ma_pgd:
        query = query.filter(current.ma_pgd == scope.ma_pgd)
    if customer_ids is not None:
        query = query.filter(current.ma_kh.in_(select(customer_ids.c.ma_kh)))
    if detail_keyword:
        pattern = f"%{detail_keyword.strip()}%"
        query = query.filter(or_(current.ma_kh.ilike(pattern), current.ten_kh.ilike(pattern)))
    total = int(query.with_entities(func.count(func.distinct(current.ma_kh))).scalar() or 0)
    total_value = float(query.with_entities(func.coalesce(func.sum(current_deposit - previous_deposit), 0)).scalar() or 0)
    rows = query.order_by(current_deposit - previous_deposit, current.ma_kh).offset((page - 1) * page_size).limit(page_size).all()
    items = []
    for current_row, previous_row in rows:
        current_value = float((current_row.so_du_tien_gui or 0) + (current_row.so_du_tgtt_binh_quan or 0))
        previous_value = float((previous_row.so_du_tien_gui or 0) + (previous_row.so_du_tgtt_binh_quan or 0))
        items.append({
            "ma_kh": current_row.ma_kh, "ten_kh": current_row.ten_kh,
            "branch_code": getattr(current_row, "branch_code", None) or getattr(current_row, "primary_branch_code", None),
            "officer_code": current_row.ma_cb, "officer_name": current_row.ten_can_bo,
            "deposit": current_value, "previous_deposit": previous_value,
            "change": current_value - previous_value,
            "change_pct": (current_value - previous_value) / abs(previous_value) * 100 if previous_value else None,
        })
    return {"metric": "deposit_drop", "label": label, "total": total, "total_value": total_value, "page": page, "page_size": page_size, "previous_period": previous_period, "items": items}


@router.get("/business-drilldown", dependencies=[Depends(require_any_permission("dashboard:drilldown", "analytics:view"))])
def dashboard_business_drilldown(
    period_key: str = Query(...),
    metric: str = Query(default="deposit"),
    detail_keyword: str | None = Query(default=None),
    detail_branch_code: str | None = Query(default=None),
    detail_customer_type: str | None = Query(default=None),
    detail_officer: str | None = Query(default=None),
    detail_has_deposit: bool | None = Query(default=None),
    detail_has_loan: bool | None = Query(default=None),
    detail_has_fee: bool | None = Query(default=None),
    detail_multi_branch: bool | None = Query(default=None),
    sort_by: str | None = Query(default=None),
    sort_dir: str = Query(default="desc", pattern="^(asc|desc)$"),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=500),
    filters: dict = Depends(_advanced_filters),
    scope: BranchScope = Depends(get_branch_scope),
    db: Session = Depends(get_db),
):
    filters = _scope_filters(scope, filters)
    query, model = _analytics_source(db, period_key, scope.ma_cn, scope.ma_pgd)
    customer_ids = _matching_customer_ids(db, period_key, scope, filters)
    if metric in FUNDING_DRILLDOWN_METRICS:
        return _funding_drilldown(
            db, period_key, metric, scope, customer_ids if filters or scope.ma_pgd else None,
            detail_keyword, page, page_size,
        )
    if metric in ACCOUNT_MOVEMENT_METRICS:
        return _account_movement_drilldown(
            db, period_key, metric, scope, customer_ids,
            detail_keyword, page, page_size,
        )
    if metric == "deposit_drop":
        return _deposit_drop_drilldown(
            db, period_key, scope, customer_ids,
            detail_keyword, page, page_size,
        )
    query = _restrict_to_matching_customers(query, model, customer_ids)
    query = query.filter(_business_drilldown_condition(model, metric, db, period_key))
    if detail_keyword:
        pattern = f"%{detail_keyword.strip()}%"
        # Tìm khách hàng và tìm cán bộ là hai tiêu chí độc lập.
        query = query.filter(or_(model.ma_kh.ilike(pattern), model.ten_kh.ilike(pattern)))
    if detail_branch_code:
        effective_branch = scope.ma_cn or detail_branch_code.strip()
        branch_column = model.branch_code if model is CustomerPeriodBranchDetail else model.primary_branch_code
        query = query.filter(branch_column == effective_branch)
    if detail_customer_type:
        customer_types = [value.strip() for value in detail_customer_type.split(",") if value.strip()]
        if customer_types:
            query = query.filter(model.loai_khach_hang.in_(customer_types))
    if detail_officer:
        officer_values = [value.strip() for value in detail_officer.split(",") if value.strip()]
        officer_conditions = [model.ma_cb.in_(officer_values)]
        if model is CustomerPeriodBranchDetail:
            officer_conditions.append(model.officer_employee_code.in_(officer_values))
        if model is CustomerPeriodProfile:
            officer_conditions.append(model.ma_kh.in_(
                db.query(CustomerPeriodBranchDetail.ma_kh).filter(
                    CustomerPeriodBranchDetail.period_key == period_key,
                    or_(
                        CustomerPeriodBranchDetail.ma_cb.in_(officer_values),
                        CustomerPeriodBranchDetail.officer_employee_code.in_(officer_values),
                    ),
                )
            ))
        query = query.filter(or_(*officer_conditions))
    deposit_expr = func.coalesce(model.so_du_tien_gui, 0) + func.coalesce(model.so_du_tgtt_binh_quan, 0)
    detail_fee_expr = sum(func.coalesce(getattr(model, field), 0) for field in FEE_FIELDS)
    if detail_has_deposit is not None:
        query = query.filter(deposit_expr > 0 if detail_has_deposit else deposit_expr <= 0)
    if detail_has_loan is not None:
        query = query.filter(func.coalesce(model.so_du_tien_vay, 0) > 0 if detail_has_loan else func.coalesce(model.so_du_tien_vay, 0) <= 0)
    if detail_has_fee is not None:
        query = query.filter(detail_fee_expr != 0 if detail_has_fee else detail_fee_expr == 0)
    if detail_multi_branch is not None:
        if model is CustomerPeriodProfile:
            query = query.filter(func.coalesce(model.branch_count, 0) > 1 if detail_multi_branch else func.coalesce(model.branch_count, 0) <= 1)
        else:
            relationship_query = db.query(CustomerPeriodProfile.ma_kh).filter(
                CustomerPeriodProfile.period_key == period_key,
                CustomerPeriodProfile.branch_count > 1 if detail_multi_branch else CustomerPeriodProfile.branch_count <= 1,
            )
            query = query.filter(model.ma_kh.in_(relationship_query))
    total = query.with_entities(func.count(func.distinct(model.ma_kh))).scalar() or 0
    fee_expr = sum(func.coalesce(getattr(model, field), 0) for field in FEE_FIELDS)
    risk_expr = func.coalesce(model.dprr_chung_lk, 0) + func.coalesce(model.dprr_cuthe_lk, 0)
    service_expr = sum(func.coalesce(getattr(model, field), 0) for field in ACTIVE_SERVICE_KEYS)
    order_expr = {
        "all": func.cast(0, model.so_du_tien_vay.type),
        "deposit": func.coalesce(model.so_du_tien_gui, 0) + func.coalesce(model.so_du_tgtt_binh_quan, 0),
        "term_deposit": model.so_du_tien_gui,
        "casa": model.so_du_tgtt_binh_quan,
        "payment_turnover": model.doanh_so_chuyen_tien_ve_tk if model is CustomerPeriodProfile else model.doanh_so_cramt,
        "loan": model.so_du_tien_vay, "fee": fee_expr, "risk": risk_expr,
        "individual_loan": model.so_du_tien_vay,
        "legal_loan": model.so_du_tien_vay,
        "provision_accumulated": risk_expr,
        "short_loan": model.du_no_ngan_han,
        "medium_long_loan": model.du_no_trung_dai_han,
        "overdraft": model.du_no_thau_chi,
        "guarantee_fee": model.phi_bao_lanh,
        "transfer_fee": model.phi_chuyen_tien,
        "digital_fee": model.phi_nhdt,
        "abic_fee": model.abic_batd,
        "fx_fee": model.phi_kdnt,
        "lc_fee": model.phi_lc,
        "international_fee": model.phi_ttqt,
        "card_fee": model.phi_the,
        "other_fee": model.phi_khac,
        "written_off": getattr(model, "du_no_xlrr", model.so_du_tien_vay),
        "general_provision_period": model.dprr_chung_tt,
        "specific_provision_period": model.dprr_cuthe_tt,
        "pos_new": model.pos_moi,
        "pos_inactive": model.pos_khong_hoat_dong,
        "pos_stopped": model.pos_ngung_hoat_dong,
    }.get(metric, model.so_du_tien_vay if metric.startswith("debt_group:") else service_expr)
    selected_query = query.with_entities(
        model.ma_kh, model.ten_kh, model.loai_khach_hang,
        model.so_du_tien_gui, model.so_du_tgtt_binh_quan, model.so_du_tien_vay,
        fee_expr.label("fee"), risk_expr.label("provision"), service_expr.label("service_count"),
        getattr(model, "du_no_xlrr", func.cast(0, model.so_du_tien_vay.type)).label("written_off"),
        (model.branch_code if model is CustomerPeriodBranchDetail else model.primary_branch_code).label("branch_code"),
        model.ma_cb, model.ten_can_bo,
        (CustomerPeriodProfile.branch_details if model is CustomerPeriodProfile else literal(None)).label("branch_details"),
        model.dprr_chung_tt, model.dprr_chung_lk, model.dprr_cuthe_tt, model.dprr_cuthe_lk,
        *(getattr(model, key) for key in ACTIVE_SERVICE_KEYS),
        model.so_thiet_bi_pos,
        model.pos_moi,
        model.pos_khong_hoat_dong,
        model.pos_ngung_hoat_dong,
    )
    sort_columns = {
        "branch": model.branch_code if model is CustomerPeriodBranchDetail else model.primary_branch_code,
        "customer": model.ten_kh,
        "deposit": model.so_du_tien_gui,
        "casa": model.so_du_tgtt_binh_quan,
        "loan": model.so_du_tien_vay,
        "fee": fee_expr,
        "officer": model.ten_can_bo,
    }
    selected_sort = sort_columns.get(sort_by or "")
    if selected_sort is not None:
        selected_query = selected_query.order_by(
            selected_sort.asc().nullslast() if sort_dir == "asc" else selected_sort.desc().nullslast(),
            model.ma_kh,
        )
    elif metric == "all":
        selected_query = selected_query.order_by(
            desc(case((func.coalesce(model.so_du_tien_gui, 0) > 0, 1), else_=0)),
            desc(case((func.coalesce(model.so_du_tien_vay, 0) > 0, 1), else_=0)),
            desc(case((func.coalesce(model.so_du_tgtt_binh_quan, 0) > 0, 1), else_=0)),
            desc(case((model.ten_can_bo.isnot(None), 1), else_=0)),
            desc(func.coalesce(model.so_du_tien_gui, 0) + func.coalesce(model.so_du_tien_vay, 0) + func.coalesce(model.so_du_tgtt_binh_quan, 0)),
            model.ma_kh,
        )
    else:
        selected_query = selected_query.order_by(desc(order_expr), model.ma_kh)
    selected = selected_query.offset((page - 1) * page_size).limit(page_size).all()
    selected_customer_codes = [row[0] for row in selected]
    debt_group_balance_by_customer: dict[str, float] = {}
    debt_group_total_value: float | None = None
    if metric.startswith("debt_group:"):
        debt_group = metric.split(":", 1)[1].strip()
        debt_group_query = _ln01_query(db, period_key, scope.ma_cn, scope.ma_pgd)
        if customer_ids is not None:
            debt_group_query = debt_group_query.filter(LN01Loan.custseq.in_(select(customer_ids.c.ma_kh)))
        if debt_group == "Chưa xác định":
            debt_group_query = debt_group_query.filter(or_(LN01Loan.debt_group.is_(None), func.trim(LN01Loan.debt_group) == ""))
        else:
            debt_group_query = debt_group_query.filter(func.trim(LN01Loan.debt_group) == debt_group)
        debt_group_total_value = float(
            debt_group_query.with_entities(func.coalesce(func.sum(LN01Loan.du_no), 0)).scalar() or 0
        )
        if selected_customer_codes:
            debt_group_balance_by_customer = {
                str(customer_code): float(balance or 0)
                for customer_code, balance in debt_group_query.filter(
                    LN01Loan.custseq.in_(selected_customer_codes)
                ).with_entities(
                    LN01Loan.custseq,
                    func.coalesce(func.sum(LN01Loan.du_no), 0),
                ).group_by(LN01Loan.custseq).all()
            }
    relationship_by_customer = {
        row.ma_kh: row
        for row in db.query(
            CustomerPeriodProfile.ma_kh,
            CustomerPeriodProfile.primary_branch_code,
            CustomerPeriodProfile.branch_codes,
            CustomerPeriodProfile.branch_count,
        ).filter(
            CustomerPeriodProfile.period_key == period_key,
            CustomerPeriodProfile.ma_kh.in_(selected_customer_codes),
        ).all()
    } if selected_customer_codes else {}
    items = []
    for row in selected:
        officer_code, officer_name = row[11], row[12]
        if not officer_code and not officer_name and isinstance(row[13], list):
            primary_detail = next((
                item for item in row[13]
                if str(item.get("branch_code") or "").strip() == str(row[10] or "").strip()
            ), None)
            if primary_detail:
                officer_code = primary_detail.get("ma_cb") or primary_detail.get("officer_employee_code")
                officer_name = primary_detail.get("ten_can_bo")
        relationship = relationship_by_customer.get(row[0])
        active_services = [
            {"key": key, "label": SERVICE_LABELS.get(key, key)}
            for index, key in enumerate(ACTIVE_SERVICE_KEYS)
            if int(row[18 + index] or 0) > 0
        ]
        pos_offset = 18 + len(ACTIVE_SERVICE_KEYS)
        items.append({
            "ma_kh": row[0], "ten_kh": row[1], "customer_type": row[2],
            "deposit": float(row[3] or 0), "casa": float(row[4] or 0),
            "loan": debt_group_balance_by_customer.get(str(row[0]), float(row[5] or 0)),
            "fee": float(row[6] or 0),
            "provision": float((row[14] if metric == "general_provision_period" else row[16] if metric == "specific_provision_period" else row[7]) or 0),
            "service_count": int(row[8] or 0),
            "so_thiet_bi_pos": int(row[pos_offset] or 0),
            "pos_moi": int(row[pos_offset + 1] or 0),
            "pos_khong_hoat_dong": int(row[pos_offset + 2] or 0),
            "pos_ngung_hoat_dong": int(row[pos_offset + 3] or 0),
            "written_off": float(row[9] or 0), "branch_code": row[10],
            "officer_code": officer_code, "officer_name": officer_name,
            "dprr_chung_tt": float(row[14] or 0), "dprr_chung_lk": float(row[15] or 0),
            "dprr_cuthe_tt": float(row[16] or 0), "dprr_cuthe_lk": float(row[17] or 0),
            "active_services": active_services,
            "primary_branch_code": relationship.primary_branch_code if relationship else row[10],
            "branch_codes": relationship.branch_codes if relationship else row[10],
            "branch_count": int(relationship.branch_count or 0) if relationship else 1,
        })
    if metric.startswith("service:"):
        label = f"Khách hàng đang dùng {SERVICE_LABELS.get(metric.split(':', 1)[1], metric)}"
    elif metric.startswith("no_service:"):
        label = f"Khách hàng chưa dùng {SERVICE_LABELS.get(metric.split(':', 1)[1], metric)}"
    elif metric.startswith("debt_group:"):
        debt_group = metric.split(":", 1)[1]
        label = f"Khách hàng thuộc nhóm nợ {debt_group}" if debt_group != "Chưa xác định" else "Khách hàng chưa xác định nhóm nợ"
    else:
        label = BUSINESS_DRILLDOWN_LABELS.get(metric, metric)
    total_value = debt_group_total_value if debt_group_total_value is not None else query.with_entities(func.coalesce(func.sum(order_expr), 0)).scalar() or 0
    return {
        "metric": metric, "label": label, "total": int(total),
        "total_value": float(total_value), "page": page, "page_size": page_size,
        "scope": {"branch_code": scope.ma_cn, "pgd_code": scope.ma_pgd},
        "items": items,
    }


@router.get("/fee-drilldown", dependencies=[Depends(require_any_permission("analytics:view", "customer:profile:view"))])
def dashboard_fee_drilldown(
    period_key: str = Query(...),
    level: str = Query(default="account", pattern="^(account|customer|transaction)$"),
    category: str | None = Query(default=None),
    account_code: str | None = Query(default=None),
    customer_code: str | None = Query(default=None),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=200),
    filters: dict = Depends(_advanced_filters),
    scope: BranchScope = Depends(get_branch_scope),
    user: CurrentUser = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Trace a fee from category to account, customer and (admin-only) KH02 row."""
    if level == "transaction" and "admin" not in set(user.permissions or []):
        raise HTTPException(status_code=403, detail="Chỉ quản trị viên được xem bút toán KH02 gốc")
    if category and category not in {"classified", "unclassified"} and category not in FEE_CATEGORY_PREFIXES:
        raise HTTPException(status_code=400, detail="Nhóm phí không hợp lệ")

    filters = _scope_filters(scope, filters)
    customer_ids = _matching_customer_ids(db, period_key, scope, filters)
    candidate = _account_prefix_condition(KH02CustomerTransaction.account_code, FEE_CANDIDATE_PREFIXES)
    classified = _all_classified_fee_condition(KH02CustomerTransaction.account_code)
    category_expr = _fee_category_expression(KH02CustomerTransaction.account_code)
    normalized_account = func.substring(func.trim(func.coalesce(KH02CustomerTransaction.account_code, "")), 1, 6)
    normalized_customer = func.trim(func.coalesce(KH02CustomerTransaction.customer_code, ""))
    credit = func.coalesce(KH02CustomerTransaction.credit_amount, 0)
    debit = func.coalesce(KH02CustomerTransaction.debit_amount, 0)
    net = credit - debit

    query = db.query(KH02CustomerTransaction).filter(
        KH02CustomerTransaction.period_key == period_key,
        candidate,
        normalized_customer.in_(select(customer_ids.c.ma_kh)),
    )
    if scope.ma_cn:
        query = query.filter(func.trim(KH02CustomerTransaction.branch_code) == scope.ma_cn)
    if category == "classified":
        query = query.filter(classified)
    elif category == "unclassified":
        query = query.filter(~classified)
    elif category:
        query = query.filter(_fee_category_condition(KH02CustomerTransaction.account_code, category))
    if account_code:
        query = query.filter(normalized_account == str(account_code).strip()[:6])
    if customer_code:
        query = query.filter(normalized_customer == str(customer_code).strip())

    summary = query.with_entities(
        func.count(KH02CustomerTransaction.id),
        func.count(func.distinct(normalized_customer)),
        func.coalesce(func.sum(credit), 0),
        func.coalesce(func.sum(debit), 0),
        func.coalesce(func.sum(net), 0),
    ).one()
    items = []
    total = 0
    if level == "account":
        grouped = query.with_entities(
            normalized_account.label("account_code"),
            category_expr.label("category"),
            func.count(KH02CustomerTransaction.id).label("records"),
            func.count(func.distinct(normalized_customer)).label("customers"),
            func.count(func.distinct(func.trim(KH02CustomerTransaction.branch_code))).label("branches"),
            func.coalesce(func.sum(credit), 0).label("credit"),
            func.coalesce(func.sum(debit), 0).label("debit"),
            func.coalesce(func.sum(net), 0).label("net"),
        ).group_by(normalized_account, category_expr).subquery()
        total = int(db.query(func.count()).select_from(grouped).scalar() or 0)
        rows = db.query(grouped).order_by(desc(func.abs(grouped.c.net)), grouped.c.account_code).offset((page - 1) * page_size).limit(page_size).all()
        items = [{
            "account_code": row.account_code,
            "category": row.category,
            "category_label": FEE_CATEGORY_LABELS.get(row.category, "Chưa phân loại"),
            "records": int(row.records or 0), "customers": int(row.customers or 0), "branches": int(row.branches or 0),
            "credit": float(row.credit or 0), "debit": float(row.debit or 0), "net": float(row.net or 0),
        } for row in rows]
    elif level == "customer":
        grouped = query.with_entities(
            normalized_customer.label("customer_code"),
            func.trim(KH02CustomerTransaction.branch_code).label("branch_code"),
            func.max(KH02CustomerTransaction.customer_name).label("customer_name"),
            func.count(KH02CustomerTransaction.id).label("records"),
            func.coalesce(func.sum(credit), 0).label("credit"),
            func.coalesce(func.sum(debit), 0).label("debit"),
            func.coalesce(func.sum(net), 0).label("net"),
        ).group_by(normalized_customer, func.trim(KH02CustomerTransaction.branch_code)).subquery()
        total = int(db.query(func.count()).select_from(grouped).scalar() or 0)
        rows = db.query(grouped).order_by(desc(func.abs(grouped.c.net)), grouped.c.customer_code).offset((page - 1) * page_size).limit(page_size).all()
        items = [{
            "customer_code": row.customer_code, "customer_name": row.customer_name, "branch_code": row.branch_code,
            "records": int(row.records or 0), "credit": float(row.credit or 0), "debit": float(row.debit or 0), "net": float(row.net or 0),
        } for row in rows]
    else:
        total = int(summary[0] or 0)
        rows = query.join(ImportFile, ImportFile.id == KH02CustomerTransaction.import_file_id).with_entities(
            KH02CustomerTransaction.id,
            KH02CustomerTransaction.transaction_date,
            KH02CustomerTransaction.branch_code,
            KH02CustomerTransaction.customer_code,
            KH02CustomerTransaction.customer_name,
            KH02CustomerTransaction.account_code,
            KH02CustomerTransaction.business_code,
            KH02CustomerTransaction.transaction_code,
            KH02CustomerTransaction.transaction_sequence,
            KH02CustomerTransaction.debit_amount,
            KH02CustomerTransaction.credit_amount,
            ImportFile.original_filename,
        ).order_by(desc(KH02CustomerTransaction.transaction_date), desc(KH02CustomerTransaction.id)).offset((page - 1) * page_size).limit(page_size).all()
        items = [{
            "id": row[0], "transaction_date": row[1].isoformat() if row[1] else None,
            "branch_code": row[2], "customer_code": row[3], "customer_name": row[4], "account_code": row[5],
            "business_code": row[6], "transaction_code": row[7], "transaction_sequence": row[8],
            "debit": float(row[9] or 0), "credit": float(row[10] or 0), "net": float((row[10] or 0) - (row[9] or 0)),
            "source_file": row[11],
        } for row in rows]

    return {
        "period_key": period_key, "level": level, "category": category,
        "category_label": FEE_CATEGORY_LABELS.get(
            category,
            "Phí đã phân loại" if category == "classified" else
            "Phí chưa phân loại" if category == "unclassified" else "Tất cả nhóm phí",
        ),
        "account_code": account_code, "customer_code": customer_code,
        "total": total, "page": page, "page_size": page_size,
        "summary": {"records": int(summary[0] or 0), "customers": int(summary[1] or 0), "credit": float(summary[2] or 0), "debit": float(summary[3] or 0), "net": float(summary[4] or 0)},
        "can_view_transactions": "admin" in set(user.permissions or []),
        "items": items,
    }


@router.get("/business-export", dependencies=[Depends(require_any_permission("dashboard:export", "analytics:export"))])
def dashboard_business_export(
    request: Request,
    period_key: str = Query(...),
    metric: str = Query(default="deposit"),
    detail_keyword: str | None = Query(default=None),
    detail_branch_code: str | None = Query(default=None),
    detail_customer_type: str | None = Query(default=None),
    detail_officer: str | None = Query(default=None),
    sort_by: str | None = Query(default=None),
    sort_dir: str = Query(default="desc", pattern="^(asc|desc)$"),
    filters: dict = Depends(_advanced_filters),
    scope: BranchScope = Depends(get_branch_scope),
    db: Session = Depends(get_db),
    user: CurrentUser = Depends(get_current_user),
):
    export_task = begin_export(request, user, "Đang truy vấn danh sách phân tích")
    filters = _scope_filters(scope, filters)
    payload = dashboard_business_drilldown(
        period_key=period_key, metric=metric, detail_keyword=detail_keyword,
        detail_branch_code=detail_branch_code, detail_customer_type=detail_customer_type,
        detail_officer=detail_officer, sort_by=sort_by, sort_dir=sort_dir,
        page=1, page_size=50_000,
        filters=filters, scope=scope, db=db,
    )
    update_export(export_task, 55, f"Đã lọc {payload['total']:,} khách hàng; đang tạo Excel")
    workbook = Workbook()
    info = workbook.active
    info.title = "Thong tin"
    info.append(["BÁO CÁO PHÂN TÍCH KHÁCH HÀNG C360"])
    info.merge_cells("A1:D1")
    info["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    info["A1"].fill = PatternFill("solid", fgColor="8F1438")
    info["A1"].alignment = Alignment(horizontal="center")
    info.append(["Chỉ tiêu", BUSINESS_DRILLDOWN_LABELS.get(metric, metric)])
    info.append(["Kỳ dữ liệu", period_key])
    info.append(["Chi nhánh", scope.ma_cn or "Toàn hệ thống"])
    info.append(["Phòng ban", scope.ma_pgd or "Tất cả"])
    info.append(["Từ khóa khách hàng", detail_keyword or "Không áp dụng"])
    info.append(["Tổng số bản ghi theo bộ lọc", payload["total"]])
    info.append(["Số bản ghi trong file", len(payload["items"])])
    if payload["total"] > len(payload["items"]):
        info.append(["Lưu ý", "File giới hạn 50.000 dòng để bảo đảm hiệu năng; hãy thu hẹp chi nhánh/PGD trước khi xuất."])
    sheet = workbook.create_sheet()
    sheet.title = "Phan tich nghiep vu"
    if metric in ACCOUNT_MOVEMENT_METRICS:
        headers = ["Mã KH", "Tên khách hàng", "Chi nhánh", "Cán bộ", "Số tài khoản biến động", "Danh sách tài khoản", "Tiền gửi kỳ trước", "Tiền gửi kỳ này", "Biến động số dư"]
        export_rows = [[
            item.get("ma_kh"), item.get("ten_kh"), item.get("branch_code"), item.get("officer_name") or item.get("officer_code"),
            item.get("movement_account_count"), ", ".join(item.get("movement_accounts") or []), item.get("previous_deposit"), item.get("deposit"), item.get("change"),
        ] for item in payload["items"]]
        money_start = 7
    elif metric == "deposit_drop":
        headers = ["Mã KH", "Tên khách hàng", "Chi nhánh", "Cán bộ", "Tiền gửi kỳ trước", "Tiền gửi kỳ này", "Biến động số dư", "Tỷ lệ biến động (%)"]
        export_rows = [[
            item.get("ma_kh"), item.get("ten_kh"), item.get("branch_code"), item.get("officer_name") or item.get("officer_code"),
            item.get("previous_deposit"), item.get("deposit"), item.get("change"), item.get("change_pct"),
        ] for item in payload["items"]]
        money_start = 5
    else:
        headers = ["Mã KH", "Tên khách hàng", "Loại KH", "Chi nhánh", "Cán bộ", "Tiền gửi CKH", "TGTT bình quân", "Dư nợ", "Thu phí", "Dự phòng", "Số SP", "Dư nợ XLRR"]
        export_rows = [[item.get("ma_kh"), item.get("ten_kh"), item.get("customer_type"), item.get("branch_code"), item.get("officer_name") or item.get("officer_code"), item.get("deposit"), item.get("casa"), item.get("loan"), item.get("fee"), item.get("provision"), item.get("service_count"), item.get("written_off")] for item in payload["items"]]
        money_start = 6
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="8F1438")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, export_row in enumerate(export_rows, 1):
        sheet.append(export_row)
        if index % 1_000 == 0 or index == len(export_rows):
            update_export(export_task, 55 + int(35 * index / max(len(export_rows), 1)), f"Đã ghi {index:,}/{len(export_rows):,} dòng")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for row_index in range(2, sheet.max_row + 1):
        if row_index % 2 == 0:
            for cell in sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor="FFF7F9")
        for column_index in range(money_start, len(headers) + 1):
            sheet.cell(row_index, column_index).number_format = '#,##0.00'
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 36)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    update_export(export_task, 96, "Đã tạo file; đang gửi về trình duyệt")
    filename = f"phan_tich_{metric}_{period_key}.xlsx"
    record_security_event(
        request,
        user,
        "business_export",
        "business_export",
        filename,
        "Xuất báo cáo phân tích nghiệp vụ",
        {
            "period_key": period_key,
            "metric": metric,
            "branch_code": scope.ma_cn,
            "department_code": scope.ma_pgd,
            "row_count": len(export_rows),
            "result_total": payload.get("total", 0),
        },
    )
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={
        "Content-Disposition": f'attachment; filename="{filename}"',
        "Content-Length": str(output.getbuffer().nbytes),
    })


@router.get("/table-export", dependencies=[Depends(require_any_permission("dashboard:export"))])
def dashboard_table_export(
    request: Request,
    period_key: str = Query(...),
    dataset: str = Query(pattern="^(branches|anomalies|top_changes)$"),
    anomaly_keyword: str | None = Query(default=None),
    anomaly_type: str | None = Query(default=None),
    filters: dict = Depends(_advanced_filters),
    scope: BranchScope = Depends(get_branch_scope),
    db: Session = Depends(get_db),
    user: CurrentUser = Depends(get_current_user),
):
    """Xuất các bảng điều hành theo đúng phạm vi và bộ lọc chung."""
    export_task = begin_export(request, user, "Đang truy vấn bảng điều hành")
    filters = _scope_filters(scope, filters)
    if dataset == "branches":
        result = dashboard_business_analytics(
            period_key=period_key, include_rankings=False,
            filters=filters, scope=scope, db=db,
        )
        headers = ["Chi nhánh", "Khách hàng", "Tiền gửi CKH", "TGTT bình quân", "Dư nợ", "Thu phí", "Cán bộ"]
        rows = [[item.get("branch_code"), item.get("customers"), item.get("deposit"), item.get("casa"), item.get("loan"), item.get("fee"), item.get("officers")] for item in result.get("branches", [])]
        title = "Kết quả theo chi nhánh"
    elif dataset == "anomalies":
        result = dashboard_insights(
            period_key=period_key, include_top_changes=False, anomalies_only=True,
            anomaly_page=1, anomaly_page_size=50_000, anomaly_include_total=True,
            anomaly_keyword=anomaly_keyword, anomaly_type=anomaly_type,
            filters=filters, scope=scope, db=db,
        )
        headers = ["Chi nhánh", "Mã KH", "Tên khách hàng", "Cán bộ", "Cảnh báo", "Tiền gửi kỳ này", "Tiền gửi kỳ trước", "Dư nợ kỳ này", "Dư nợ kỳ trước"]
        rows = [[item.get("viewing_branch_code") or item.get("primary_branch_code"), item.get("ma_kh"), item.get("ten_kh"), item.get("viewing_officer_name") or item.get("viewing_officer_code"), ", ".join(item.get("flags") or []), item.get("deposit"), item.get("previous_deposit"), item.get("loan"), item.get("previous_loan")] for item in result.get("abnormal", {}).get("items", [])]
        title = "Khách hàng có biến động bất thường"
    else:
        result = dashboard_insights(
            period_key=period_key, include_top_changes=True, anomalies_only=False,
            anomaly_page=1, anomaly_page_size=12, anomaly_include_total=False,
            anomaly_keyword=None, anomaly_type=None,
            filters=filters, scope=scope, db=db,
        )
        headers = ["Nhóm biến động", "Chi nhánh", "Mã KH", "Tên khách hàng", "Kỳ trước", "Kỳ này", "Chênh lệch", "Tỷ lệ (%)"]
        rows = []
        labels = {"deposit_increase": "Tăng tiền gửi", "deposit_decrease": "Giảm tiền gửi", "loan_increase": "Tăng dư nợ", "fee": "Thu phí lớn"}
        for key, items in result.get("top_changes", {}).items():
            rows.extend([[labels.get(key, key), item.get("primary_branch_code"), item.get("ma_kh"), item.get("ten_kh"), item.get("previous"), item.get("current"), item.get("change"), item.get("change_pct")] for item in items])
        title = "Top biến động trong kỳ"

    update_export(export_task, 55, f"Đã lọc {len(rows):,} dòng; đang tạo Excel")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bao cao"
    sheet.append([title])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="8F1438")
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.append([f"Kỳ {period_key}", f"Chi nhánh: {scope.ma_cn or 'Toàn hệ thống'}", f"Phòng ban: {scope.ma_pgd or 'Tất cả'}"])
    sheet.append(headers)
    for cell in sheet[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="A51F40")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index, values in enumerate(rows, start=4):
        sheet.append(values)
        if (row_index - 3) % 1_000 == 0 or row_index - 3 == len(rows):
            update_export(export_task, 55 + int(35 * (row_index - 3) / max(len(rows), 1)), f"Đã ghi {row_index - 3:,}/{len(rows):,} dòng")
        if row_index % 2 == 0:
            for cell in sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor="FFF7F9")
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:{sheet.cell(3, len(headers)).coordinate}"
    sheet.sheet_view.showGridLines = False
    for column_index, column in enumerate(sheet.columns, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
    output = BytesIO(); workbook.save(output); output.seek(0)
    update_export(export_task, 96, "Đã tạo file; đang gửi về trình duyệt")
    filename = f"c360_{dataset}_{period_key}.xlsx"
    record_security_event(
        request,
        user,
        "dashboard_export",
        "dashboard_export",
        filename,
        "Xuất bảng dữ liệu điều hành",
        {
            "period_key": period_key,
            "dataset": dataset,
            "branch_code": scope.ma_cn,
            "department_code": scope.ma_pgd,
            "row_count": len(rows),
        },
    )
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={
        "Content-Disposition": f'attachment; filename="{filename}"',
        "Content-Length": str(output.getbuffer().nbytes),
    })


@router.get("/insights", dependencies=[Depends(require_any_permission("dashboard:view"))])
def dashboard_insights(
    period_key: str = Query(...),
    include_top_changes: bool = Query(default=False),
    anomalies_only: bool = Query(default=False),
    anomaly_page: int = Query(default=1, ge=1),
    anomaly_page_size: int = Query(default=12, ge=5, le=100),
    anomaly_include_total: bool = Query(default=True),
    anomaly_keyword: str | None = Query(default=None),
    anomaly_type: str | None = Query(default=None),
    filters: dict = Depends(_advanced_filters),
    scope: BranchScope = Depends(get_branch_scope),
    db: Session = Depends(get_db),
):
    db.execute(text("SET LOCAL max_parallel_workers_per_gather = 2"))
    filters = _scope_filters(scope, filters)
    filter_key = tuple(sorted((key, str(value)) for key, value in filters.items()))
    cache_key = (period_key, scope.ma_cn, f"fee-v2:{scope.ma_pgd or ''}:top={int(include_top_changes)}:alerts_only={int(anomalies_only)}:{anomaly_page}:{anomaly_page_size}:total={int(anomaly_include_total)}:q={anomaly_keyword or ''}:type={anomaly_type or ''}:{filter_key}")
    cached = _INSIGHTS_CACHE.get(cache_key)
    if cached and monotonic() - cached[0] < _INSIGHTS_CACHE_TTL_SECONDS:
        return cached[1]
    shared_key = (*cache_key, _analysis_data_version(db, period_key))
    shared = get_shared_analysis_cache("insights", shared_key)
    if shared is not None:
        _INSIGHTS_CACHE[cache_key] = (monotonic(), shared)
        return shared

    current = aliased(CustomerPeriodProfile)
    previous = aliased(CustomerPeriodProfile)
    previous_period = (
        db.query(func.max(CustomerPeriodProfile.period_key))
        .filter(CustomerPeriodProfile.period_key < period_key)
        .scalar()
    )
    joined = (
        db.query(current, previous)
        .outerjoin(
            previous,
            and_(
                previous.period_key == previous_period,
                previous.ma_kh == current.ma_kh,
            ),
        )
        .filter(current.period_key == period_key)
    )
    customer_ids = _customer_filter_ids(db, period_key, scope, filters)
    if customer_ids is not None:
        joined = joined.filter(current.ma_kh.in_(select(customer_ids.c.ma_kh)))
    if scope.ma_cn:
        joined = joined.filter(current.branch_codes.ilike(f"%{scope.ma_cn}%"))
    # Phạm vi phòng ban đã được xác định trong customer_ids bằng cán bộ/user hợp lệ.
    # Không lọc lại pgd_codes vì trường này có thể phản ánh đơn vị phát sinh nguồn.

    current_deposit = (
        func.coalesce(current.so_du_tien_gui, 0)
        + func.coalesce(current.so_du_tgtt_binh_quan, 0)
    )
    previous_deposit = (
        func.coalesce(previous.so_du_tien_gui, 0)
        + func.coalesce(previous.so_du_tgtt_binh_quan, 0)
    )
    current_services = sum(func.coalesce(getattr(current, key), 0) for key in ACTIVE_SERVICE_KEYS)
    previous_services = sum(func.coalesce(getattr(previous, key), 0) for key in ACTIVE_SERVICE_KEYS)

    overview = joined.with_entities(
        func.count(current.id),
        func.sum(case((current_deposit > 0, 1), else_=0)),
        func.sum(case((func.coalesce(current.so_du_tien_vay, 0) > 0, 1), else_=0)),
        func.sum(case((current_services > 0, 1), else_=0)),
        func.sum(case((func.coalesce(current.branch_count, 0) > 1, 1), else_=0)),
        func.coalesce(func.sum(current_deposit), 0),
        func.coalesce(func.sum(previous_deposit), 0),
        func.coalesce(func.sum(current.so_du_tien_vay), 0),
        func.coalesce(func.sum(previous.so_du_tien_vay), 0),
        func.sum(case((and_(previous.id.isnot(None), current_deposit <= previous_deposit * 0.7, previous_deposit > 0), 1), else_=0)),
        func.sum(case((and_(previous.id.isnot(None), func.coalesce(current.so_du_tien_vay, 0) >= func.coalesce(previous.so_du_tien_vay, 0) * 1.3, func.coalesce(previous.so_du_tien_vay, 0) > 0), 1), else_=0)),
        func.sum(case((and_(previous.id.isnot(None), current_services < previous_services), 1), else_=0)),
    ).one()

    scoped_change_counts = None
    if scope.ma_cn and previous_period:
        current_detail = aliased(CustomerPeriodBranchDetail)
        previous_detail = aliased(CustomerPeriodBranchDetail)
        scoped_changes = (
            db.query(current_detail, previous_detail)
            .outerjoin(previous_detail, and_(
                previous_detail.period_key == previous_period,
                previous_detail.ma_kh == current_detail.ma_kh,
                previous_detail.branch_code == current_detail.branch_code,
                func.coalesce(previous_detail.ma_pgd, '') == func.coalesce(current_detail.ma_pgd, ''),
            ))
            .filter(
                current_detail.period_key == period_key,
                current_detail.branch_code == scope.ma_cn,
                current_detail.ma_kh.in_(select(customer_ids.c.ma_kh)),
            )
        )
        # customer_ids đã giới hạn khách hàng thuộc phòng quản lý được cấu hình.
        current_detail_deposit = func.coalesce(current_detail.so_du_tien_gui, 0) + func.coalesce(current_detail.so_du_tgtt_binh_quan, 0)
        previous_detail_deposit = func.coalesce(previous_detail.so_du_tien_gui, 0) + func.coalesce(previous_detail.so_du_tgtt_binh_quan, 0)
        current_detail_services = sum(func.coalesce(getattr(current_detail, key), 0) for key in ACTIVE_SERVICE_KEYS)
        previous_detail_services = sum(func.coalesce(getattr(previous_detail, key), 0) for key in ACTIVE_SERVICE_KEYS)
        scoped_change_counts = scoped_changes.with_entities(
            func.sum(case((and_(previous_detail.id.isnot(None), previous_detail_deposit > 0, current_detail_deposit <= previous_detail_deposit * 0.7), 1), else_=0)),
            func.sum(case((and_(previous_detail.id.isnot(None), func.coalesce(previous_detail.so_du_tien_vay, 0) > 0, func.coalesce(current_detail.so_du_tien_vay, 0) >= func.coalesce(previous_detail.so_du_tien_vay, 0) * 1.3), 1), else_=0)),
            func.sum(case((and_(previous_detail.id.isnot(None), current_detail_services < previous_detail_services), 1), else_=0)),
        ).one()

    anomaly_conditions = {
        "deposit_drop": and_(previous.id.isnot(None), previous_deposit > 0, current_deposit <= previous_deposit * 0.7),
        "loan_increase": and_(
            previous.id.isnot(None),
            func.coalesce(previous.so_du_tien_vay, 0) > 0,
            func.coalesce(current.so_du_tien_vay, 0) >= func.coalesce(previous.so_du_tien_vay, 0) * 1.3,
        ),
        "service_drop": and_(previous.id.isnot(None), current_services < previous_services),
    }
    anomaly_condition = anomaly_conditions[anomaly_type] if anomaly_type in anomaly_conditions else or_(*anomaly_conditions.values())
    anomaly_query = joined.filter(anomaly_condition)
    if anomaly_keyword:
        keyword_pattern = f"%{anomaly_keyword.strip()}%"
        anomaly_query = anomaly_query.filter(or_(current.ma_kh.ilike(keyword_pattern), current.ten_kh.ilike(keyword_pattern)))
    anomaly_total = anomaly_query.count() if anomaly_include_total else None
    anomaly_rows = (
        anomaly_query
        .order_by(desc(func.abs(current_deposit - previous_deposit)))
        .offset((anomaly_page - 1) * anomaly_page_size)
        .limit(anomaly_page_size)
        .all()
    )
    scoped_current = {}
    scoped_previous = {}
    if scope.ma_cn and anomaly_rows:
        anomaly_customer_codes = [row.ma_kh for row, _ in anomaly_rows]

        def scoped_finance(target_period):
            if not target_period:
                return {}
            scoped_query = db.query(
                CustomerPeriodBranchDetail.ma_kh,
                func.coalesce(func.sum(CustomerPeriodBranchDetail.so_du_tien_gui), 0),
                func.coalesce(func.sum(CustomerPeriodBranchDetail.so_du_tgtt_binh_quan), 0),
                func.coalesce(func.sum(CustomerPeriodBranchDetail.so_du_tien_vay), 0),
                func.max(CustomerPeriodBranchDetail.ma_cb),
                func.max(CustomerPeriodBranchDetail.ten_can_bo),
            ).filter(
                CustomerPeriodBranchDetail.period_key == target_period,
                CustomerPeriodBranchDetail.branch_code == scope.ma_cn,
                CustomerPeriodBranchDetail.ma_kh.in_(anomaly_customer_codes),
            )
            # Không dùng ma_pgd từ file nguồn để thu hẹp lại danh mục quản lý.
            return {item[0]: item for item in scoped_query.group_by(CustomerPeriodBranchDetail.ma_kh).all()}

        scoped_current = scoped_finance(period_key)
        scoped_previous = scoped_finance(previous_period)
    anomalies = []
    for row, old in anomaly_rows:
        flags = []
        current_scope_row = scoped_current.get(row.ma_kh)
        previous_scope_row = scoped_previous.get(row.ma_kh)
        now_deposit = float((current_scope_row[1] or 0) + (current_scope_row[2] or 0)) if current_scope_row else float((row.so_du_tien_gui or 0) + (row.so_du_tgtt_binh_quan or 0))
        old_deposit = float((previous_scope_row[1] or 0) + (previous_scope_row[2] or 0)) if previous_scope_row else (float((old.so_du_tien_gui or 0) + (old.so_du_tgtt_binh_quan or 0)) if old else 0)
        now_loan = float(current_scope_row[3] or 0) if current_scope_row else float(row.so_du_tien_vay or 0)
        old_loan = float(previous_scope_row[3] or 0) if previous_scope_row else (float(old.so_du_tien_vay or 0) if old else 0)
        if old_deposit > 0 and now_deposit <= old_deposit * 0.7:
            flags.append("Tiền gửi giảm mạnh")
        if old_loan > 0 and now_loan >= old_loan * 1.3:
            flags.append("Dư nợ tăng nhanh")
        now_services = sum(float(getattr(row, key) or 0) for key in ACTIVE_SERVICE_KEYS)
        old_services = sum(float(getattr(old, key) or 0) for key in ACTIVE_SERVICE_KEYS) if old else 0
        if old and now_services < old_services:
            flags.append("Giảm sản phẩm")
        anomalies.append({
            "id": row.id,
            "ma_kh": row.ma_kh,
            "ten_kh": row.ten_kh,
            "primary_branch_code": row.primary_branch_code,
            "viewing_branch_code": scope.ma_cn,
            "viewing_officer_code": current_scope_row[4] if current_scope_row else None,
            "viewing_officer_name": current_scope_row[5] if current_scope_row else None,
            "branch_codes": row.branch_codes,
            "deposit": now_deposit,
            "previous_deposit": old_deposit,
            "loan": now_loan,
            "previous_loan": old_loan,
            "flags": flags,
        })

    if anomalies_only:
        result = {
            "period_key": period_key,
            "previous_period": previous_period,
            "abnormal": {
                **({"total": int(anomaly_total or 0)} if anomaly_total is not None else {}),
                "page": anomaly_page,
                "page_size": anomaly_page_size,
                "items": anomalies,
            },
        }
        _INSIGHTS_CACHE[cache_key] = (monotonic(), result)
        set_shared_analysis_cache("insights", shared_key, result)
        return result

    def serialize_change_rows(rows, value_kind: str):
        items = []
        for row, old in rows:
            if value_kind == "deposit":
                current_value = float((row.so_du_tien_gui or 0) + (row.so_du_tgtt_binh_quan or 0))
                previous_value = float((old.so_du_tien_gui or 0) + (old.so_du_tgtt_binh_quan or 0)) if old else 0
            elif value_kind == "loan":
                current_value = float(row.so_du_tien_vay or 0)
                previous_value = float(old.so_du_tien_vay or 0) if old else 0
            else:
                current_value = float(sum(getattr(row, key) or 0 for key in FEE_FIELDS))
                previous_value = float(sum(getattr(old, key) or 0 for key in FEE_FIELDS)) if old else 0
            items.append({
                "ma_kh": row.ma_kh, "ten_kh": row.ten_kh,
                "primary_branch_code": row.primary_branch_code,
                "current": current_value, "previous": previous_value,
                "change": current_value - previous_value,
                "change_pct": ((current_value - previous_value) / abs(previous_value) * 100) if previous_value else None,
            })
        return items

    deposit_difference = current_deposit - previous_deposit
    loan_difference = func.coalesce(current.so_du_tien_vay, 0) - func.coalesce(previous.so_du_tien_vay, 0)
    top_deposit_increase = joined.filter(previous.id.isnot(None), deposit_difference > 0).order_by(desc(deposit_difference)).limit(10).all() if include_top_changes else []
    top_deposit_decrease = joined.filter(previous.id.isnot(None), deposit_difference < 0).order_by(deposit_difference).limit(10).all() if include_top_changes else []
    top_loan_increase = joined.filter(previous.id.isnot(None), loan_difference > 0).order_by(desc(loan_difference)).limit(10).all() if include_top_changes else []
    fee_current = sum(func.coalesce(getattr(current, key), 0) for key in FEE_FIELDS)
    top_fee = joined.filter(fee_current != 0).order_by(desc(func.abs(fee_current))).limit(10).all() if include_top_changes else []

    pf_previous_period = (
        db.query(func.max(PF14AccountBalance.period_key))
        .filter(PF14AccountBalance.period_key < period_key)
        .scalar()
    )
    new_accounts = 0
    closed_accounts = 0
    if pf_previous_period:
        current_accounts_query = db.query(
            PF14AccountBalance.trbrcd.label("branch"),
            PF14AccountBalance.custseq.label("customer"),
            PF14AccountBalance.accountno.label("account"),
        ).filter(PF14AccountBalance.period_key == period_key)
        previous_accounts_query = db.query(
            PF14AccountBalance.trbrcd.label("branch"),
            PF14AccountBalance.custseq.label("customer"),
            PF14AccountBalance.accountno.label("account"),
        ).filter(PF14AccountBalance.period_key == pf_previous_period)
        if scope.ma_cn:
            current_accounts_query = current_accounts_query.filter(PF14AccountBalance.trbrcd == scope.ma_cn)
            previous_accounts_query = previous_accounts_query.filter(PF14AccountBalance.trbrcd == scope.ma_cn)
        if customer_ids is not None:
            current_accounts_query = current_accounts_query.filter(PF14AccountBalance.custseq.in_(select(customer_ids.c.ma_kh)))
            previous_accounts_query = previous_accounts_query.filter(PF14AccountBalance.custseq.in_(select(customer_ids.c.ma_kh)))
        current_accounts = current_accounts_query.subquery()
        previous_accounts = previous_accounts_query.subquery()
        account_changes = (
            db.query(
                func.sum(case((previous_accounts.c.account.is_(None), 1), else_=0)),
                func.sum(case((current_accounts.c.account.is_(None), 1), else_=0)),
            )
            .select_from(current_accounts)
            .join(
                previous_accounts,
                and_(
                    previous_accounts.c.branch == current_accounts.c.branch,
                    previous_accounts.c.customer == current_accounts.c.customer,
                    previous_accounts.c.account == current_accounts.c.account,
                ),
                full=True,
            )
            .one()
        )
        new_accounts = int(account_changes[0] or 0)
        closed_accounts = int(account_changes[1] or 0)

    ln_query = _ln01_query(db, period_key, scope.ma_cn, scope.ma_pgd)
    if customer_ids is not None:
        ln_query = ln_query.filter(LN01Loan.custseq.in_(select(customer_ids.c.ma_kh)))
    period_date = datetime.strptime(period_key, "%Y%m%d").date()
    next_month_start = date(period_date.year + (period_date.month == 12), 1 if period_date.month == 12 else period_date.month + 1, 1)
    next_month_end = date(next_month_start.year + (next_month_start.month == 12), 1 if next_month_start.month == 12 else next_month_start.month + 1, 1)
    obligations = ln_query.with_entities(
        func.coalesce(func.sum(case((and_(LN01Loan.next_repayment_date >= next_month_start, LN01Loan.next_repayment_date < next_month_end), LN01Loan.next_repayment_amount), else_=0)), 0),
        func.coalesce(func.sum(case((and_(LN01Loan.next_interest_repayment_date >= next_month_start, LN01Loan.next_interest_repayment_date < next_month_end), LN01Loan.total_interest_repayment_amount), else_=0)), 0),
        func.coalesce(func.sum(LN01Loan.pastdue_interest_amount), 0),
        func.count(func.distinct(case((func.coalesce(LN01Loan.pastdue_interest_amount, 0) > 0, LN01Loan.custseq)))),
        func.count(LN01Loan.next_repayment_date),
    ).one()

    result = {
        "period_key": period_key,
        "previous_period": previous_period,
        "customer_overview": {
            "total": int(overview[0] or 0),
            "with_deposit": int(overview[1] or 0),
            "with_loan": int(overview[2] or 0),
            "with_digital_service": int(overview[3] or 0),
            "multi_branch": int(overview[4] or 0),
        },
        "deposit": {
            "total": float(overview[5] or 0),
            "previous_total": float(overview[6] or 0),
            "change": float((overview[5] or 0) - (overview[6] or 0)),
            "new_accounts": int(new_accounts or 0),
            "closed_accounts": int(closed_accounts or 0),
            "large_drop_customers": int((scoped_change_counts[0] if scoped_change_counts else overview[9]) or 0),
        },
        "credit": {
            "total": float(overview[7] or 0),
            "previous_total": float(overview[8] or 0),
            "change": float((overview[7] or 0) - (overview[8] or 0)),
            "principal_due_next_month": float(obligations[0] or 0),
            "interest_due_next_month": float(obligations[1] or 0),
            "overdue_interest": float(obligations[2] or 0),
            "overdue_customers": int(obligations[3] or 0),
            "obligation_source_available": bool(obligations[4]),
        },
        "abnormal": {
            "total": int(anomaly_total or 0),
            "page": anomaly_page,
            "page_size": anomaly_page_size,
            "deposit_drop_count": int((scoped_change_counts[0] if scoped_change_counts else overview[9]) or 0),
            "loan_increase_count": int((scoped_change_counts[1] if scoped_change_counts else overview[10]) or 0),
            "service_drop_count": int((scoped_change_counts[2] if scoped_change_counts else overview[11]) or 0),
            "items": anomalies,
        },
        "top_changes": {
            "deposit_increase": serialize_change_rows(top_deposit_increase, "deposit"),
            "deposit_decrease": serialize_change_rows(top_deposit_decrease, "deposit"),
            "loan_increase": serialize_change_rows(top_loan_increase, "loan"),
            "fee": serialize_change_rows(top_fee, "fee"),
        },
    }
    _INSIGHTS_CACHE[cache_key] = (monotonic(), result)
    set_shared_analysis_cache("insights", shared_key, result)
    return result
