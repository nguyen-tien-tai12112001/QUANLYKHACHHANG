from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.database import SessionLocal
from app.models import (
    OrgBranch,
    OrgDepartment,
    SystemPermission,
    SystemRole,
    SystemRolePermission,
    SystemUser,
)
from app.security import hash_password


ROOT_DIR = Path(__file__).resolve().parents[2]
DOCUMENTS_DIR = ROOT_DIR / "documents"
DEFAULT_PASSWORD = "1"


ROLE_DEFINITIONS = [
    ("ADMIN", "Quản trị hệ thống", "Toàn quyền hệ thống, bao gồm quản trị dữ liệu, tổ chức và phân quyền."),
    ("HEAD_OFFICE_LEADER", "Lãnh đạo Hội sở", "Theo dõi và điều hành toàn Hội sở cùng các chi nhánh loại II trực thuộc."),
    ("BRANCH_MANAGER", "Lãnh đạo chi nhánh loại II", "Theo dõi toàn bộ phòng ban, PGD và khách hàng thuộc chi nhánh."),
    ("DEPARTMENT_MANAGER", "Lãnh đạo phòng/PGD", "Theo dõi cán bộ và khách hàng thuộc phòng ban hoặc PGD được giao."),
    ("USER", "Cán bộ quản lý khách hàng", "Theo dõi các khách hàng được phân công trực tiếp."),
]

PERMISSION_DEFINITIONS = [
    ("dashboard:view", "Xem Dashboard", "Tổng quan"),
    ("dashboard:drilldown", "Xem danh sách chi tiết tạo ra KPI", "Tổng quan"),
    ("dashboard:export", "Xuất dữ liệu Dashboard", "Tổng quan"),
    ("dashboard:health", "Xem trạng thái hệ thống", "Tổng quan"),
    ("customer:view", "Xem danh sách khách hàng", "Khách hàng C360"),
    ("customer:profile:view", "Xem hồ sơ khách hàng", "Khách hàng C360"),
    ("customer:deposit:view", "Xem tiền gửi khách hàng", "Khách hàng C360"),
    ("customer:credit:view", "Xem tiền vay và rủi ro", "Khách hàng C360"),
    ("customer:income:view", "Xem phí và thu nhập", "Khách hàng C360"),
    ("customer:export", "Xuất dữ liệu khách hàng", "Khách hàng C360"),
    ("analytics:view", "Xem phân tích nghiệp vụ", "Phân tích nghiệp vụ"),
    ("analytics:export", "Xuất phân tích nghiệp vụ", "Phân tích nghiệp vụ"),
    ("warehouse:view", "Xem kho dữ liệu", "Kho dữ liệu"),
    ("warehouse:import", "Import dữ liệu", "Kho dữ liệu"),
    ("warehouse:replace", "Thay thế file dữ liệu", "Kho dữ liệu"),
    ("warehouse:delete", "Xóa file dữ liệu", "Kho dữ liệu"),
    ("warehouse:summarize", "Tổng hợp kỳ dữ liệu", "Kho dữ liệu"),
    ("cif:view", "Xem Kho CIF", "Kho CIF"),
    ("cif:import", "Import dữ liệu CIF", "Kho CIF"),
    ("cif:review", "Rà soát và xác nhận xung đột CIF", "Kho CIF"),
    ("cif:override", "Ghi đè thủ công dữ liệu CIF", "Kho CIF"),
    ("processing:view", "Xem trạng thái xử lý khách hàng", "Xử lý dữ liệu"),
    ("processing:run", "Chạy xử lý khách hàng", "Xử lý dữ liệu"),
    ("processing:recover", "Khôi phục job bị lỗi hoặc bị kẹt", "Xử lý dữ liệu"),
    ("reconciliation:view", "Xem khách hàng chưa khớp CIF", "Đối chiếu dữ liệu"),
    ("reconciliation:review", "Cập nhật kết quả đối chiếu CIF", "Đối chiếu dữ liệu"),
    ("mapping:view", "Xem từ điển và mapping", "Quản trị dữ liệu"),
    ("mapping:write", "Cấu hình mapping và công thức", "Quản trị dữ liệu"),
    ("report:view", "Xem báo cáo", "Báo cáo"),
    ("report:summarize", "Chạy tổng hợp báo cáo", "Báo cáo"),
    ("report:export", "Xuất báo cáo", "Báo cáo"),
    ("admin:branch:view", "Xem chi nhánh", "Quản trị hệ thống"),
    ("admin:branch:write", "Thêm/sửa/xóa chi nhánh", "Quản trị hệ thống"),
    ("admin:department:view", "Xem phòng ban", "Quản trị hệ thống"),
    ("admin:department:write", "Thêm/sửa/xóa phòng ban", "Quản trị hệ thống"),
    ("admin:user:view", "Xem người dùng", "Quản trị hệ thống"),
    ("admin:user:write", "Thêm/sửa/xóa người dùng", "Quản trị hệ thống"),
    ("admin:role:view", "Xem nhóm quyền", "Quản trị hệ thống"),
    ("admin:role:write", "Cấu hình nhóm quyền", "Quản trị hệ thống"),
    ("admin:config:view", "Xem cấu hình nghiệp vụ", "Quản trị hệ thống"),
    ("admin:config:write", "Thêm sửa xóa cấu hình nghiệp vụ", "Quản trị hệ thống"),
    ("admin:audit:view", "Xem nhật ký thao tác", "Quản trị hệ thống"),
]


def norm_code(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text.zfill(2) if text.isdigit() and len(text) < 2 else text


def norm_text(value) -> str:
    return str(value or "").strip()


def read_excel_rows(path: Path) -> list[dict]:
    if not path.exists():
        return []
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = [norm_text(value) for value in next(rows)]
    data = []
    for row in rows:
        item = {headers[index]: row[index] if index < len(row) else None for index in range(len(headers))}
        if any(value is not None and norm_text(value) for value in item.values()):
            data.append(item)
    workbook.close()
    return data


def get_or_create_role(db: Session, code: str, name: str, description: str) -> SystemRole:
    role = db.query(SystemRole).filter(SystemRole.role_code == code).first()
    if not role:
        role = SystemRole(role_code=code, role_name=name, description=description, is_system=True)
        db.add(role)
        db.flush()
    else:
        role.role_name = name
        role.description = description
    return role


def seed_roles_permissions(db: Session) -> dict[str, SystemRole]:
    roles = {code: get_or_create_role(db, code, name, description) for code, name, description in ROLE_DEFINITIONS}

    permissions = {}
    for code, name, group in PERMISSION_DEFINITIONS:
        permission = db.query(SystemPermission).filter(SystemPermission.permission_code == code).first()
        if not permission:
            permission = SystemPermission(permission_code=code, permission_name=name, permission_group=group)
            db.add(permission)
            db.flush()
        else:
            permission.permission_name = name
            permission.permission_group = group
        permissions[code] = permission

    role_permission_codes = {
        "ADMIN": list(permissions.keys()),
        "HEAD_OFFICE_LEADER": [
            "dashboard:view",
            "dashboard:drilldown",
            "dashboard:export",
            "dashboard:health",
            "customer:view", "customer:profile:view", "customer:deposit:view", "customer:credit:view", "customer:income:view", "customer:export",
            "analytics:view", "analytics:export",
            "report:view",
            "report:summarize",
        ],
        "BRANCH_MANAGER": [
            "dashboard:view", "dashboard:drilldown", "dashboard:export", "dashboard:health",
            "customer:view", "customer:profile:view", "customer:deposit:view", "customer:credit:view", "customer:income:view", "customer:export",
            "analytics:view", "analytics:export",
            "report:view", "report:summarize", "report:export",
        ],
        "DEPARTMENT_MANAGER": [
            "dashboard:view", "dashboard:drilldown", "dashboard:export",
            "customer:view", "customer:profile:view", "customer:deposit:view", "customer:credit:view", "customer:income:view", "customer:export",
            "analytics:view", "analytics:export", "report:view", "report:export",
        ],
        "USER": ["dashboard:view", "dashboard:drilldown", "customer:view", "customer:profile:view", "customer:deposit:view", "customer:credit:view", "customer:income:view", "analytics:view", "report:view"],
    }
    for role_code, permission_codes in role_permission_codes.items():
        role = roles[role_code]
        desired_permission_ids = {permissions[code].id for code in permission_codes}
        db.query(SystemRolePermission).filter(
            SystemRolePermission.role_id == role.id,
            ~SystemRolePermission.permission_id.in_(desired_permission_ids),
        ).delete(synchronize_session=False)
        for permission_code in permission_codes:
            exists = (
                db.query(SystemRolePermission)
                .filter(
                    SystemRolePermission.role_id == role.id,
                    SystemRolePermission.permission_id == permissions[permission_code].id,
                )
                .first()
            )
            if not exists:
                db.add(SystemRolePermission(role_id=role.id, permission_id=permissions[permission_code].id))

    return roles


def upsert_branch(db: Session, branch_code: str, branch_name: str) -> OrgBranch:
    branch = db.query(OrgBranch).filter(OrgBranch.branch_code == branch_code).first()
    if not branch:
        branch = OrgBranch(branch_code=branch_code, branch_name=branch_name or f"Chi nhánh {branch_code}")
        db.add(branch)
        db.flush()
    else:
        branch.branch_name = branch_name or branch.branch_name
        branch.status = "active"
    return branch


def upsert_department(db: Session, branch: OrgBranch, department_code: str, department_name: str) -> OrgDepartment:
    department = (
        db.query(OrgDepartment)
        .filter(OrgDepartment.branch_id == branch.id, OrgDepartment.department_code == department_code)
        .first()
    )
    if not department:
        department = OrgDepartment(
            branch_id=branch.id,
            department_code=department_code,
            department_name=department_name or f"Phòng ban {department_code}",
        )
        db.add(department)
        db.flush()
    else:
        department.department_name = department_name or department.department_name
        department.status = "active"
    return department


def seed_organization(db: Session, default_password_hash: str) -> None:
    branch_rows = read_excel_rows(DOCUMENTS_DIR / "Danh sach chi nhanh.xlsx")
    branches: dict[str, OrgBranch] = {}
    departments: dict[tuple[str, str], OrgDepartment] = {}

    for row in branch_rows:
        branch_code = norm_code(row.get("BRCD"))
        branch_name = norm_text(row.get("Name"))
        department_code = norm_code(row.get("TRSTCD"))
        department_name = norm_text(row.get("TRSTName"))
        if not branch_code:
            continue
        branch = upsert_branch(db, branch_code, branch_name)
        branches[branch_code] = branch
        if department_code:
            departments[(branch_code, department_code)] = upsert_department(db, branch, department_code, department_name)

    staff_rows = read_excel_rows(DOCUMENTS_DIR / "Danh sách chi tiết cán bộ.xlsx")
    role_user = db.query(SystemRole).filter(SystemRole.role_code == "USER").first()

    for row in staff_rows:
        branch_code = norm_code(row.get("BRCD"))
        branch = branches.get(branch_code)
        if not branch and branch_code:
            branch = upsert_branch(db, branch_code, f"Chi nhánh {branch_code}")
            branches[branch_code] = branch

        department_code = norm_code(row.get("TRSTCD"))
        department_name = norm_text(row.get("TRSTName"))
        department = None
        if branch and department_code:
            department = departments.get((branch_code, department_code))
            if not department:
                department = upsert_department(db, branch, department_code, department_name)
                departments[(branch_code, department_code)] = department

        employee_code = norm_code(row.get("Mã NV"))
        ipcas_username = norm_text(row.get("User IPCAS")).upper()
        full_name = norm_text(row.get("Tên NV"))
        if not employee_code and not ipcas_username:
            continue

        username = (employee_code or ipcas_username).lower()
        user = db.query(SystemUser).filter(SystemUser.username == username).first()
        if not user and employee_code:
            user = db.query(SystemUser).filter(SystemUser.employee_code == employee_code).first()

        payload = {
            "username": username,
            "employee_code": employee_code or None,
            "credit_officer_code": norm_code(row.get("Mã CBTD")) or None,
            "ipcas_username": ipcas_username or None,
            "full_name": full_name or username,
            "branch_id": branch.id if branch else None,
            "department_id": department.id if department else None,
            "is_active": True,
        }

        if not user:
            user = SystemUser(
                password_hash=default_password_hash,
                role_id=role_user.id if role_user else None,
                **payload,
            )
            db.add(user)
        else:
            for key, value in payload.items():
                setattr(user, key, value)
            # Đồng bộ danh mục cán bộ không được xóa vai trò đã phân công hoặc
            # đặt lại mật khẩu mỗi lần backend khởi động.


def seed_admin(db: Session, roles: dict[str, SystemRole], default_password_hash: str) -> None:
    admin = db.query(SystemUser).filter(SystemUser.username == "admin").first()
    if not admin:
        admin = SystemUser(
            username="admin",
            password_hash=default_password_hash,
            full_name="Quản trị hệ thống C360",
            role_id=roles["ADMIN"].id,
            is_active=True,
            is_superuser=True,
        )
        db.add(admin)
    else:
        admin.full_name = "Quản trị hệ thống C360"
        admin.role_id = roles["ADMIN"].id
        admin.is_active = True
        admin.is_superuser = True
        admin.password_hash = default_password_hash


def seed_initial_data() -> None:
    db = SessionLocal()
    try:
        default_password_hash = hash_password(DEFAULT_PASSWORD)
        roles = seed_roles_permissions(db)
        seed_admin(db, roles, default_password_hash)
        seed_organization(db, default_password_hash)
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()
