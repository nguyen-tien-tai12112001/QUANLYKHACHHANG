import csv
import codecs
import re
from pathlib import Path

import openpyxl
import polars as pl


DEFAULT_CHUNK_SIZE = 20000
UTF8_BOM = b"\xef\xbb\xbf"
BOM_MARKERS = ("\ufeff", "ï»¿", "Ï»¿")
ZERO_WIDTH_CHARS = ("\u200b", "\u200c", "\u200d", "\u2060")


def normalize_header(header) -> str:
    text = str(header or "")
    for marker in BOM_MARKERS:
        text = text.replace(marker, "")
    for marker in ZERO_WIDTH_CHARS:
        text = text.replace(marker, "")
    text = text.replace("\xa0", " ")
    text = text.strip().strip('"').strip("'").strip()
    text = re.sub(r"\s+", "_", text)
    return text.upper()


def detect_delimiter(sample: str) -> str:
    if not sample:
        return ","
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def detect_csv_encoding(file_path: str | Path) -> str:
    sample = Path(file_path).read_bytes()[:65536]
    if sample.startswith(UTF8_BOM):
        return "utf-8-sig"

    encodings = ("utf-8-sig", "utf-8", "cp1258", "latin-1")
    last_error = None
    for encoding in encodings:
        try:
            decoder = codecs.getincrementaldecoder(encoding)("strict")
            decoder.decode(sample, final=False)
            return encoding
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"Không đọc được encoding CSV: {last_error}")


def detect_csv_dialect(file_path: str | Path, encoding: str) -> tuple[str, int]:
    path = Path(file_path)
    with path.open("r", encoding=encoding, errors="replace", newline="") as file:
        first_line = file.readline()
        clean_first_line = first_line.lstrip("\ufeff").strip()
        if clean_first_line.lower().startswith("sep=") and len(clean_first_line) >= 5:
            return clean_first_line[4], 1
        sample = first_line + file.read(8192)
        return detect_delimiter(sample), 0


def read_csv_rows(file_path: str | Path) -> list[dict]:
    path = Path(file_path)
    encoding = detect_csv_encoding(path)
    delimiter, skip_rows = detect_csv_dialect(path, encoding)

    # Keep every column as text so codes, account numbers and leading zeroes
    # stay exactly as the source file provides them.
    df = pl.read_csv(
        path,
        encoding=encoding,
        separator=delimiter,
        skip_rows=skip_rows,
        infer_schema_length=0,
        null_values=[""],
        ignore_errors=False,
        truncate_ragged_lines=False,
    )

    df = df.rename({column: normalize_header(column) for column in df.columns})
    df = df.drop([column for column in df.columns if not column])
    return df.to_dicts()


def iter_csv_row_chunks(file_path: str | Path, chunk_size: int = DEFAULT_CHUNK_SIZE):
    path = Path(file_path)
    encoding = detect_csv_encoding(path)
    delimiter, skip_rows = detect_csv_dialect(path, encoding)

    with path.open("r", encoding=encoding, errors="replace", newline="") as file:
        for _ in range(skip_rows):
            next(file, None)
        reader = csv.reader(file, delimiter=delimiter)
        try:
            headers = [normalize_header(value) for value in next(reader)]
        except StopIteration:
            return

        rows = []
        for raw_row in reader:
            item = {}
            has_value = False
            for index, header in enumerate(headers):
                if not header:
                    continue
                value = raw_row[index] if index < len(raw_row) else None
                if value is not None and str(value).strip() != "":
                    has_value = True
                item[header] = value
            if not has_value:
                continue
            rows.append(item)
            if len(rows) >= chunk_size:
                yield rows
                rows = []
        if rows:
            yield rows


def read_excel_rows(file_path: str | Path) -> list[dict]:
    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    worksheet = workbook.worksheets[0]
    rows = worksheet.iter_rows(values_only=True)
    headers = [normalize_header(value) for value in next(rows)]

    data = []
    for row in rows:
        item = {}
        has_value = False
        for index, header in enumerate(headers):
            if not header:
                continue
            value = row[index] if index < len(row) else None
            if value is not None and str(value).strip() != "":
                has_value = True
            item[header] = value
        if has_value:
            data.append(item)

    workbook.close()
    return data


def iter_excel_row_chunks(file_path: str | Path, chunk_size: int = DEFAULT_CHUNK_SIZE):
    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = [normalize_header(value) for value in next(rows)]
        except StopIteration:
            return

        data = []
        for row in rows:
            item = {}
            has_value = False
            for index, header in enumerate(headers):
                if not header:
                    continue
                value = row[index] if index < len(row) else None
                if value is not None and str(value).strip() != "":
                    has_value = True
                item[header] = value
            if not has_value:
                continue
            data.append(item)
            if len(data) >= chunk_size:
                yield data
                data = []
        if data:
            yield data
    finally:
        workbook.close()


def read_file_rows(file_path: str | Path) -> list[dict]:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".csv":
        return read_csv_rows(file_path)
    if suffix == ".xlsx":
        return read_excel_rows(file_path)
    raise ValueError("Định dạng file chưa được hỗ trợ")


def iter_file_row_chunks(file_path: str | Path, chunk_size: int = DEFAULT_CHUNK_SIZE):
    suffix = Path(file_path).suffix.lower()
    if suffix == ".csv":
        yield from iter_csv_row_chunks(file_path, chunk_size=chunk_size)
        return
    if suffix == ".xlsx":
        yield from iter_excel_row_chunks(file_path, chunk_size=chunk_size)
        return
    raise ValueError("Định dạng file chưa được hỗ trợ")
